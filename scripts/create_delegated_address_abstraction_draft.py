from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documents" / "PairCT_address_abstraction_draft02.docx"


TITLE = "PairCT: Privacy-Preserving Address Abstraction With Identity-Level Traceability"

ABSTRACT = (
    "Blockchain users are commonly identified through chain-specific addresses, resulting in fragmented "
    "accounts across blockchains and services. Reusing an address improves continuity but enables "
    "cross-context linkability, whereas creating independent addresses increases management and recovery "
    "burdens. Existing account-abstraction systems improve transaction programmability, and "
    "credential-based address-abstraction schemes reduce address fragmentation, but they do not jointly "
    "provide identity-provider-assisted authentication, identity-level accountability, and revocable "
    "session-based execution. We present PairCT, a privacy-preserving address-abstraction framework that "
    "derives RP- and "
    "chain-scoped smart accounts from a common identity-provider-authenticated identity. For a relying-party "
    "identifier and a namespace-qualified chain context, PairCT derives a pairwise pseudonymous identifier "
    "without revealing the user's identity-provider-side identifier. A zero-knowledge proof binds the derived "
    "identifier, the target context, an identity-provider-signed authentication statement, and a "
    "wallet-generated session key. This construction enables one IdP-authenticated identity to authorize "
    "accounts across supported relying parties and blockchains while avoiding a reusable global address. "
    "PairCT further provides identity-level traceability: when opening is authorized for a disputed session, "
    "retained authentication records can associate the corresponding account with its identity-provider-scoped "
    "identity (iss, uid), while ordinary activity remains pseudonymous. Finally, session-bound authorization "
    "amortizes authentication and proof-generation costs across multiple transactions and supports selective "
    "session or account revocation through a Merkle-accumulator state anchored on-chain. PairCT therefore "
    "combines address abstraction, accountable pseudonymity, and efficient session lifecycle "
    "management in a unified Web3 authentication framework."
)

INDEX_TERMS = (
    "Address abstraction, identity-provider-assisted authentication, identity-level traceability, pairwise pseudonymous "
    "identifier, session revocation, zero-knowledge proofs."
)

INTRODUCTION = [
    (
        "Web3 users interact with services through blockchain accounts whose addresses and authorization "
        "rules are tied to particular chains. A user who accesses several services on several chains may "
        "therefore accumulate many addresses, signing keys, recovery procedures, and service-side account "
        "records. Reusing one address reduces this management burden, but exposes a stable public identifier "
        "that permits activity correlation across services and chains. Generating a fresh address for every "
        "context improves separation, but leaves the user and each service to reconstruct account continuity "
        "on top of unrelated blockchain identifiers."
    ),
    (
        "Account abstraction addresses part of this problem by moving authorization logic from a fixed "
        "externally owned account into a programmable smart account. This enables recovery policies, "
        "sponsored execution, batched calls, and alternative signature checks. It does not by itself determine "
        "how one authenticated user should obtain accounts across different relying parties (RPs) and chains, "
        "nor how those accounts should remain unlinkable outside their intended contexts. Chain abstraction "
        "similarly simplifies cross-chain interaction, but commonly retains chain-specific addresses and "
        "cryptographic authorization underneath the abstraction layer."
    ),
    (
        "Credential-based address abstraction takes a different approach: a certificate or credential acts as "
        "the root from which blockchain-facing identifiers are derived. Existing constructions show that this "
        "can reduce address fragmentation, particularly when the credential is issued once and subsequently "
        "presented in a self-sovereign manner. Identity-provider-assisted Web authentication follows a "
        "delegated-authentication model with a different lifecycle. An "
        "identity provider (IdP) authenticates the user at login time and issues a fresh signed statement for "
        "the current session. This model benefits from established authentication and account-recovery "
        "infrastructure, but introduces privacy and accountability questions that an issuer-offline credential "
        "model does not directly resolve."
    ),
    (
        "An IdP-assisted address-abstraction protocol must reconcile three requirements. First, one authenticated "
        "identity should be sufficient to derive and authorize accounts across supported RPs and chains without "
        "requiring a new long-term blockchain identity for every context. Second, the resulting convenience "
        "must not create a globally reusable identifier that allows RPs, chains, or public observers to link "
        "the user's activity. Third, pseudonymity must not eliminate accountability: when misuse is reported "
        "and opening is authorized, the system should identify the IdP-scoped account responsible for the "
        "disputed session rather than returning only another opaque blockchain address."
    ),
    (
        "Efficiency introduces a further requirement. If every blockchain transaction repeats IdP "
        "authentication and zero-knowledge proof generation, the authentication layer becomes the dominant "
        "cost of ordinary account use. A session mechanism can amortize this cost by binding one successful "
        "IdP authentication to a wallet-generated session key that authorizes multiple transactions "
        "within a bounded lifetime. Such a session must be revocable before natural expiration when its key is "
        "lost, compromised, or otherwise disputed, without forcing unrelated sessions or all accounts derived "
        "from the same identity to be disabled."
    ),
    (
        "We present PairCT, a privacy-preserving address-abstraction framework. PairCT begins from an "
        "IdP-authenticated identity and derives a pairwise identifier for a specific RP and namespace-qualified "
        "chain context. Conceptually, PPID_(r,c) = H(domain, iss, uid, rid, chainContext_c, salt), where iss and "
        "uid identify the IdP account, rid identifies the RP, and the wallet-held salt prevents public "
        "recomputation. Domain separation and the explicit chain context ensure that the same authenticated "
        "identity produces deterministic but distinct identifiers across RPs and chains. Each supported chain "
        "maps its scoped identifier to a chain-local programmable account; it need not expose one universal "
        "address across heterogeneous ledgers."
    ),
    (
        "A zero-knowledge relation connects this scoped identifier to an IdP-signed authentication statement "
        "and a wallet-generated session public key without revealing uid or the wallet salt to the RP or the "
        "blockchain. The resulting smart account accepts execution authorized by the session key only while "
        "the corresponding authentication statement remains valid. PairCT thus abstracts authorization away "
        "from possession of a reusable EOA key: control of the chain-local account follows from a current "
        "IdP-authenticated session whose context and expiration are cryptographically bound."
    ),
    (
        "PairCT couples this pseudonymous execution model with identity-level traceability. Ordinary "
        "authentication and transaction execution reveal only context-scoped identifiers. For a specifically "
        "disputed session, an externally authorized opening procedure joins the retained RP transcript with "
        "the corresponding IdP authentication record and returns the responsible IdP-scoped identity (iss, "
        "uid). Opening is therefore session-scoped rather than a general discovery interface, and the protocol "
        "does not give an IdP or RP an unrestricted capability to enumerate a user's accounts across contexts."
    ),
    (
        "To support early invalidation, PairCT represents revoked sessions and accounts in a Merkle-based "
        "accumulator whose compact root is anchored on-chain. A wallet proves that the current session and "
        "account are absent from the applicable revocation set, while the underlying revocation identifiers "
        "remain hidden witnesses. Session entries expire with their authentication window and can be pruned, "
        "keeping the active structure proportional to recent exceptional revocations rather than to all users "
        "or all previously issued sessions."
    ),
    (
        "The contributions of this work are threefold. First, PairCT provides identity-level traceability for "
        "pseudonymous smart accounts by linking an authorized disputed-session opening to its IdP-scoped "
        "identity. Second, it introduces privacy-preserving address abstraction, allowing a common "
        "IdP-authenticated identity to derive deterministic RP- and chain-scoped accounts without exposing a "
        "globally reusable identifier. Third, it provides session-bound execution that amortizes IdP "
        "authentication and proof costs across transactions, together with selective and efficiently "
        "maintainable session and account revocation. The construction targets supported programmable "
        "blockchains and does not assume that one address format or one smart-contract implementation is "
        "natively portable to every ledger."
    ),
]

CONCLUSION = [
    (
        "This paper presented PairCT as a privacy-preserving address-abstraction framework for Web3 "
        "authentication. Rather than treating a blockchain address or a long-term EOA key as the user's "
        "identity, PairCT derives deterministic RP- and chain-scoped smart accounts from a common "
        "IdP-authenticated identity. The construction preserves account continuity within an intended context "
        "while avoiding a globally reusable address that would make activity directly linkable across RPs and "
        "chains."
    ),
    (
        "PairCT combines this abstraction with identity-level traceability and session-bound authorization. "
        "Ordinary execution remains pseudonymous, whereas an externally authorized opening of a specific "
        "disputed session can identify the corresponding IdP-scoped account (iss, uid). A successful IdP "
        "authentication is bound to a wallet-generated session key, allowing its cost to be amortized across "
        "multiple transactions within a bounded lifetime. The Merkle-based revocation mechanism supports early "
        "invalidation of a compromised session or account while keeping the on-chain state compact and allowing "
        "expired revocation entries to be pruned."
    ),
    (
        "The proposed abstraction is intentionally context-scoped: it provides one authentication foundation "
        "for supported services and chains, not one publicly reusable identifier or one universally identical "
        "address. Future work includes formalizing the new cross-context unlinkability and opening definitions, "
        "implementing namespace-qualified derivation across heterogeneous smart-contract platforms, measuring "
        "the amortized session cost under realistic workloads, and evaluating revocation publication, witness "
        "distribution, and root-freshness policies under production conditions."
    ),
]


def set_columns(section, count):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols_element = cols[0]
    else:
        cols_element = OxmlElement("w:cols")
        sect_pr.append(cols_element)
    cols_element.set(qn("w:num"), str(count))
    if count > 1:
        cols_element.set(qn("w:space"), "360")


def configure_section(section, columns):
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    set_columns(section, columns)


def set_run_font(run, size=10, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_body_paragraph(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.16)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    set_run_font(paragraph.add_run(text), 9.5)
    return paragraph


def add_section_heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text), 10, bold=False)
    return paragraph


def build_document():
    document = Document()
    configure_section(document.sections[0], 1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(9.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run(TITLE), 20)

    abstract = document.add_paragraph()
    abstract.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract.paragraph_format.left_indent = Inches(0.45)
    abstract.paragraph_format.right_indent = Inches(0.45)
    abstract.paragraph_format.space_after = Pt(4)
    set_run_font(abstract.add_run("ABSTRACT "), 9, bold=True)
    set_run_font(abstract.add_run(ABSTRACT), 9)

    index_terms = document.add_paragraph()
    index_terms.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    index_terms.paragraph_format.left_indent = Inches(0.45)
    index_terms.paragraph_format.right_indent = Inches(0.45)
    index_terms.paragraph_format.space_after = Pt(8)
    set_run_font(index_terms.add_run("INDEX TERMS "), 9, bold=True)
    set_run_font(index_terms.add_run(INDEX_TERMS), 9, italic=True)

    body_section = document.add_section(WD_SECTION.CONTINUOUS)
    configure_section(body_section, 2)

    add_section_heading(document, "I. INTRODUCTION")
    for paragraph in INTRODUCTION:
        add_body_paragraph(document, paragraph)

    add_section_heading(document, "II. CONCLUSION")
    for paragraph in CONCLUSION:
        add_body_paragraph(document, paragraph)

    document.core_properties.title = TITLE
    document.core_properties.subject = "Privacy-preserving address abstraction research draft"
    document.core_properties.keywords = INDEX_TERMS
    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
