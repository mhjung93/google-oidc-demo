from docx import Document
F = 'documents/PairCT_research_article_20260706_092133_with_figures_final47.docx'
d = Document(F)
P = d.paragraphs

def find(sub, style=None):
    for i,p in enumerate(P):
        if sub in p.text and (style is None or p.style.name == style):
            return i
    raise SystemExit(f'못 찾음: {sub}')

# ── 1) 프로토콜 절 G: 폐기 ────────────────────────────────────────────────
anchor = find('SECURITY AND PRIVACY ANALYSIS', 'H1_List (Space)')
blocks = [
 ('Normal', 'G. CREDENTIAL REVOCATION'),
 ('PARA',
  'Expiration alone is not sufficient. A signed login statement stays usable until max_height, so a leaked session '
  'key cannot be cut off before that window closes, and an account that must be blocked keeps whatever statements '
  'it already holds. The IdP can refuse future issuance, but it cannot recall what it has already signed. What makes '
  'this harder than in the off-chain setting is the verifier: for the address-abstraction path of Section VI-C the '
  'verifier is a smart contract, which can neither hold a revocation list nor be handed one per transaction without '
  'revealing which entry the user is checking.'),
 ('PARA',
  'PairCT therefore proves non-membership in a revocation set inside the same Groth16 relation that already proves '
  'key binding, and publishes only a commitment to that set on-chain. Two layers are revoked independently. A '
  'session leaf H(TAG_SESSION, r_token) withdraws one issued statement; an account leaf H(TAG_ACCOUNT, auid) '
  'withdraws every statement of one account, where auid is the same per-account value already used for the '
  'one-active-binding check of Section IV-B. Because both leaves are hashes of hidden witnesses and the proof shows '
  'their absence rather than their presence, a non-revoked user reveals nothing: the anonymity set is every account '
  'that has not been revoked, including accounts that never authenticated through this IdP.'),
 ('PARA',
  'The revocation set is held as indexed Merkle trees supporting O(log n) non-membership witnesses. A single tree '
  'would make every proof depend on one root, so any user’s revocation would invalidate every other user’s proof '
  'the moment it was published. PairCT splits the structure instead. The lower level is a forest of subtrees, '
  'Poseidon-hashed and verified inside the relation; the upper level commits to the subtree roots with keccak and is '
  'verified by the contract, outside the relation. A proof is then bound only to its own subtree root, so a '
  'revocation in any other subtree leaves it valid and only the upper-level sibling hashes—ordinary calldata—have '
  'to be refreshed. The two layers are sharded on different keys: session leaves by the credential’s own '
  'max_height, which is already a public signal, and account leaves by low-order bits of the account leaf, because '
  'the expiry of an account revocation is by construction not known to the prover.'),
 ('PARA',
  'Revocations are batched. The IdP queues them, computes the next root, the operator publishes that root, and only '
  'then does the IdP advance its served state; the registry accepts a superseded root for a bounded number of '
  'blocks so that transactions already in flight are not invalidated mid-submission. Revocation latency is therefore '
  'bounded by that grace window rather than by the batching period, and it does not depend on how often revocations '
  'occur. The cost is symmetric and deliberate: a revoked credential remains usable for at most that many blocks.'),
 ('PARA',
  'Two boundaries should be stated plainly. First, publication introduces a standing role that the opening path does '
  'not: whoever is authorized to update the on-chain registry can withhold an update, republish a superseded root '
  'and thereby reinstate revoked credentials, or publish an unusable root and block every wallet. The contract '
  'checks only that a root is current, not that it was derived correctly, so this authority is trusted for '
  'revocation integrity; making the update itself verifiable is left as future work. Second, the account shard index '
  'is a public signal, and auid does not depend on rid, so the same account exposes the same index at every RP. '
  'This is a bounded weakening of cross-service unlinkability rather than a break, and Property 5 is stated with '
  'that carve-out; the shard count is kept small for exactly this reason, which is affordable because account-level '
  'revocations are far rarer than session-level ones.'),
]
for style, text in blocks:
    P[anchor].insert_paragraph_before(text, style=style)
print('OK 절 G 삽입 (문단 %d 앞)' % anchor)

# ── 2) Property 9 ─────────────────────────────────────────────────────────
P = d.paragraphs
anchor2 = find('IMPLEMENTATION AND SCOPE', 'H1_List (Space)')
blocks2 = [
 ('Normal', 'N. REVOCATION SOUNDNESS'),
 ('PARA',
  'Scope. Revocation soundness concerns what a revoked credential can still do once the corresponding root has been '
  'published and the grace window of Section G has elapsed. It does not claim that revocation is instantaneous, and '
  'it does not cover an issuer that declines to revoke or to publish; those are availability assumptions on the '
  'issuer and the publisher, stated in Section G and revisited in the limitations.'),
 ('PARA',
  'Property 9 (Revocation soundness). Under the stated system and security model, for every PPT adversary A,'),
 ('PARA',
  'Adv_revoke(A) <= Adv_G16^ksnd(B_g16) + Adv_H^coll(B_coll) + negl(lambda),'),
 ('PARA',
  'where Adv_revoke(A) is the probability that A produces an accepted on-chain execution whose session or account '
  'leaf is contained in the revocation state committed to by a root the verifier accepts.'),
 ('PARA',
  'Argument summary. Acceptance requires a Groth16 proof for a relation that includes two non-membership statements '
  'against subtree roots supplied as public signals, and the contract independently recomputes the published '
  'commitment from those roots and the shard indices before verifying the proof. By knowledge soundness, an accepted '
  'proof yields witnesses satisfying the non-membership relation for the committed subtrees; by collision resistance '
  'of the hash, no such witness exists for a leaf that the issuer has inserted, since producing one would require a '
  'second preimage on the Merkle path. Two conditions carry the argument and are enforced rather than assumed: the '
  'relation binds the shard index to the leaf, so a prover cannot present an empty subtree in place of its own, and '
  'the contract derives the session shard from max_height itself rather than accepting a caller-supplied index. '
  'Absent either, the non-membership statement would be vacuous.'),
]
for style, text in blocks2:
    P[anchor2].insert_paragraph_before(text, style=style)
print('OK Property 9 삽입 (문단 %d 앞)' % anchor2)

d.save(F)
print('저장:', F)
