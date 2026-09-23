#!/usr/bin/env python3
"""Build a deck on the BAAR-based revocation work: 폐기 기준, BAAR와의 차이,
root를 온체인에 두는 이유, 1억 사용자 확장성, 그리고 그 분석에서 이어진
설계 재검토(크레덴셜 폐기 + ephemeral 트리).

Reuses the POSTECH template from 260731_meeting_MHJ.pptx, same as
build_meeting_slides_260818.py.

All constraint counts in this deck are measured, not estimated — see
"측정 방법" slide for the exact commands.
"""

import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(SCRIPT_DIR)

SRC = os.path.join(REPO_ROOT, "documents", "260731_meeting_MHJ.pptx")
LOGO = os.path.join(SCRIPT_DIR, "postech_logo.png")
OUT = os.path.join(REPO_ROOT, "documents", "260826_revocation_baar_scaling.pptx")

DARK = RGBColor(0x44, 0x54, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(SRC)
LAYOUT_TITLE = prs.slide_layouts[0]
LAYOUT_CONTENT = prs.slide_layouts[1]
LAYOUT_TITLE_ONLY = prs.slide_layouts[7]

xml_slides = prs.slides._sldIdLst
for sld_id in list(xml_slides):
    rId = sld_id.get(qn('r:id'))
    prs.part.drop_rel(rId)
    xml_slides.remove(sld_id)


def add_bullets(tf, items):
    tf.clear()
    first = True
    for item in items:
        level, text = item[0], item[1]
        bold = item[2] if len(item) > 2 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.text = text
        for run in p.runs:
            run.font.bold = bold


def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(LAYOUT_TITLE)
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    slide.shapes.add_picture(LOGO, Emu(3971764), Emu(4951853), Emu(4248472), Emu(611893))
    return slide


def add_bullet_slide(title, items):
    slide = prs.slides.add_slide(LAYOUT_CONTENT)
    slide.shapes.title.text = title
    add_bullets(slide.placeholders[1].text_frame, items)
    return slide


def add_table_slide(title, headers, rows, col_widths, top=Inches(1.5), height=Inches(4.5),
                    font_size=13, header_size=13, left=Inches(0.5), width=Inches(12.3),
                    note=None):
    slide = prs.slides.add_slide(LAYOUT_TITLE_ONLY)
    slide.shapes.title.text = title
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gframe = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = gframe.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = w
    for c_i, h in enumerate(headers):
        table.cell(0, c_i).text = h
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            table.cell(r_i, c_i).text = val
    for r_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            for p in tf.paragraphs:
                sz = header_size if r_i == 0 else font_size
                p.font.size = Pt(sz)
                if r_i == 0:
                    p.font.bold = True
                    p.font.color.rgb = WHITE
                for run in p.runs:
                    run.font.size = Pt(sz)
                    if r_i == 0:
                        run.font.bold = True
                        run.font.color.rgb = WHITE
            if r_i == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = DARK
    if note:
        box = slide.shapes.add_textbox(left, top + height + Inches(0.15), width, Inches(0.9))
        tf = box.text_frame
        tf.word_wrap = True
        add_bullets(tf, note)
        for p in tf.paragraphs:
            p.font.size = Pt(12)
            for run in p.runs:
                run.font.size = Pt(12)
    return slide


# ============================================================
# 1. Title
# ============================================================
add_title_slide(
    "BAAR 기반 Revocation —\n폐기 기준, 확장성, 설계 재검토",
    "\n무엇을 기준으로 폐기하는가 · BAAR와의 차이 · 1억 규모 비용 · 크레덴셜 폐기로의 전환\nMoonhyeon Chung, POSTECH",
)

# ============================================================
# 2. 배경
# ============================================================
add_bullet_slide("배경 — 이번에 구현한 것", [
    (0, "이전 상태: revocation 미구현. max_height 만료뿐 (폐기 분류 ③번 단기 크레덴셜, 네 범주 중 가장 약함)", True),
    (0, "구현: IdP가 폐기 리프의 Indexed Merkle Tree를 관리하고 root만 온체인에 게시", True),
    (1, "pi_pk_i 회로가 세션과 계정 둘 다에 대해 비멤버십을 증명"),
    (1, "PPIDWallet.execute()가 증명에 쓰인 root의 registry 등록 여부를 확인"),
    (0, "세 계층이 각각 검증됨", True),
    (1, "IdP 폐기 → witness 발급 불가 / 회로가 폐기된 세션·계정을 거부 / 온체인이 만료 root를 거부"),
    (0, "전 구간 온체인 E2E 통과 — 실제 6-signal 증명이 배포된 PPIDWallet을 통과 (gas 297,837)", True),
    (0, "그런데 확장성 분석 과정에서 설계 자체를 다시 봐야 할 지점이 드러났다 (후반부)", True),
])

# ============================================================
# 3. 폐기 기준
# ============================================================
add_bullet_slide("폐기 기준 — 무엇을 리프로 삼는가", [
    (0, "두 종류가 한 트리에 도메인 태그로 구분되어 들어간다", True),
    (1, "세션 리프 = Poseidon(TAG_SESSION, r_token),  r_token = Poseidon(pk_i, max_height, rp_nonce)"),
    (1, "계정 리프 = Poseidon(TAG_ACCOUNT, auid),      auid = Poseidon(uid, salt)"),
    (0, "r_token — (그 로그인, 그 RP) 하나를 정확히 지목", True),
    (1, "세션키·유효창·RP 논스를 묶은 값 → 유출된 세션 하나만 끊을 수 있다"),
    (1, "계정 폐기만 있으면 세션 하나 유출에 그 사용자의 모든 RP·기기를 통째로 죽여야 함"),
    (0, "auid — 계정 전체, RP와 무관", True),
    (1, "rid가 안 들어가므로 어느 RP에서든 같은 값 → 리프 하나로 전역 폐기"),
    (0, "PPID = Poseidon(uid, rid, salt)는 리프로 쓰이지 않는다", True),
    (1, "IdP가 PPID를 모르고(rid·salt·ppid 전부 미지), RP별 값이라 전역 폐기도 불가"),
])

# ============================================================
# 4. 왜 그 값이어야 하는가
# ============================================================
add_table_slide(
    "왜 그 값이어야 하는가 — 아는 쪽이 갈려 있다",
    ["값", "IdP가 아는가", "체인이 아는가", "역할"],
    [
        ["r_token", "○ (token_nonce로 공개)", "✕ (pi_pk_i에서 private)", "세션 폐기 대상"],
        ["auid", "○ (로그인 시 지갑이 제시)", "✕ (회로 내부 유도)", "계정 폐기 대상"],
        ["pk_i", "✕ (pi_arid_i에서 private)", "○ (ecrecover에 필요)", "세션키 — 폐기 대상 불가"],
        ["PPID", "✕", "○ (지갑 주소 유도)", "지갑 식별 — 폐기 대상 불가"],
    ],
    [Inches(1.6), Inches(3.6), Inches(3.6), Inches(3.5)],
    height=Inches(2.6), font_size=13,
    note=[
        (0, "IdP가 아는 것과 체인이 아는 것이 의도적으로 겹치지 않는다 — 폐기 대상 둘이 모두 \"IdP만 아는\" 쪽에 있는 건 우연이 아니다."),
        (0, "pk_i를 리프로 삼으면: IdP가 값을 모를 뿐 아니라, 체인의 pk_i로 Poseidon(TAG, pk_i)를 계산해 공개 폐기 목록과 대조하면"),
        (0, "어느 온체인 세션이 폐기됐는지 누구나 알아낼 수 있다. r_token은 preimage에 rp_nonce가 있어 이 대조가 불가능하다."),
    ],
)

# ============================================================
# 5. BAAR 원본
# ============================================================
add_bullet_slide("BAAR 원본 구조", [
    (0, "Ahmed, Ahmad, Zeshan, Akram — PLOS ONE 21(3): e0343696 (2026-03-31)", True),
    (0, "구성 요소", True),
    (1, "Pedersen vector commitment — 속성 벡터를 C = Σ(aᵢ·Gᵢ) + r·H 로 커밋"),
    (1, "Schnorr ZKP — Σ-protocol을 Fiat-Shamir로 비대화형화, 선택적 속성 공개"),
    (1, "Merkle 동적 accumulator — 크레덴셜 식별자를 리프에 매핑, accumulator 값 = root"),
    (0, "폐기 방식: 유효 집합에서 리프를 제거 → 갱신된 root에 대해 기존 witness가 실패", True),
    (0, "데이터 흐름: 오프체인 증명 + 온체인 컴팩트 root — 개별 커밋먼트는 온체인에 저장하지 않음", True),
    (0, "보고된 성능: 검증 ~200ms, 폐기 200–310ms(5,000개까지), 폐기 gas ~30K, 검증 gas ~60K", True),
])

# ============================================================
# 6. 비교표
# ============================================================
add_table_slide(
    "무엇을 바꿨나",
    ["", "BAAR", "본 구현"],
    [
        ["누적기 의미", "유효 집합 멤버십\n(폐기 = 리프 제거)", "폐기 집합 비멤버십\n(폐기된 것만 트리에)"],
        ["증명 시스템", "Pedersen 커밋 + Schnorr Σ-protocol", "PS 서명 + Groth16"],
        ["Merkle 경로 검증", "회로 밖 (트랜스크립트에 witness 동봉)", "회로 안 (circom 제약)"],
        ["트리 종류", "일반 Merkle 동적 accumulator", "Indexed Merkle Tree\n(정렬 + low < target < next)"],
        ["검증자", "리스트 보유 가능", "스마트컨트랙트 — 리스트 보유 불가"],
        ["폐기 단위", "크레덴셜 1종", "세션 + 계정 동시\n(도메인 태그 분리, 한 트리·한 root)"],
        ["신선도 정책", "최신 root", "K개 슬롯 + 블록 기반 만료\n(GRACE_BLOCKS)"],
    ],
    [Inches(2.0), Inches(4.8), Inches(5.5)],
    height=Inches(4.6), font_size=12,
)

# ============================================================
# 7. 방향을 뒤집은 이유
# ============================================================
add_bullet_slide("가장 큰 차이 — 누적기의 방향을 뒤집었다", [
    (0, "BAAR: 유효 집합 멤버십을 증명 → 폐기는 리프 제거", True),
    (0, "본 구현: 폐기 집합 비멤버십을 증명 → 폐기는 리프 추가", True),
    (0, "왜 뒤집었나", True),
    (1, "유효 집합 방식은 발급할 때마다 root가 바뀐다"),
    (1, "→ 발급 1건이 전 사용자의 witness를 무효화한다"),
    (1, "발급은 잦고 폐기는 드물다 → 폐기 집합 쪽이 root 변경 빈도가 훨씬 낮다"),
    (1, "root 변경이 드물어야 지갑이 증명을 캐시할 수 있고, grace window가 의미를 갖는다"),
    (0, "BAAR가 5,000개 규모 프로토타입이라 이 문제가 드러나지 않았을 가능성", True),
    (0, "IMT를 쓴 이유: 비멤버십에 sparse Merkle tree를 쓰면 깊이가 필드 비트수(254) → Poseidon 254회", True),
    (1, "IMT는 정렬된 리프에서 부등식 두 개 → 깊이를 실제 원소 수에 맞출 수 있음"),
])

# ============================================================
# 8. Groth16이 강제한 것
# ============================================================
add_bullet_slide("Groth16이 강제한 것 — 경로 검증이 회로 안으로", [
    (0, "Schnorr는 Merkle witness를 트랜스크립트에 얹으면 된다 (검증자가 직접 해시 확인)", True),
    (0, "Groth16은 경로 검증이 회로 제약으로 들어가야 한다", True),
    (1, "검증자가 스마트컨트랙트라 증명 밖의 데이터를 신뢰할 수 없음"),
    (0, "그 대가", True),
    (1, "pi_pk_i: 4,820 → 19,049 제약 (약 4배), 증명 시간 실측 799.61 ms (±85.69 ms, n=10)"),
    (0, "부수적으로 얻은 것 — 폐기 대상이 회로 밖으로 나오지 않는다", True),
    (1, "r_token이나 auid가 public이 되면 폐기가 linkability oracle이 되어 conditional privacy가 무너짐"),
    (1, "spec §3.4가 \"r_token을 public으로 올려 컨트랙트가 리스트를 조회\"하는 안을 기각한 이유"),
    (1, "→ 이 제약이 설계의 불변식. 회로 4배 비용은 그 대가다"),
])

# ============================================================
# 9. root를 왜 온체인에
# ============================================================
add_bullet_slide("root를 왜 온체인에 올리나? IdP 홈페이지면 안 되나?", [
    (0, "먼저: 온체인에 올라가는 건 32바이트 root 하나뿐이다", True),
    (1, "폐기 리프 목록은 IdP가 HTTP로 서빙, 지갑이 로컬에서 트리 재구성 → 리스트는 이미 오프체인"),
    (0, "root가 온체인이어야 하는 이유 셋", True),
    (1, "① 결정적 이유 — 검증자가 스마트컨트랙트다. PPIDWallet.execute()는 HTTPS를 호출할 수 없다"),
    (2, "검증자가 오프체인 RP 서버였다면 홈페이지 게시로 충분 — W3C Bitstring Status List가 그 방식"),
    (1, "② 비-equivocation — 웹서버는 클라이언트마다 다른 응답을 줄 수 있다"),
    (2, "IdP가 A에겐 \"폐기됨\", B에겐 \"유효함\"을 보여줘도 아무도 모름. 체인은 단일 공개 로그"),
    (2, "PairCT는 IdP를 신뢰하되 검증 가능하게 두려는 설계 → 이 성질이 중요"),
    (1, "③ 공통 시계 — GRACE_BLOCKS가 block.number 기준으로 동작"),
    (2, "홈페이지 게시엔 모두가 동의하는 시계가 없어 root의 신선도를 강제할 수 없음"),
])

# ============================================================
# 10. 리스트는 오히려 홈페이지가 낫다
# ============================================================
add_bullet_slide("답 — 리스트는 오히려 홈페이지/CDN에 올리는 게 낫다", [
    (0, "현재: 지갑이 트랜잭션 직전에 IdP에 폐기 목록을 직접 조회", True),
    (0, "문제 — 내용이 아니라 접근 패턴이 샌다", True),
    (1, "IdP는 로그인 시점에 uid와 그 클라이언트의 네트워크 신원을 안다"),
    (1, "IdP 로그 + 체인 관찰 → \"t에 세션 X 조회, t+Δ에 지갑 W에서 tx\" 상관"),
    (1, "→ 트랜잭션 2~3건이면 W ↔ uid가 사실상 확정"),
    (0, "이는 spec §3.4가 r_token public을 기각한 바로 그 연결을 다른 경로로 재도입한 것", True),
    (1, "§6.1은 트리 \"내용물\"만 논증하고 \"접근 패턴\"을 빠뜨렸음"),
    (0, "완화(적용됨): 조회를 증명 캐시 판정 뒤로 이동 → 캐시 적중 시 IdP 요청 없음", True),
    (0, "근본 해결: 리스트를 CDN·IPFS에 미러링 → IdP가 조회 패턴을 아예 볼 수 없음 (spec §5.5)", True),
])

# ============================================================
# 11. 1억 전제
# ============================================================
add_bullet_slide("1억 사용자면 depth와 회로 크기는?", [
    (0, "먼저 전제를 바로잡아야 한다 — 트리는 사용자가 아니라 폐기 항목을 담는다", True),
    (1, "사용자 1억 명 ≠ 리프 1억 개"),
    (0, "트리에 들어가는 것", True),
    (1, "계정 폐기 리프 — 현재 구현에서는 영구적"),
    (1, "세션 폐기 리프 — 크레덴셜이 max_height(1시간)로 만료되면 의미 없어짐"),
    (0, "현재 depth 20 = 1,048,576 슬롯", True),
    (0, "다음: 깊이를 늘리면 회로가 얼마나 커지는지 — 추정이 아니라 실측", True),
])

# ============================================================
# 12. 측정 결과
# ============================================================
add_table_slide(
    "실측 — depth별 회로 크기",
    ["depth", "슬롯 수", "pi_pk_i 제약", "depth 20 대비", "증명 시간"],
    [
        ["20 (현재)", "1,048,576", "19,049", "—", "799.61 ms (실측)"],
        ["24", "16.8M", "20,993", "+10.2%", "~880 ms (추정)"],
        ["27", "134M", "22,451", "+17.9%", "~940 ms (추정)"],
        ["30", "1.07B", "23,909", "+25.5%", "~1.00 s (추정)"],
        ["32", "4.29B", "24,881", "+30.6%", "~1.04 s (추정)"],
    ],
    [Inches(1.6), Inches(2.2), Inches(2.4), Inches(2.4), Inches(3.7)],
    height=Inches(2.9), font_size=13,
    note=[
        (0, "정확히 선형:  pi_pk_i(d) = 19,049 + 486 × (d − 20)"),
        (0, "레벨당 243제약(Poseidon 240 + mux 2 + 비트 1), 비멤버십 컴포넌트가 2개라 486"),
        (0, "증명 시간은 제약 수에 대략 선형이라 가정한 외삽 — 실측은 depth 20뿐"),
    ],
)

# ============================================================
# 13. 측정 방법
# ============================================================
add_bullet_slide("측정 방법 (재현 가능)", [
    (0, "pi_pk_i.circom을 깊이만 바꿔 복제하고 circom으로 직접 컴파일", True),
    (1, "IMTNonMembership(20) → (d)"),
    (1, "sess/acct_pathElements[20], pathIndices[20] → [d]"),
    (1, "배선 루프 i < 20 → i < d  (셋 중 하나라도 빠지면 컴파일 에러)"),
    (0, "circom <file> --r1cs -l circomlib/circuits -l circuits -l circuits/lib", True),
    (0, "검증: depth 20 컴파일 결과가 19,049로 저장소의 커밋된 .r1cs 값과 일치", True),
    (1, "→ 측정 절차 자체가 옳다는 확인"),
    (0, "비멤버십 템플릿 단독 측정도 동일한 기울기: nonmembership(d) = 243 × d + 1,767", True),
    (0, "임시 회로는 스크래치패드에만 생성, 저장소의 build/ 산출물은 건드리지 않음", True),
])

# ============================================================
# 14. 깊이는 병목이 아니다
# ============================================================
add_bullet_slide("깊이는 병목이 아니다 — 진짜 병목은 따로 있다", [
    (0, "최악 가정(전원 폐기) 1억 리프 → depth 27, 회로 +17.9%", True),
    (1, "로그 스케일이라 사용자 128배에 회로 1.18배 — 값싸다"),
    (0, "① 트리에서 리프를 지우는 경로가 없다", True),
    (1, "lib/imt.js의 API는 insert / getRoot / getNonMembershipWitness 셋뿐"),
    (1, "구조적 제약은 아님 — values는 정렬 배열이고 build()가 매번 전체를 재구성하므로 삭제는 몇 줄"),
    (1, "어려운 건 정책: 리프가 해시라 어느 세션인지·언제 만료되는지 트리만 봐선 모른다"),
    (0, "② 지갑이 트랜잭션마다 전체 트리를 재구성한다", True),
    (1, "getRoot()/getNonMembershipWitness()가 호출마다 build() → tx당 3회 재구성"),
    (0, "③ 용량 가드가 없다 — 2^depth 초과 시 getRoot()가 부분 root를 조용히 반환 → 전원 증명 불가", True),
    (0, "정확성 자체는 안전하다 — 과잉 포함은 fail-closed. 비용·가용성 문제다", True),
])

# ============================================================
# 15. 드러난 근본 문제
# ============================================================
add_bullet_slide("확장성 분석이 드러낸 것 — 폐기의 의미를 정한 적이 없다", [
    (0, "발견된 우회 경로 (수정 완료, 커밋 3881ac5)", True),
    (1, "auid = Poseidon(uid, salt) → salt를 갈아끼우면 auid가 바뀌어 자기 폐기를 통과"),
    (1, "이를 막는 검사가 /sso_with_credentials에만 있었고 주 경로 /authorize/login에는 없었음"),
    (0, "그런데 그 수정이 정당한 복구 경로도 같이 막았다", True),
    (1, "\"공격자의 salt 회전\"과 \"키 분실 후 재등록\"이 IdP 눈에는 동일하다 — 같은 uid, 다른 auid"),
    (1, "→ 키 분실이 곧 계정 영구 상실"),
    (0, "삭제 API 부재 + 이 수정이 겹쳐 현재 의미론이 \"영구 추방\"으로 굳었다", True),
    (1, "그렇게 정한 적이 없다 — 설계 문서에 근거가 없고 우연에 가깝다"),
    (0, "두 해석이 갈린다", True),
    (1, "(a) 크레덴셜 폐기 — \"이 uid가 이 지갑에 묶인 바인딩을 무효화\" → 재등록 허용"),
    (1, "(b) 계정 차단 — \"이 사람은 시스템을 쓸 수 없다\" → 재등록 차단 (현재 상태)"),
])

# ============================================================
# 16. (a)를 택하는 이유
# ============================================================
add_bullet_slide("결론 — (a) 크레덴셜 폐기가 맞다", [
    (0, "① 트리가 유일하게 할 수 있는 일이 (a)다", True),
    (1, "미래 로그인을 막는 데는 트리가 필요 없다 — IdP가 토큰을 안 내주면 된다 (계정 disabled 플래그)"),
    (1, "IdP가 못 하는 것은 이미 발급돼 나간 크레덴셜의 회수 — 그게 트리의 존재 이유"),
    (1, "(b)를 트리로 하는 건 더 간단한 수단이 있는데 회로 19,049 제약을 동원하는 것"),
    (0, "② 온체인 검증자는 \"사람\"을 모른다", True),
    (1, "execute()가 하는 일은 \"이 크레덴셜의 핸들이 폐기 집합에 없는가\" — 무상태 크레덴셜 상태 조회"),
    (0, "③ 체인에서는 가역성이 더 중요하다", True),
    (1, "오폐기에 항소 수단이 없고, 되돌릴 수 없으면 그 지갑 자산이 영구 동결된다"),
    (0, "두 층으로 분리한다", True),
    (1, "발급된 크레덴셜 무효화 → 폐기 트리 (가역)   /   사람 차단 → IdP 계정 비활성화"),
    (0, "남는 성질: 새 salt는 새 PPID → 새 지갑. 옛 지갑 자산은 폐기가 풀려야 접근 가능 (논문에 명시)", True),
])

# ============================================================
# 17. (a)의 귀결
# ============================================================
add_bullet_slide("(a)의 귀결 — 트리 전체가 ephemeral해진다", [
    (0, "(a)의 정의를 그대로 따라가면", True),
    (1, "트리의 일 = 이미 발급돼 나간 것을 만료 전에 무효화"),
    (1, "max_height가 지나면 이미 발급된 것들은 전부 죽는다"),
    (1, "앞으로 발급하지 않는 것은 IdP 계정 층이 담당"),
    (1, "⟹ 계정 리프도 max_height 이후엔 뺄 수 있다"),
    (0, "계정 리프가 영구히 남아야 했던 건 (b) 해석 때문이었다", True),
    (0, "트리 크기가 \"역대 폐기 수\"에서 \"최근 한 창 안의 긴급 폐기 수\"로 바뀐다", True),
    (0, "그리고 이건 heartbeat와 같은 작업이다", True),
    (1, "GRACE_BLOCKS마다 root를 재게시해야 하는 건 어차피 필요(fail-closed 전환 때문)"),
    (1, "그 sweep에서 만료 리프를 같이 걷어내면 된다 → cron 하나가 둘 다 수행"),
    (0, "가역성이 공짜로 따라온다 — 폐기 해제를 위한 별도 기능 없이, 갱신하지 않으면 된다", True),
])

# ============================================================
# 18. 규칙
# ============================================================
add_table_slide(
    "규칙 — 같은 술어를 두 시점에 적용한다",
    ["", "만료 블록", "입구 검사 (insert)", "sweep (주기)"],
    [
        ["세션\n(r_token)", "그 크레덴셜의\nmax_height",
         "발급 기록을 조회 → 기록이 없거나\n이미 만료면 거부", "만료되면 제거"],
        ["계정\n(auid)", "폐기 시각 + max_height\n(폐기 직전 발급분까지 덮기 위해)",
         "항상 통과 (만료가 항상 미래)", "만료되면 제거\n= 폐기 해제"],
    ],
    [Inches(1.6), Inches(3.2), Inches(4.3), Inches(3.2)],
    height=Inches(2.2), font_size=12,
    note=[
        (0, "술어는 하나:  만료블록 > 현재블록.  입구에서 쓰면 태어날 때부터 죽은 것을 막고, sweep에서 쓰면 살다가 죽은 것을 치운다."),
        (0, "입구 검사의 공짜 이득: 만료를 판단하려면 \"이 r_token을 내가 발급했는가\"를 먼저 조회해야 한다"),
        (0, "→ 발급한 적 없는 값은 거부되므로, 임의의 숫자를 넣어 트리를 부풀리는 오염이 원천 차단된다."),
    ],
)

# ============================================================
# 19. 해소되는 것들
# ============================================================
add_table_slide(
    "한꺼번에 해소되는 것들",
    ["문제", "현재", "(a) + ephemeral 트리"],
    [
        ["트리 무한 증가", "영구 누적", "소멸 — 한 창 분량으로 유계"],
        ["tx당 O(n) 재구성", "리프 100만이면 치명적", "n이 작아져 무해"],
        ["용량 초과", "넘으면 조용히 전원 장애", "한계에서 한참 멀어짐"],
        ["폐기 시 unlinkability 상실", "노출 창 = 영구", "노출 창 = 한 max_height"],
        ["가역성", "없음 (오폐기 복구 불가)", "프루닝이 곧 폐기 해제 — 별도 기능 불필요"],
        ["필요 depth", "1억이면 27", "20도 과할 수 있음"],
    ],
    [Inches(3.2), Inches(4.0), Inches(5.1)],
    height=Inches(3.4), font_size=13,
    note=[
        (0, "대가: IdP 계정 층의 영속화가 필수가 된다(지속적 차단이 그쪽으로 옮겨가므로). root 변경이 잦아져 지갑 캐시 적중률이 떨어진다."),
    ],
)

# ============================================================
# 20. 정리
# ============================================================
add_bullet_slide("정리", [
    (0, "폐기 기준: 세션은 r_token, 계정은 auid — IdP가 아는 값이면서 체인에는 공개되지 않는 값", True),
    (0, "BAAR 대비 핵심 변경: 누적기 방향을 멤버십 → 비멤버십으로 뒤집음", True),
    (1, "발급 때마다 전 사용자 witness가 깨지는 문제를 회피. Groth16이라 경로 검증이 회로 안으로(4배 비용)"),
    (0, "root가 온체인인 이유는 검증자가 스마트컨트랙트이기 때문", True),
    (1, "리스트는 이미 오프체인이며, CDN 미러링이 프라이버시상 오히려 바람직"),
    (0, "1억 확장은 회로 깊이 문제가 아니다 (+18%) — 트리 수명 관리가 실제 제약이었다", True),
    (0, "그 분석이 설계 재검토로 이어졌다: 폐기 = 크레덴셜 폐기(a), 트리는 ephemeral", True),
    (1, "확장성·가역성·프라이버시 문제가 함께 해소되고, sweep이 heartbeat와 통합된다"),
    (0, "논문에 명시할 성질: 폐기 효력 지연 상한 = GRACE_BLOCKS / 새 salt는 새 지갑(옛 지갑 자산 동결)", True),
])

# ============================================================
# 21. 남은 과제
# ============================================================
add_bullet_slide("남은 과제 (재설계 반영)", [
    (0, "설계 전환에 필요한 작업", True),
    (1, "issuanceLog가 max_height도 저장 (지금은 r_token → uid만)"),
    (1, "/idp/revoke가 발급 기록을 조회해 (리프, 만료 블록) 기록 + 입구 검사"),
    (1, "sweep = heartbeat 통합: 만료 리프 제거 후 root 재게시"),
    (1, "lib/imt.js에 remove() / users[]에 disabled 플래그"),
    (1, "pinAuidToAccount는 유지하되 IdP 인증 복구 흐름에서 해제 가능하게"),
    (0, "재설계와 무관하게 남는 것", True),
    (1, "IdP 폐기 집합·계정 상태 영속화 — 재시작 시 소실되면 차단이 통째로 풀림 (중요도 상승)"),
    (1, "onlyIdP 주소와 IdP EdDSA 키가 암호학적으로 연결돼 있지 않음"),
    (1, "회로 건전성 테스트가 clean checkout에서 실행 불가 (/tmp 경로 하드코딩)"),
    (1, ".env·CLAUDE.md 갱신, 벤치마크 재측정"),
])

prs.save(OUT)
print(f"wrote {OUT} ({len(prs.slides._sldIdLst)} slides)")
