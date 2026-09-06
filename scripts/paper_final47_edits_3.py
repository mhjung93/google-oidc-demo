from docx import Document
from docx.shared import Pt
F = 'documents/PairCT_research_article_20260706_092133_with_figures_final47.docx'
d = Document(F)

# ── Table 1: Revocation 열 추가 + BAAR 행 추가 ────────────────────────────
t = d.tables[0]
ncols_before = len(t.columns)
col = t.add_column(Pt(90))
cells = [r.cells[-1] for r in t.rows]
values = [
  'Credential Revocation',
  'Not targeted',
  'Not targeted (ephemeral key expiry only)',
  'Not targeted',
  'Not targeted',
  'Partial (committee-mediated key rotation)',
  'Stated as future work',
  'Not targeted',
  'Yes — non-membership proven in-circuit; only a root on-chain / publisher is a standing role',
]
assert len(values) == len(t.rows), f'{len(values)} != {len(t.rows)}'
for c, v in zip(cells, values):
    c.text = v
print(f'열 추가: {ncols_before} -> {len(t.columns)}')

# BAAR 행
row = t.add_row()
baar = [
  'BAAR [20]',
  'Not targeted (single issuing authority; no IdP/RP split)',
  'Claimed; the presentation sends the credential commitment and its witness in the clear, so a verifier sees a stable per-credential identifier',
  'No (Schnorr credential, no wallet-bound session key)',
  'Not targeted / issuing authority is a standing role',
  'Yes — membership in a Merkle accumulator of valid credentials',
]
assert len(baar) == len(row.cells), f'{len(baar)} != {len(row.cells)}'
for c, v in zip(row.cells, baar):
    c.text = v
print('BAAR 행 추가')

# ── Threats to Validity: 폐기 관련 항목 추가 ─────────────────────────────
P = d.paragraphs
idx = None
for i,p in enumerate(P):
    if p.text.startswith('The implementation results are subject to four threats to validity'):
        idx = i; break
assert idx is not None, 'Threats 문단 못 찾음'
p = P[idx]
p.runs[0].text = p.text.replace('four threats to validity', 'six threats to validity', 1)
for r in p.runs[1:]: r.text = ''
P[idx+1].insert_paragraph_before(
  'Fifth, the revocation results assume a revocation rate far below the login rate. The structure of Section G is '
  'sized for that regime: if every logout were revoked, the live revocation set would grow at the login rate and the '
  'non-membership formulation would lose its advantage over a validity-set formulation. Treating expiry, not '
  'revocation, as the normal way a session ends is therefore a deployment requirement rather than an optimization. '
  'Sixth, the on-chain measurements were taken against a local development chain; per-block publication is only '
  'economical on a rollup or an application-specific chain, and the gas figures should be read as relative to the '
  'non-revoking baseline (330,079 versus 276,709 for a first execution) rather than as absolute deployment costs.',
  style='PARA')
print('OK Threats 확장')

# ── References: BAAR 추가 ────────────────────────────────────────────────
P = d.paragraphs
last_ref = None
for i,p in enumerate(P):
    if p.style.name == 'Ref':
        last_ref = i
assert last_ref is not None
P[last_ref].insert_paragraph_before(
  '[20] M. Ahmed, A. Ahmad, F. Zeshan, and S. Akram, "BAAR: A framework for blockchain-based anonymous and '
  'revocable user authentication scheme," PLOS ONE, vol. 21, no. 3, e0343696, Mar. 2026, '
  'doi: 10.1371/journal.pone.0343696.',
  style='Ref')
print('OK 참고문헌 [20] 추가')

d.save(F)
print('저장:', F)
