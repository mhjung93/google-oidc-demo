from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documents" / "PairCT_address_abstraction_draft03.docx"

TITLE = "PairCT: Privacy-Preserving Address Abstraction With Identity-Level Traceability"
AUTHORS = "MOONHYEON CHUNG AND CHANIK PARK"
AFFILIATION = "Department of Computer Science and Engineering, Pohang University of Science and Technology, Pohang 37673, South Korea"

ABSTRACT = (
    "Blockchain users are commonly identified through chain-specific addresses, resulting in fragmented accounts "
    "across blockchains and services. Reusing an address improves continuity but enables cross-context linkability, "
    "whereas creating independent addresses increases management and recovery burdens. Existing account-abstraction "
    "systems improve transaction programmability, and credential-based address-abstraction schemes reduce address "
    "fragmentation, but they do not jointly provide identity-provider-assisted authentication, identity-level "
    "accountability, and revocable session-based execution. We present PairCT, a privacy-preserving address-abstraction "
    "framework that derives RP- and chain-scoped smart accounts from a common identity-provider-authenticated identity. "
    "For a relying-party identifier and a namespace-qualified chain context, PairCT derives a pairwise pseudonymous "
    "identifier without revealing the user's identity-provider-side identifier. A zero-knowledge proof binds the derived "
    "identifier, target context, identity-provider-signed authentication statement, and wallet-generated session key. "
    "This construction enables one IdP-authenticated identity to authorize accounts across supported relying parties and "
    "blockchains while avoiding a reusable global address. PairCT further provides identity-level traceability: when "
    "opening is authorized for a disputed session, retained authentication records can associate the corresponding "
    "account with its identity-provider-scoped identity (iss, uid), while ordinary activity remains pseudonymous. "
    "Finally, session-bound authorization amortizes authentication and proof-generation costs across multiple transactions "
    "and supports selective session or account revocation through a Merkle-accumulator state anchored on-chain."
)

INDEX_TERMS = (
    "Address abstraction, identity-provider-assisted authentication, identity-level traceability, pairwise pseudonymous "
    "identifier, session revocation, zero-knowledge proofs."
)

SECTIONS = [
    ("I. INTRODUCTION", [
        "Web3 users interact with services through blockchain accounts whose addresses and authorization rules are tied to particular chains. A user accessing several services on several chains may accumulate many addresses, signing keys, recovery procedures, and service-side account records. Reusing one address reduces this burden but exposes a stable public identifier that permits activity correlation. Generating a fresh address for every context improves separation but leaves users and services to reconstruct continuity on top of unrelated identifiers.",
        "Account abstraction moves authorization logic from a fixed externally owned account into a programmable smart account. This enables alternative verification rules, recovery policies, sponsored execution, and batched calls. It does not by itself determine how one authenticated person obtains accounts across relying parties (RPs) and chains, or how those accounts remain unlinkable outside their intended contexts. Chain abstraction simplifies interaction across ledgers but commonly retains chain-specific identifiers and cryptography beneath the interface.",
        "Credential-based address abstraction instead treats a certificate or credential as the root of blockchain-facing identifiers. zkAA demonstrates this approach using Web2 certificates and zero-knowledge proofs [17]. Its issuer participates at credential issuance, after which the user presents the credential in a self-sovereign manner. PairCT studies a different lifecycle: an identity provider (IdP) authenticates the user for each login and issues a fresh signed statement for a bounded session. This model can reuse established authentication and recovery infrastructure, but it introduces IdP tracking, session management, and accountability questions.",
        "An IdP-assisted construction must reconcile three requirements. First, one authenticated identity should authorize context-specific accounts without a new long-term blockchain identity for every RP and chain. Second, this convenience must not produce a globally reusable identifier. Third, pseudonymity must not eliminate accountability: an authorized investigation of a disputed session should identify the responsible IdP-scoped account rather than return only another opaque address.",
        "Efficiency adds a fourth concern. Repeating IdP authentication and zero-knowledge proof generation for every transaction would place interactive authentication on the transaction path. PairCT instead binds one successful login to a wallet-generated session key. Multiple transactions can then reuse the resulting authorization within a bounded lifetime, while session- and account-level revocation provide early invalidation.",
        "PairCT derives a context-scoped identifier conceptually as PPID_(r,c) = H(domain, iss, uid, rid, chainContext_c, salt), where iss and uid identify the IdP account, rid identifies the RP, and chainContext is a namespace-qualified ledger context. The wallet-held salt prevents public recomputation. Each supported chain maps this identifier to a chain-local programmable account; the construction does not require one identical address format across heterogeneous ledgers.",
        "The contributions are threefold. First, PairCT supports identity-level traceability by joining a retained disputed-session transcript with the corresponding IdP issuance record under external authorization. Second, it provides privacy-preserving address abstraction from a common IdP-authenticated identity to deterministic RP- and chain-scoped accounts. Third, it amortizes authentication and proof costs through session-bound execution and supports selective, efficiently maintainable session and account revocation."
    ]),
    ("II. BACKGROUND AND RELATED FOUNDATIONS", [
        "An externally owned account authorizes transactions through possession of a long-term private key. A smart account can replace this fixed rule with programmable validation. ERC-4337 provides one prominent account-abstraction architecture, but PairCT's core construction is independent of a particular transaction-entry framework: its essential requirement is a programmable account capable of checking a session signature, a validity bound, and a zero-knowledge proof.",
        "OIDC allows an RP to authenticate a user through an IdP-issued ID token [4], [5]. Conventional flows expose a client identifier, redirect URI, or audience to the IdP. PairCT therefore uses an OIDC-inspired authorization-code flow with an IdP extension: a common wallet client, pushed authorization data, PKCE, and a loopback callback separate the browser-visible authorization reference from the hidden RP relation. PairCT is not claimed to be an unmodified OIDC profile.",
        "Groth16 provides succinct noninteractive arguments suitable for compact verification [13]. PairCT uses zero knowledge to show that a context-scoped identifier, a signed session statement, and revocation witnesses share consistent hidden values. The construction relies on a correctly generated proving system and on domain-separated hashes for PPID derivation, session binding, and revocation leaves.",
        "PairCT's revocation structure is an indexed Merkle accumulator. A nonmembership witness identifies adjacent ordered leaves low and next such that low < target < next and proves their authenticated path to a published root. Separate domain tags distinguish session and account leaves even when their underlying field values coincide."
    ]),
    ("III. SYSTEM AND SECURITY MODEL", [
        "The system contains a user, wallet, IdP, RP frontend and backend, programmable blockchain, smart-account factory, revocation registry, and an external opening-authorization process. The wallet protects uid-related wallet state, a high-entropy salt, and ephemeral session secret keys. The IdP authenticates uid and signs session statements. The RP maintains its local account under PPID and retains accepted transcripts when accountability is required.",
        "The IdP is honest in authentication and signing but may be curious about the accessed RP. RPs may compare their own pseudonymous views, but RP-IdP collusion outside an authorized opening is excluded. The wallet is trusted to validate authenticated RP metadata against the user-visible origin and to protect its secrets. Network-level anonymity is outside the model: an IdP may observe an IP address and timing, so RP hiding concerns protocol identifiers and transcript contents rather than traffic-analysis resistance.",
        "The adversary may control an RP frontend, replay messages, substitute contexts, observe public blockchain data, and submit transactions. Security requires authentication soundness, RP-side identity privacy, IdP-side RP hiding, cross-context unlinkability, session binding, replay resistance, revocation soundness, and correct transcript-scoped opening. Denial of service, compromised wallets, compromised setup parameters, and opening after required records have expired are outside the claimed guarantees.",
        "PairCT distinguishes protocol claims from deployment policy. Cryptography can bind a disputed account to retained records, but it cannot by itself decide whether an opening request is legally or organizationally authorized. A deployment must authenticate opening requests, enforce retention and access-control rules, and audit the result."
    ]),
    ("IV. PAIRCT CONSTRUCTION", [
        "Setup assigns each RP an identifier rid and authenticated origin metadata. A chain context is encoded as chainContext = Encode(namespace, reference), avoiding ambiguous numeric chain identifiers across ledger families. The IdP has a statement-signing key, and each supported chain has a verifier, smart-account factory, and revocation-root registry configured with the trusted IdP key.",
        "The wallet maintains a stable high-entropy salt for an enrolled IdP account. It computes PPID_(r,c) from domain, issuer, uid, rid, chainContext, and salt. Stability within the same context provides account continuity; rid and chainContext separation produce different values across contexts. Including iss prevents equal local uid values at different IdPs from colliding.",
        "At login the RP issues a fresh challenge rho. The wallet generates an ephemeral session key pair (sk_i, pk_i), selects an expiration bound h_i, and derives a token binding r_i = H_session(pk_i, h_i, rho). It constructs an IdP-facing zero-knowledge proof showing that the hidden rid and origin belong to a valid RP credential and that the blinded session values are consistent with the authenticated uid and wallet salt.",
        "The IdP authenticates the user independently, inserts its authenticated uid as the required public identity input, verifies the proof, and signs a statement binding issuer, common wallet audience, authorization nonce, blinded RP/account values, r_i, expiration, and chain context. The IdP learns uid and the signed blinded values but not rid, PPID, or the resulting smart-account address.",
        "The wallet receives the statement through an authorization-code exchange protected by state and PKCE. The RP receives PPID, pk_i, the signed statement, and an RP-facing proof. It checks its session challenge, context, expiration, IdP signature, and the proof that PPID was derived for its rid. Successful verification creates or resumes the service-local account represented by PPID."
    ]),
    ("V. ADDRESS AND ACCOUNT ABSTRACTION", [
        "PairCT abstracts identity from a raw blockchain address in two stages. The identity layer derives PPID_(r,c) from one IdP-authenticated identity while preserving context separation. The account layer maps PPID_(r,c) to a chain-local programmable account. On an EVM chain, a factory can use CREATE2 with PPID_(r,c) as a salt, producing a counterfactual address before deployment and deploying the account lazily on first use.",
        "The term address abstraction does not mean that one byte-identical address must exist on every ledger. Factory address, bytecode, constructor parameters, and address rules differ by chain. PairCT instead provides one authentication foundation and a deterministic derivation rule for each qualified context. This formulation avoids claiming portability where a ledger lacks compatible programmable-account support.",
        "The smart account authorizes a payload only if its stored nonce matches, a signature recovers pk_i, the IdP key is trusted, the current height does not exceed h_i, the supplied revocation root is recent, and the session proof verifies. The proof connects pk_i and PPID to the same IdP-signed session while hiding uid, rid, salt, and revocation targets from the chain.",
        "This design removes a reusable EOA key from the logical authorization condition, although an EOA, bundler, or relayer may still transport the deployment and execution transactions and pay gas. The transport payer does not gain authority over the PPID account unless it also possesses the valid session authorization."
    ]),
    ("VI. SESSION-BOUND EXECUTION", [
        "A PairCT session begins with one IdP authentication and one session proof. The wallet signs each call payload with sk_i, while the account nonce prevents replay. The same session proof can be cached and reused while the session key and accepted revocation root remain unchanged. A root change or a new login invalidates the cache key and requires a fresh proof.",
        "The validity bound limits exposure if a session key is lost. Binding r_i to pk_i, h_i, and the RP challenge prevents a statement issued for one session from authorizing another key or validity interval. Chain context must also be included in the signed relation and, in a complete chain-scoped implementation, checked against the executing chain's canonical context.",
        "Amortization is the principal performance benefit: interactive IdP authentication and proof generation occur once per session rather than once per transaction. The achievable gain depends on the number of calls per session, root-update frequency, and whether proof verification occurs directly in a smart account or through an account-abstraction entry framework."
    ]),
    ("VII. REVOCATION", [
        "PairCT supports session revocation and account revocation. A session leaf is H_rev(TAG_SESSION, r_i). An account leaf is H_rev(TAG_ACCOUNT, auid), where auid = H_account(iss, uid, salt) is stable for an enrolled wallet-IdP binding and independent of rid. The wallet proves that both targets are absent from the accumulator represented by the root supplied to the smart account.",
        "Revocation requests first enter a pending set. A prepare operation computes the next leaf set by adding pending entries and pruning expired entries, then returns an expected root without changing the published IdP view. After that root is anchored on-chain, commit advances the IdP's published state. This ordering avoids distributing witnesses for a root that the chain does not yet recognize.",
        "A bounded registry retains several recently published roots to tolerate in-flight proofs. A block-based freshness window prevents an old root from remaining valid indefinitely, while heartbeat publication refreshes an unchanged root when no revocations occur. The resulting delay is explicit: a revoked credential may remain usable until the relevant old root leaves the accepted window. Deployment parameters must be chosen from confirmation latency and incident-response requirements rather than treated as universal constants.",
        "Session leaves expire at their session maximum height. Account-revocation leaves need remain only long enough to cover credentials issued before revocation; afterward they can be removed unless policy requires permanent account disablement. Consequently the active accumulator scales with recent exceptional revocations rather than all users or all historical sessions."
    ]),
    ("VIII. IDENTITY-LEVEL TRACEABILITY", [
        "Ordinary execution exposes a context-scoped account and session public key but not uid. For a disputed transaction, the RP identifies the corresponding accepted account/session transcript. The RP record supplies a session join value such as auid_i, while the IdP issuance log maps that value to the authenticated identity (iss, uid). Joining the integrity-protected records yields identity-level rather than address-level traceability.",
        "Opening is scoped to an already identified transcript. It is not a query that starts from uid and enumerates all RPs or accounts. IdP-side RP hiding and RP-side uid privacy remain meaningful outside opening because neither party alone holds both sides of the join under the non-collusion model.",
        "Authorization is an external precondition, not an automatic consequence of possessing a join value. A production deployment should require an authenticated authorization artifact identifying the disputed transcript, requester, purpose, and validity interval; restrict the returned result to the authorized recipient; and append an auditable opening record. The current prototype validates record joining but does not cryptographically enforce this policy boundary."
    ]),
    ("IX. SECURITY ANALYSIS", [
        "Authentication soundness follows from the IdP signature, Groth16 knowledge soundness, session-key signature, and canonical context checks. An accepted account execution implies a valid statement for the hidden account relation, possession of sk_i for the payload, a nonexpired height, a fresh account nonce, and nonmembership under an accepted root, except with the failure probability of the underlying primitives.",
        "RP-side identity privacy follows from zero knowledge and the entropy of salt: the RP sees stable PPID_(r,c) by design but cannot recover uid from public information alone. IdP-side RP hiding follows because rid, origin, PPID, and account address remain hidden witnesses; the IdP sees only a common client and blinded values. This claim excludes network metadata and RP-IdP collusion.",
        "Cross-context unlinkability follows from including rid and chainContext in the PPID derivation. Colluding RPs or public observers see independently scoped outputs unless they learn the wallet salt, obtain uid and all context inputs, break the hash assumption, or combine views with the IdP outside the model. Same-context continuity is intentionally visible.",
        "Session binding follows from the signed relation over pk_i, expiration, challenge, and chain context. Payload nonces prevent transaction replay; authorization-code state and PKCE protect the browser/loopback exchange. Revocation soundness additionally requires the contract to accept only recent registry roots and the proof to establish nonmembership for correctly domain-separated targets.",
        "Opening correctness requires integrity-protected RP and IdP records and a collision-resistant join value. Exclusivity is conditional on the authorization system and non-collusion assumption: the cryptographic construction establishes joinability, while deployment controls determine who may invoke it and receive uid."
    ]),
    ("X. PROTOTYPE", [
        "The prototype separates an RP frontend/backend, a local wallet agent, a MetaMask Snap approval interface, a custom IdP, Groth16 circuits, a deterministic PPIDWalletFactory, PPIDWallet smart accounts, and a RevocationRegistry. The local agent stores the wallet salt, generates fresh secp256k1 session keys, produces proofs, performs PAR and token exchange, and constructs execution calldata. The RP frontend never receives uid, salt, or sk_i.",
        "The authorization path uses a common wallet client, pushed authorization request, authorization code, PKCE, state, and a dynamically selected 127.0.0.1 loopback callback. The RP frontend opens the IdP authorization page only after receiving the opaque request reference. The IdP authenticates uid and verifies the hidden-RP proof before issuing the signed statement.",
        "The current implementation uses PPID = Poseidon(uid, rid, salt); chainContext has not yet been added to that circuit. It binds chain_id inside the IdP-signed token, but the current on-chain public inputs do not compare it directly with block.chainid. Accordingly, RP- and chain-scoped derivation is the proposed construction, whereas the prototype currently validates RP-scoped derivation and chain-bound statement plumbing.",
        "The deployed PPIDWallet is a minimal programmable smart account with one arbitrary external call, not a complete ERC-4337 EntryPoint/UserOperation implementation. A browser EOA currently sends deployment and execution transactions and pays gas. The trace endpoint demonstrates RP/IdP record joining but lacks production authorization, retention, and audit enforcement. These boundaries prevent prototype behavior from being overstated as a complete deployment."
    ]),
    ("XI. PRELIMINARY EVALUATION AND DISCUSSION", [
        "Existing warm-state measurements cover three proof-generation paths. The previously recorded means over ten warmed runs were 201.30 ms for pi_arid_i, 85.00 ms for pi_ppid, and 239.20 ms for pi_pk_i, with a sequential per-run sum of 525.50 ms. These values are prototype circuit-path measurements, not end-to-end login latency, and must be remeasured after the chain-context and revocation relations are finalized.",
        "A Groth16 proof occupies 256 bytes in the prototype's ABI-oriented representation; total calldata additionally depends on the number of public inputs. JSON proof sizes around 723 bytes are serialization artifacts rather than raw proof sizes. Revocation adds a compact public root on-chain while keeping the two nonmembership targets and their paths private in the proof witness.",
        "PairCT trades compatibility for privacy and accountability. An unmodified IdP-visible OIDC audience cannot simultaneously hide the actual RP, so the construction requires an IdP extension and a common wallet client. Stable wallet salt also creates enrollment and recovery obligations. Changing salt must be treated as an explicit re-binding or account-reset event; otherwise account continuity and account-level revocation can be bypassed.",
        "The revocation design trades immediate invalidation for availability of in-flight sessions. Root buffers, freshness windows, and publication cadence are separate parameters: the buffer limits how many publication states are represented, while the freshness window limits how long any matching state remains acceptable. Production evaluation should measure publication delay, chain finality, witness distribution, root-cache hit rate, and accumulator size under realistic revocation workloads."
    ]),
    ("XII. RELATED WORK", [
        "SIWE standardizes nonce-bearing wallet-address login [1] but does not abstract identity from a reused address. zkLogin combines OIDC credentials and zero-knowledge proofs for blockchain authorization [3]; its audience remains visible to the IdP and its primary goal differs from RP-hiding, transcript-scoped opening, and selective session/account revocation.",
        "zkAA is the closest address-abstraction comparison [17]. It targets identifier fragmentation through an issuer-once, certificate-based design and presents a unified credential to supported chains. PairCT retains online IdP authentication per login, derives pairwise RP- and chain-scoped accounts rather than a reusable global identifier, binds a bounded session key, and provides an RP/IdP record-joining path to identity-level traceability.",
        "Privacy-preserving SSO systems such as SPRESSO, EL PASSO, UPPRESSO, and MISO reduce IdP tracking or inter-RP linkability [6]-[9]. EL PASSO supports accountability through ciphertexts and threshold decryption authorities, while PairCT uses existing RP and IdP records under an external authorization boundary. ARPSSO hides the RP in OIDC-style login but leaves tracking and de-anonymization as future work [18]. Han et al. provide traceable anonymous SSO through a dedicated central verifier [19].",
        "Anonymous credentials provide selective disclosure, unlinkability, and optional anonymity revocation [14]-[16]. PairCT addresses a narrower Web3 execution lifecycle: online IdP authentication is bound to a context-scoped programmable account and short-lived session key, with exceptional revocation represented by a chain-anchored accumulator."
    ]),
    ("XIII. LIMITATIONS", [
        "PairCT assumes a trusted wallet, high-entropy stable salt, authenticated RP metadata, non-collusion outside authorized opening, integrity-protected records, and secure Groth16 setup. It does not provide network anonymity, global Sybil resistance across IdP accounts, or opening after records have expired.",
        "The proposed chain-scoped relation and namespace encoding require circuit and contract completion. Supporting heterogeneous ledgers requires chain-specific account adapters and does not imply identical addresses. The prototype's current custom smart account does not establish full ERC-4337 compatibility, paymaster behavior, batched execution, or production recovery.",
        "Conditional traceability has an intentional privacy cost. Once an authorized RP receives uid, it can associate uid with history already stored under that RP's stable PPID. PairCT limits cross-context discovery; it does not erase the requesting RP's own local history."
    ]),
    ("XIV. CONCLUSION", [
        "This paper presented PairCT as a privacy-preserving address-abstraction framework for Web3 authentication. It derives deterministic RP- and chain-scoped programmable accounts from a common IdP-authenticated identity while avoiding a globally reusable address.",
        "PairCT combines this abstraction with identity-level traceability and session-bound authorization. Ordinary execution remains pseudonymous, while an externally authorized opening of a disputed session joins retained RP and IdP records to identify (iss, uid). A successful login is bound to a wallet-generated session key, amortizing authentication and proof costs over multiple transactions. A Merkle-based mechanism supports selective early invalidation and pruning of expired revocations.",
        "The current prototype validates RP-scoped derivation, hidden-RP authentication flow, session-bound smart-account execution, record joining, and batched revocation. Completing namespace-qualified chain derivation, on-chain chain-context enforcement, production opening authorization, and broader evaluation remains future work."
    ]),
]

REFERENCES = [
    '[1] W. Chang et al., "ERC-4361: Sign-In with Ethereum," Ethereum Improvement Proposals, no. 4361, 2021.',
    '[2] K. Yan, X. Zhang, and W. Diao, "Stealing Trust: Unraveling Blind Message Attacks in Web3 Authentication," Proc. ACM CCS, 2024, pp. 555-569.',
    '[3] F. Baldimtsi et al., "zkLogin: Privacy-preserving blockchain authentication with existing credentials," Proc. ACM CCS, 2024, pp. 3182-3196.',
    '[4] N. Sakimura et al., "OpenID Connect Core 1.0 incorporating errata set 1," OpenID Foundation, 2014.',
    '[5] D. Hardt, "The OAuth 2.0 Authorization Framework," RFC 6749, 2012.',
    '[6] D. Fett, R. Kusters, and G. Schmitz, "SPRESSO: A secure, privacy-respecting single sign-on system for the Web," Proc. ACM CCS, 2015.',
    '[7] Z. Zhang et al., "EL PASSO: Efficient and lightweight privacy-preserving single sign on," Proc. PETS, vol. 2021, no. 2, pp. 70-87.',
    '[8] C. Guo et al., "UPPRESSO: Untraceable and unlinkable privacy-preserving single sign-on services," CoRR, abs/2110.10396, 2021.',
    '[9] R. Xu et al., "MISO: Legacy-compatible privacy-preserving single sign-on using trusted execution environments," Proc. IEEE EuroS&P, 2023.',
    '[10] D. Fett, R. Kusters, and G. Schmitz, "The Web SSO standard OpenID Connect: In-depth formal security analysis and security guidelines," Proc. IEEE CSF, 2017.',
    '[11] D. Fett, R. Kusters, and G. Schmitz, "A comprehensive formal security analysis of OAuth 2.0," Proc. ACM CCS, 2016.',
    '[12] D. Maram et al., "CanDID: Can-do decentralized identity with legacy compatibility, Sybil-resistance, and accountability," Proc. IEEE S&P, 2021.',
    '[13] J. Groth, "On the Size of Pairing-Based Non-interactive Arguments," EUROCRYPT, 2016, pp. 305-326.',
    '[14] D. Chaum, "Security without identification," Commun. ACM, vol. 28, no. 10, 1985.',
    '[15] J. Camenisch and A. Lysyanskaya, "An efficient system for non-transferable anonymous credentials with optional anonymity revocation," EUROCRYPT, 2001.',
    '[16] D. Chaum, "Blind signatures for untraceable payments," Advances in Cryptology, 1983.',
    '[17] S. Park et al., "Beyond the blockchain address: Zero-knowledge address abstraction," Proc. ACM SAC, 2025, pp. 366-374.',
    '[18] J. He et al., "ARPSSO: An OIDC-compatible privacy-preserving SSO scheme based on RP anonymization," ESORICS, 2024.',
    '[19] J. Han et al., "Anonymous single-sign-on for n designated services with traceability," ESORICS, 2018.',
]


def font(run, size=9.5, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def columns(section, count):
    cols = section._sectPr.xpath("./w:cols")
    element = cols[0] if cols else OxmlElement("w:cols")
    if not cols:
        section._sectPr.append(element)
    element.set(qn("w:num"), str(count))
    if count > 1:
        element.set(qn("w:space"), "360")


def configure(section, count):
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    columns(section, count)


def body(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.16)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    font(p.add_run(text))


def heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), 10)


def build():
    doc = Document()
    configure(doc.sections[0], 1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(9.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(TITLE), 20)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(AUTHORS), 10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(AFFILIATION), 8.5, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.right_indent = Inches(0.45)
    font(p.add_run("ABSTRACT "), 9, bold=True)
    font(p.add_run(ABSTRACT), 9)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.right_indent = Inches(0.45)
    font(p.add_run("INDEX TERMS "), 9, bold=True)
    font(p.add_run(INDEX_TERMS), 9, italic=True)

    section = doc.add_section(WD_SECTION.CONTINUOUS)
    configure(section, 2)
    for title, paragraphs in SECTIONS:
        heading(doc, title)
        for paragraph in paragraphs:
            body(doc, paragraph)

    heading(doc, "REFERENCES")
    for reference in REFERENCES:
        body(doc, reference, indent=False)

    doc.core_properties.title = TITLE
    doc.core_properties.subject = "PairCT address-abstraction full paper draft"
    doc.core_properties.keywords = INDEX_TERMS
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
