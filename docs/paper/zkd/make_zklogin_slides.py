# zk-Delegation vs zkLogin 비교 슬라이드 — 교수님 덱(260917_OVERALL.pptx) 템플릿 위에 새로 만든다.
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC = 'documents/260917_OVERALL.pptx'
DST = 'documents/260918_zkDelegation_vs_zkLogin.pptx'
prs = Presentation(SRC)
# 기존 슬라이드 제거(템플릿만 남긴다)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)

L_TITLE, L_BODY = prs.slide_layouts[0], prs.slide_layouts[1]

def title_slide(t, sub):
    s = prs.slides.add_slide(L_TITLE)
    s.shapes.title.text = t
    s.placeholders[1].text = sub
    return s

def bullets(title, items, size=20):
    s = prs.slides.add_slide(L_BODY)
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.clear()
    first = True
    for lvl, txt in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = txt; p.level = lvl
        for r in p.runs: r.font.size = Pt(size - 3 * lvl)
    return s

def table_slide(title, header, rows, widths, size=13):
    s = prs.slides.add_slide(L_BODY)
    s.shapes.title.text = title
    ph = s.placeholders[1]
    left, top, width, height = ph.left, ph.top, ph.width, ph.height
    ph._element.getparent().remove(ph._element)
    tbl = s.shapes.add_table(len(rows) + 1, len(header), left, top, width, Emu(int(height * 0.95))).table
    total = sum(widths)
    for j, w in enumerate(widths): tbl.columns[j].width = Emu(int(width * w / total))
    def put(cell, text, bold=False, sz=size):
        cell.text = ''
        tf = cell.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.text = text
        for r in p.runs: r.font.size = Pt(sz); r.font.bold = bold
        cell.margin_left = cell.margin_right = Emu(50000); cell.margin_top = cell.margin_bottom = Emu(30000)
    for j, h in enumerate(header): put(tbl.cell(0, j), h, bold=True)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row): put(tbl.cell(i, j), v, bold=(j == 0))
    return s

title_slide('zk-Delegation vs zkLogin', '무엇이 다르고, 왜 그렇게 선택했나\n2026-09-18')

bullets('한 줄 요약', [
    (0, 'zkLogin: 기존 OIDC 제공자(Google 등)를 건드리지 않고 Web2 로그인을 한 체인의 계정에 붙인다'),
    (0, 'zk-Delegation: 협력하는 발급자(AA)를 두는 대신, 제공자·서비스·체인 어느 쪽도 남의 몫을 보지 못하게 하고 폐기와 조건부 추적까지 갖춘다'),
    (0, '차이의 뿌리는 하나 — 발급자가 협력하느냐'),
    (1, 'zkLogin은 제공자에게 아무것도 요구할 수 없다 → JWT를 그대로 회로에서 검증, 주소·만료를 지갑이 정함'),
    (1, '우리는 AA가 협력한다 → "무엇에 서명하는지 모르는" 커밋 서명, 폐기 트리, 트레이스 태그가 전부 여기서 나온다'),
])

H = ['항목', 'zkLogin', 'zk-Delegation', '왜 이렇게 골랐나']
W = [1.1, 2.2, 2.4, 3.0]
table_slide('비교 1 — 발급자·가시성·회로·주소', H, [
    ['발급자', '무수정 OIDC 제공자. JWT를 그대로 씀', '협력하는 AA. 사용자 커밋 C에 EdDSA 서명', '제공자가 협력해야만 "무엇에 서명하는지 모르게" 만들 수 있다. G4·G5·폐기·개봉이 전부 이 전제 위'],
    ['제공자가 보는 것', '클라이언트(aud)·사용자 sub', 'uid·C·chainid·allowAgent. 서비스·가명·속성은 못 봄', 'OIDC의 근본 누수(어느 서비스에 로그인하는가)를 없애는 것이 목표'],
    ['회로 안 서명 검증', '제공자의 RSA(JWT). 제약 수백만, 외부 증명 서비스가 보통 필요', 'Baby Jubjub EdDSA-Poseidon. 제약 25,560, 지갑에서 1초 미만', '지갑이 스스로 증명하려면 회로 친화 서명이어야 한다. 협력 발급자라 가능'],
    ['주소', 'H(iss, aud, sub, salt). 앱(aud)당 하나, 같은 앱 안 서비스 간 공유', 'Poseidon(uid, s_u, chainid, arid). 서비스·체인마다 다르고 연결 불가', 'G3: 서비스끼리 기록을 합쳐도 연결 안 되게. 체인은 이름의 파라미터 → 신원은 체인 무관'],
], W)

table_slide('비교 2 — 비밀·유효기간·폐기·세션', H, [
    ['salt / 사용자 비밀', 'salt 서비스가 보관·백업', 's_u는 지갑에, AA에는 커밋 cm_u만. 발급마다 같은 값임을 증명', '비밀을 제3자에 맡기지 않으면서 불변성(G2) 강제. 대가: s_u 분실 = 주소 분실(복구 미해결)'],
    ['유효기간', '지갑이 고른 max_epoch, JWT nonce에 묶임. 제공자는 모름', '지갑이 고른 양자화 max_height를 AA가 서명. 검증자가 상한 L 강제', '같은 구조(사용자 선택·발급자 서명·검증자 상한). 차이: AA가 값을 보므로 지갑이 양자화. 둘 다 체인 시계(벽시계 아님)'],
    ['폐기', '없음. 만료로만 끝남', '체인에 게시되는 append-only 트리, 성명·계정 단위, 자기 폐기', '만료 전에 끊을 수단이 필요하고, AA를 검증 경로에서 빼려면 체인에 둬야 한다(G8·G9)'],
    ['세션', '증명을 max_epoch까지 트랜잭션마다 재사용', '같은 재사용 + 오프체인 세션(요청은 서명만)', '오프체인 서비스도 대상이라 "증명 없는 요청"이 필요 → 세션 중 폐기(root 갱신 시 재검증)를 다뤄야 했다'],
], W)

table_slide('비교 3 — 온체인·서비스 인증·속성·추적·에이전트', H, [
    ['온체인 실행', 'Sui 검증자가 프로토콜 차원에서 지원', '서비스 팩토리의 CREATE2 계정 컨트랙트가 트랜잭션마다 검증', '체인 프로토콜 변경 없이 어떤 EVM 체인에서든. 대가: 트랜잭션당 검증 가스(340k–390k)'],
    ['서비스 인증', '제공자의 redirect URI 검사', 'AA 발급 cert_s를 지갑이 오리진과 대조', '제공자가 서비스를 모르니 제공자가 검사할 수 없다 → 지갑이 한다(G6)'],
    ['속성', '없음(JWT 클레임은 제공자 것)', '커밋 안 사용자 속성 4슬롯, 발급자 비가시', '서비스가 쓸 속성을 발급자에게 안 보이게 실어 둠. 술어 증명은 후속'],
    ['추적', '설계된 절차 없음(제공자+salt 서비스가 사실상 가능)', '2-of-2 트레이스 태그, 운영자 승인 개봉', 'KYC·travel rule: 한 당사자 단독으로는 못 열되 열 수는 있어야(G10)'],
    ['에이전트 플래그', '없음', 'allowAgent를 AA가 서명, 온체인 이벤트로 기록', '세션별 AI 에이전트 허용 여부를 분쟁 증거로(22장)'],
], W, size=12)

bullets('유효기간: 왜 둘 다 벽시계가 아닌가', [
    (0, 'ZKP 안의 값은 전부 증명자가 넣는다 → 회로는 "지금"을 알 수 없다. 만료는 서명된 공개 입력, 비교는 검증자'),
    (0, '검증자가 컨트랙트면 시계는 블록 높이/에포크뿐'),
    (1, 'zkLogin: JWT exp(벽시계)는 온체인에서 검사하지 않음. 지갑이 고른 max_epoch를 nonce에 묶고 체인이 비교'),
    (1, 'zk-Delegation: 지갑이 max_height = ⌈(head+T)/G⌉·G 를 고르고 AA가 서명. 컨트랙트는 block.number ≤ max_height ≤ block.number + L'),
    (0, '양자화(G=100)의 이유: 온체인 공개 입력은 AA도 읽는다. 정확한 높이는 발급 순간의 지문 → 창 안 사용자가 익명 집합. 지갑 자신을 지키는 것이라 지갑이 한다'),
    (1, '대가: 만료 정밀도 G, 최악 노출 T+G. 남는 것: 발급 직후 트랜잭션이라는 시각 상관(창을 넓혀 완화)'),
])

bullets('대가 — 우리가 새로 짊어진 것', [
    (0, 'AA를 새로 운영해야 한다(zkLogin은 Google을 재사용)'),
    (0, 'AA가 로그인 경로에 있다: AA가 죽으면 새 로그인은 멈춘다(진행 중 세션은 계속)'),
    (0, '온체인에서는 트랜잭션마다 검증 가스; root 게시와 겹친 트랜잭션은 실패하고 가스 소각'),
    (0, 'AA가 온체인 max_height를 보므로 발급 창 단위 익명 집합 — 작은 배포에서는 1일 수 있음(양자화는 지갑 몫)'),
    (0, 's_u 분실 시 주소 복구 없음(zkLogin은 salt 서비스 백업)'),
])

prs.save(DST); print('saved', DST, len(prs.slides), 'slides')
