# zk-Delegation 원고 v5 → v6 (2026-09-25): 집합 소속 술어(V7, 2026-09-23) + 세션 단위 폐기(V8, 2026-09-24)
# + 검증자 참조 코드 대조와 팩토리 상한 밴드 + 2026-09-25 전체 코드 리뷰(정규 공개 입력, 마스크 밖 슬롯 거절,
# σ_tx 다이제스트 16워드, 서비스의 set_root 정책 대조) + 실측 갱신. v5 는 건드리지 않고 v6 사본에만 적용한다.
#
# 표기 대응(코드 → 논문): set_sel / set_root 그대로 — 슬롯 번호는 논문 관례 a_1..a_4 이므로(코드는 a₀..a₃) set_sel = k 는 논문의 a_k.
# 세션 리프 mask₂₅₂(Poseidon(5, Cf_s)) = leaf(Cf_s), 사용자 리프 mask₂₅₂(Poseidon(4, Cf_u)) = leaf(Cf_u),
# 회로 ④′ = 관계 R 의 조건 4′, ⑦ = 조건 8. DOMAIN_MODE3_REVOKESESS = D_revsess.
# 팩토리 생성자 인자 maxRootAge / maxLifetime = 논문의 MAX_ROOT_AGE / 수명 상한 L.
# `AttrGate` v2 의 allowedCountriesRoot = 서비스 정책 집합의 root, minAge = 최소 나이.
# 근거: docs/paper/zkd/security_formal.md §10·§11, docs/paper/zkd/conditional_privacy_formal.md 주석(2026-09-24·09-25),
# docs/MODE3_DEMO.md, 실측 results/mode3_review_fixes_20260925.md(컨트랙트 변경 뒤 gas·로그인 왕복) 와
# results/mode3_session_revocation_20260924.md(V7→V8 회로 비교: 제약·증명 시간·zkey).
import copy, shutil
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

SRC = 'documents/2026-09-23_v5_zk-Delegation_research_article.docx'
DST = 'documents/2026-09-25_v6_zk-Delegation_research_article.docx'
shutil.copy(SRC, DST)
d = Document(DST)
XMLNS = '{http://www.w3.org/XML/1998/namespace}space'
PARAS = list(d.paragraphs)                      # 프록시를 붙들어 두어야 id 가 안정된다
IDX = {id(p._p): i for i, p in enumerate(PARAS)}
LOG = []

def pidx(p): return IDX.get(id(p._p), 'new')
def log(p, before, after, kind='edit'): LOG.append((pidx(p), kind, before, after))

def find(prefix, style=None):
    for p in PARAS:
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise KeyError(prefix)

def set_text(p, text):
    rs = p.runs
    assert len(rs) >= 1, p.text[:40]
    rs[0].text = text
    for r in rs[1:]:
        r._r.getparent().remove(r._r)

def rep(p, old, new):
    t = p.text
    assert old in t, (old[:80], t[:120])
    set_text(p, t.replace(old, new, 1))
    log(p, old, new)

def clone_after(p, text, template=None):
    el = copy.deepcopy((template or p)._p)
    ts = el.findall('.//' + qn('w:t'))
    assert len(ts) >= 1
    ts[0].text = text
    ts[0].set(XMLNS, 'preserve')
    for t in ts[1:]:
        t.getparent().remove(t)
    p._p.addnext(el)
    q = Paragraph(el, p._parent)
    LOG.append((f'after {pidx(p)}', 'insert', '', text))
    return q

def set_cell(tc, text):
    ps_ = tc.findall(qn('w:p'))
    for q in ps_[1:]: tc.remove(q)
    q = ps_[0]
    rs = q.findall(qn('w:r'))
    for r in rs[1:]: q.remove(r)
    if not rs: rs = [etree.SubElement(q, qn('w:r'))]
    ts = rs[0].findall(qn('w:t'))
    for t in ts[1:]: rs[0].remove(t)
    if not ts: ts = [etree.SubElement(rs[0], qn('w:t'))]
    ts[0].text = text; ts[0].set(XMLNS, 'preserve')

def set_row(table, key, row):
    for r in table.rows:
        if r.cells[0].text.strip() == key:
            before = ' | '.join(c.text.strip() for c in r.cells)
            for tc, text in zip(r.cells, row): set_cell(tc._tc, text)
            LOG.append(('table', 'row', before, ' | '.join(row)))
            return
    raise KeyError(key)

def row_after(table, key, row):
    tbl = table._tbl
    for tr in tbl.findall(qn('w:tr')):
        if tr.findall(qn('w:tc'))[0].xpath('string(.)').strip() == key:
            new = copy.deepcopy(tr)
            for tc, text in zip(new.findall(qn('w:tc')), row): set_cell(tc, text)
            tr.addnext(new)
            LOG.append(('table', 'row-insert', f'after "{key}"', ' | '.join(row)))
            return
    raise KeyError(key)

BODY = find('A Web3 service must recognize a returning user')          # 본문 문단 템플릿(PARA)
assert len(BODY.text) > 200 and len(BODY.runs) >= 1, BODY.text[:60]

# ---------------------------------------------------------------- Abstract (IEEE Access 상한 250 단어 — v5 는 247 단어였다)
p = find('ABSTRACT ')
assert p.text.startswith('ABSTRACT A Web3 service identifies its users by a blockchain address'), p.text[:90]   # 덮어쓰기 전 v5 초록인지 확인
ABSTRACT = ("ABSTRACT A Web3 service identifies users by a blockchain address: reuse across services and chains makes a user linkable, and a login says nothing about who holds the key. "
  "Delegating authentication to an existing account tells the authority which service the user visits. "
  "zk-Delegation redefines the address as the name under which a service knows a user, derived from one user secret for every chain and service and realized on chain as a contract account whose transactions carry the login's proof. "
  "The authentication authority (AA) verifies once, in zero knowledge, a commitment to the user's identity, secret, and attributes; per login it signs, without a proof, that credential and a session commitment with a block-height expiry, and learns neither pseudonym nor service, not even from the chain. "
  "The wallet can disclose a range of an attested attribute, or its membership in a service's set, in the same proof. "
  "Each service receives a pairwise pseudonym with a zero-knowledge proof of its derivation and non-revocation and cannot link a user across services or chains. "
  "A trace tag, the identity encrypted under a key shared by AA and service, opens a disputed session only with both and an operator. "
  "One proof serves a session. "
  "Revocation is an append-only tree of credentials and sessions whose root the AA publishes on a public chain; users can revoke their account or one session. "
  "A prototype proof has about 37,000 constraints, proves in about 1.2 s, and verifies in milliseconds off chain or 420,000 gas on chain.")
NW = len(ABSTRACT.split()) - 1
assert NW <= 250, NW
set_text(p, ABSTRACT); log(p, 'ABSTRACT (v5, 247 words)', f'ABSTRACT (v6, {NW} words)')

# ---------------------------------------------------------------- Introduction
p = find('zk-Delegation combines address abstraction with delegated authentication')
rep(p, 'the wallet can disclose a chosen range of an attribute to a service or a contract inside the proof, and the AA takes no part in that disclosure.',
    'the wallet can disclose a chosen range of an attribute, or its membership in a set the verifier names, to a service or a contract inside the proof, and the AA takes no part in that disclosure.')
p = find('Statements must also be revocable before they expire')
rep(p, 'The AA appends revoked accounts to an append-only tree and publishes its root; the wallet proves non-membership against the published root inside the same proof,',
    'The AA appends revoked credentials, and revoked single sessions, to an append-only tree and publishes its root; the wallet proves non-membership of both against the published root inside the same proof,')
rep(p, 'A user who suspects loss or theft of their account can revoke it from their own AA account page, without waiting for an administrator.',
    'A user who suspects loss or theft of their account can revoke it, or end one of its sessions, from their own AA account page, without waiting for an administrator.')
p = find('3) We give a pseudonym proof')
rep(p, 'and that it can disclose a chosen range of an attested attribute to a service or a contract.',
    'and that it can disclose a chosen range of an attested attribute, or its membership in a set the verifier names, to a service or a contract.')
p = find('4) We implement a prototype')
rep(p, 'administrator- and user-initiated revocation, on-chain execution from pseudonym accounts, selective disclosure of an attested attribute to a contract,',
    'administrator- and user-initiated revocation of an account and of a single session, on-chain execution from pseudonym accounts, selective disclosure of an attested attribute to a contract as a range or as membership in a set,')

# ---------------------------------------------------------------- Fig. 1 캡션(그림은 다시 그리지 않는다 — 거짓이 된 문장만 고친다)
p = find('FIGURE 1. One user credential')
rep(p, 'Steps 7–8 happen only for a disputed session, after an operator approves; the AA holds nothing per login until then.',
    'Steps 7–8 happen only for a disputed session, after an operator approves. Per login the AA keeps the session commitment it signed together with the expiry, the chain, the flag, and the times of issuance and, if any, revocation, which is what a session revocation names and what Section VII-C accounts for (Sections V-D and VI-D).')

# ---------------------------------------------------------------- V-D 성명 발급: AA 가 세션을 기록한다(V8)
p = find('an EdDSA signature over Baby Jubjub with Poseidon as the message hash and D_cred a domain tag')
rep(p, 'The AA returns (max_height, chainid, allowAgent, σ_AA) and records nothing: the only per-user state it holds is the active credential leaf of Section V-B\', which is what an account revocation publishes, and it holds no record of when or how often a user logged in.',
    'The AA returns (max_height, chainid, allowAgent, σ_AA) and records, under the account, the session it has just signed: Cf_s together with max_height, chainid, allowAgent, the time of issuance, and a revocation time once there is one. Every field of that record is a value the AA saw in the request, so the record tells it nothing new about the login itself; what it adds is state, the number and the issuance times of a user\'s live sessions. It exists for one purpose, to let the user or an administrator name a single session for revocation (Section VI-D), and the AA drops a record once the chain head passes its max_height. The record holds neither the service identifier nor the session key, since both stay inside C_s. The other per-user state is the active credential leaf of Section V-B\', which is what an account revocation publishes.')
p = find('The AA never sees the challenge. A replayed issuance request obtains a fresh signature')
rep(p, 'so no service and no contract accepts anything made from it, and the AA, which records nothing per issuance, is not even burdened with a duplicate. The only per-login state in the system is the service\'s own challenge record, which is what allows the AA to hold nothing that links a user to a login at all.',
    'so no service and no contract accepts anything made from it, and the AA records the same session a second time. Single use of a login rests on the service\'s own challenge record, and the AA\'s session record names a commitment rather than a login, since it carries neither the service nor the challenge.')

# ---------------------------------------------------------------- V-E 증명: 공개 입력 25개, 증인, 조건 8 과 조건 4′
p = find('The wallet derives PPID = Poseidon(uid, s_u, chainid, arid), draws a fresh 250-bit randomness r')
rep(p, 'followed by the nine disclosure inputs disc_mask, disc_lo_1..disc_lo_4, disc_hi_1..disc_hi_4 of condition 7, 23 in all. The witness is uid, s_u, blind_u, blind_s, a_1..a_4, the signature σ_AA, r, and the non-membership path of Section VI.',
    'the nine disclosure inputs disc_mask, disc_lo_1..disc_lo_4, disc_hi_1..disc_hi_4 of condition 7, and the two set-predicate inputs set_sel and set_root of condition 8, 25 in all. The witness is uid, s_u, blind_u, blind_s, a_1..a_4, the signature σ_AA, r, the path of the set predicate, and the two non-membership paths of Section VI.')
p = find('Two range checks belong to the relation.')
rep(p, 'The leaf and the neighboring leaves of the non-membership path are reduced to 252 bits before comparison,',
    'The leaves and the neighboring leaves of the two non-membership paths are reduced to 252 bits before comparison,')
p = find('Condition 5 is a hashed ElGamal encryption')
rep(p, 'and since r is fresh per proof, tags of one user are unlinkable except through the address itself.',
    'and since r is fresh per proof, and constrained in the circuit to be non-zero so that c1 cannot be the identity point, tags of one user are unlinkable except through the address itself.')
p = find('Condition 7 is selective disclosure.')
rep(p, 'The prototype supports ranges over single slots only; set membership and predicates that relate two slots or a slot to the current block are not supported (Section IX).',
    'The predicates are over single slots; predicates that relate two slots are not supported (Section IX).')

q = clone_after(p, 'Condition 8 is set membership, a second form of disclosure beside the range. The public input set_sel names the slot under the predicate, with set_sel = 0 meaning that no set predicate is made, and set_root is the root of a Poseidon Merkle tree of depth 8 whose leaves are the members of the set, sorted, deduplicated, and padded to 256 places with the value 2^64; every attribute is below 2^64, so no attribute can equal the padding and no proof can pass a padding slot off as a member. For set_sel = k ≥ 1 the circuit checks a Merkle path from a_k to set_root, with the path and the index in the witness, so a verifier learns that a_k is one of the members and not which one; for set_sel = 0 it requires set_root = 0. The predicate costs about 1,960 constraints (Section VIII-B) and reveals, in the terms of Section VII-D, that the slot holds one of the members, which is as much as an exact value when the set has one member and less than that otherwise.', template=BODY)
q = clone_after(q, 'A verifier must compare set_root with the root of its own policy set, and this comparison is as necessary as the comparison of pk_AA above. The circuit proves that the disclosed attribute belongs to the tree of set_root; it does not, and cannot, prove whose tree that is. A prover who builds a one-element tree around its own value and presents that root satisfies condition 8 and satisfies no policy, so a verifier that accepts any root accepts any value. Comparing the root binds the set but not the slot: set_sel is a separate public input and the comparison says nothing about it, so a verifier that needs a particular attribute to be in the set must check set_sel itself. The prototype\'s service checks both; on chain the account contract checks the form of set_sel and set_root, and the demonstration target compares set_root with its own and requires set_sel to name the slot its policy is about.', template=BODY)
q = clone_after(q, 'A predicate relative to the current time stays outside the circuit. The circuit has no clock, and a bound the wallet computed from its own clock would be a value the prover chose; the demonstration target therefore reads the block timestamp, converts it to a calendar year, and requires that the disclosed upper bound on the birth year plus a minimum age be at most that year. The proof attests the bound and the verifier applies the clock, which works wherever the verifier has a clock it can show it read, and on chain it has one.', template=BODY)
clone_after(q, 'Conditions 4 and 4′ are the two non-membership proofs. They are made against one and the same root and, in the wallet, against one and the same state of the tree: condition 4 for the user leaf of the credential, condition 4′ for the session leaf of C_s (Section VI-A). Because condition 2 recomputes C_s from the public arid and pk_i, and the AA signature of condition 1 covers Cf_u and Cf_s together, a prover cannot answer condition 4′ with a session commitment other than the one the statement was issued over.', template=BODY)

# ---------------------------------------------------------------- Table 1 (표기) + Table 2 (관계 R)
for t in d.tables:
    heads = [r.cells[0].text.strip() for r in t.rows]
    if heads[0] == 'Symbol':
        set_row(t, 'a_1..a_4, blind_u', ['a_1..a_4, blind_u', 'user attributes (the AA\'s account record) and credential blinding', 'yes', 'a_1..a_4', 'disclosed ranges and set memberships only'])
        row_after(t, 'a_1..a_4, blind_u', ['set_sel, set_root', 'slot named by a set predicate and the root of the set it must belong to', 'yes', 'no', 'yes'])
    if heads[0] == '#':
        row_after(t, '4', ['4′', 'leaf(Cf_s) = mask_252(Poseidon(5, Cf_s)) is not in the same tree under the same root', 'this session has not been revoked as of this root'])
        set_row(t, '5', ['5', '0 < r < 2^250,  c1 = r·B8,  K = r·pk_trace,  c2 = Poseidon(uid, arid) + Poseidon(K.x, K.y)', 'the tag encrypts a service-specific hash of this credential\'s uid under the certified combined key (verifiable encryption)'])
        row_after(t, '7', ['8', 'set_sel ∈ {0, …, 4};  set_sel = 0 ⇒ set_root = 0;  set_sel = k ≥ 1 ⇒ a_k is a leaf of the depth-8 Poseidon tree whose root is set_root', 'the named attribute belongs to the set that set_root commits to, without revealing which member it is; the verifier compares set_root with its own policy set'])

# ---------------------------------------------------------------- V-F 서비스 검사
p = find('The wallet sends the service the proof, its public inputs, and a signature')
rep(p, 'It checks that disc_mask is below 16 and records the disclosure inputs with the session; the prototype\'s service asks for no disclosure at login, and a service policy that does is a matter of configuration.',
    'It requires every public input to be a canonical decimal string, so that the values it checks against its own configuration and the values the argument is verified against are the same values; a verifier that normalizes the inputs for its policy checks but hands the strings it received to the proof verifier can be made to read one number and verify another (Section VIII-C). It checks that disc_mask is below 16, that set_sel is at most 4 and that set_root is zero when set_sel is, and, where the service has a policy set, that set_root is that set\'s root; it records the disclosure and set inputs with the session. The prototype\'s service asks for no disclosure at login, and a service policy that does is a matter of configuration.')

# ---------------------------------------------------------------- V-I 온체인 실행
p = find('A service that acts on a chain deploys, once, an account factory')
assert 'CREATE2(factory, PPID)' in p.text and 'the wallet learns the factory' in p.text, p.text[:120]   # 삽입 앵커 내용 확인
clone_after(p, 'Before it uses that factory the wallet checks what the factory is. Registration with the AA makes a service known, not trusted with the user\'s assets: the service deploys the factory and the verifier itself, and a factory whose verifier accepted every argument would let anyone move the assets of the accounts it created. The wallet therefore compares the deployed bytecode of the factory and of the verifier with its own reference build of the same contracts and proceeds only if they agree, which leaves an attacker with the problem of finding a second preimage of a keccak hash. The comparison masks the factory\'s immutable slots, which is where its constructor arguments sit, so it says nothing about MAX_ROOT_AGE and the lifetime bound L; the wallet reads those two from the factory and requires them to lie in an accepted band, 1 to 200 blocks and 1 to 1600 blocks in the prototype. The band matters in both directions. A MAX_ROOT_AGE that is too large removes the bound of Section VI-C, since the accounts would keep executing long after the AA stopped publishing revocations. MAX_ROOT_AGE = 0 freezes the account permanently: every execute() reverts on the root-age check, and because the account address is derived from the same constructor arguments, no redeployment with another value yields the same address, so whatever was already sent there cannot be reached. The band itself is an operational choice, twice and four times the prototype\'s own values, and nothing in the security argument derives it (Section IX).', template=BODY)

p = find('A transaction is execute(payload, σ, π, public inputs)')
rep(p, 'The wallet signs the payload under sk_i together with the chain identifier, the account address, the account\'s nonce, and the nine disclosure inputs of the proof it attaches, as keccak256(abi.encode(chainid, wallet, to, value, data, nonce, disc_mask, disc_lo, disc_hi)), so that neither another chain nor another account can reuse the signature and a relayer cannot pair the payload with another proof of the same session that discloses something else;',
    'The wallet signs the payload under sk_i together with the chain identifier, the account address, the account\'s nonce, and sixteen words taken from the proof it attaches, as keccak256(abi.encode(chainid, wallet, to, value, data, nonce, disc_mask, disc_lo, disc_hi, set_sel, set_root, max_height, allowAgent, c1.x, c1.y, c2)), so that neither another chain nor another account can reuse the signature and a relayer cannot pair the payload with another proof of the same session. The last five words are what a relayer could otherwise vary while the signature stayed valid: two statements of one user under one session key may agree on every disclosed predicate and still differ in the expiry, the agent flag, and the tag, and those three are what the account emits, so a swapped proof would leave on the public record an expiry, an agent permission, and an opening tag that this transaction never authorized. The digest therefore covers sixteen words while the call-data tail below carries eleven, so the two are no longer the same set of words;')
rep(p, 'that allowAgent is a bit, disc_mask is below 16, and c1 is not the identity point;',
    'that allowAgent is a bit and c1 is not the identity point, that disc_mask is below 16, that set_sel is at most 4 with set_root zero when set_sel is, and that disc_lo_k and disc_hi_k are zero for every slot whose mask bit is clear;')
rep(p, 'appends the nine disclosure inputs to the call data as trailing words in the manner of ERC-2771, performs the call, and emits the outcome together with pk_i, max_height, allowAgent, and the tag, so that a transaction is a transcript in the sense of Section V-G, and a Disclosure event when disc_mask is not zero.',
    'appends the eleven disclosure and set inputs to the call data as trailing words in the manner of ERC-2771, performs the call, and emits the outcome together with pk_i, max_height, allowAgent, and the tag, so that a transaction is a transcript in the sense of Section V-G, and a Disclosure event when disc_mask or set_sel is not zero.')

p = find('A target contract that acts on a disclosed attribute reads the last 288 bytes')
rep(p, 'reads the last 288 bytes of its call data', 'reads the last 352 bytes of its call data')
rep(p, 'The words are appended whether or not disc_mask is zero.', 'The words are appended whether or not disc_mask and set_sel are zero.')
rep(p, 'With the tail always present, the last 288 bytes of the call are the inputs of the proof the account has just verified,',
    'With the tail always present, the last 352 bytes of the call are the inputs of the proof the account has just verified,')
rep(p, 'The one exception is a call with empty payload.data and mask 0, a plain transfer:',
    'The one exception is a call with empty payload.data, mask 0, and set_sel 0, a plain transfer:')
rep(p, 'The prototype\'s target, AttrGate, accepts one claim per account if slots 1 and 2 are disclosed with a_2 equal to a configured country code and the upper bound of a_1 at most a configured birth year, that is, a residency and an age check.',
    'Condition 7 constrains only the slots inside the mask, so the bounds of a slot outside it are values the prover chose; the account contract therefore rejects a non-zero disc_lo_k or disc_hi_k on such a slot rather than zeroing it, so that whoever reads the tail or the event may take a word outside the mask to be zero. This changes neither the relation nor what a service learns, which was always that only the slots inside the mask carry meaning; what it removes is the chance that a third party reads a free word as an attested disclosure. The prototype\'s target, AttrGate, accepts one claim per account if a_1 is disclosed as a range, a_2 is shown to lie in the set whose root the target holds, and the disclosed upper bound on a_1 plus a configured minimum age is at most the year of the current block timestamp, that is, a residency and an age check with the clock on the verifier\'s side.')

p = find('The proof is the session\'s proof.')
rep(p, 'verification, however, is paid at every transaction, about 400,000 gas (Section VIII-C),',
    'verification, however, is paid at every transaction, about 420,000 gas (Section VIII-C),')

# ---------------------------------------------------------------- VI-A 리프 두 종류
p = find('The revocation set, which we also call the revoked-credential list (RCL)')
rep(p, 'which we also call the revoked-credential list (RCL), is an indexed Merkle tree',
    'which we also call the revoked-credential list (RCL), a name kept from before it held session leaves as well, is an indexed Merkle tree')
rep(p, 'Leaves are the 252-bit values leaf(Cf_u) = mask_252(Poseidon(4, Cf_u)), one per user credential, never per statement: a leaf is inserted only when a credential is revoked or retired, never at issuance, so the tree size is the number of revocations, not the number of logins or of users, and the tree carries no record of who logged in where. A service cannot recognize a leaf, since Cf_u appears in no transcript, and the AA cannot recognize a session from a leaf, since the leaf names a credential and not a login.',
    'Leaves are 252-bit values and come in two kinds: a user leaf leaf(Cf_u) = mask_252(Poseidon(4, Cf_u)), one per user credential, and a session leaf leaf(Cf_s) = mask_252(Poseidon(5, Cf_s)), one per revoked session; the domain tags 4 and 5 keep the two kinds apart. A leaf is inserted only when a credential is revoked or retired or a session is revoked, never at issuance, so the tree size is the number of revocations and not the number of logins or of users, and the tree carries no record of who logged in where. Both kinds are hash outputs, so an observer reading the log cannot tell which kind a published leaf is. A service can recognize neither: Cf_u appears in no transcript, and Cf_s, although the service holds the session it belongs to, is a commitment the service cannot compute, since it does not hold blind_s. The proof carries one non-membership witness of each kind against the same root, conditions 4 and 4′ of Table 2.')
p = find('The tree is append-only. Revoking an account inserts one leaf,')
rep(p, 'Retiring a credential because its attribute record changed inserts the same leaf without disabling the account, and the next login issues a new credential.',
    'Retiring a credential because its attribute record changed inserts the same leaf without disabling the account, and the next login issues a new credential. Revoking one session inserts that session\'s leaf and touches nothing else: the credential stays active, the user\'s other sessions run on, and a new login is unaffected.')

# ---------------------------------------------------------------- VI-D 관리자·사용자 폐기
p = find('An administrator revokes an account, or retires its credential')
rep(p, 'An administrator revokes an account, or retires its credential by changing its attribute record; there is no revocation of a single statement, since the AA holds no record of statements (Section IX).',
    'An administrator revokes an account, retires its credential by changing its attribute record, or revokes one session of an account; a user can do the first and the last. A session is named by its commitment: the request carries (uid, Cf_s) together with either an administrator secret or a signature Sign_sk_u(Poseidon(D_revsess, uid, Cf_s, nonce)) under the user\'s long-term key, with D_revsess a domain tag, and the AA verifies the authorization before it looks the session up, so an unauthenticated caller cannot learn from the answer whether a given Cf_s belongs to a given account. The AA refuses a session the chain head has already expired, since the expiry stops it in any case, and repeating a request revokes the same session again and changes nothing. A service cannot revoke a session of its own users, because it never sees Cf_s.')

# ---------------------------------------------------------------- VII-C 서비스 비관측: AA 의 세션 기록
p = find('At credential issuance the AA receives (uid, C_u) and π_u')
rep(p, 'Because the AA keeps no record of issuances, what it can later join is only what it observes at the moment of issuance and what the chain shows (next paragraph).',
    'The AA does keep the session record of Section V-D, but every field of it is a value the AA saw in the request, so what it can later join is what it saw at issuance and what the chain shows (next paragraph); what the record changes is that the AA need not have been watching at the moment of issuance in order to hold those values.')
p = find('The chain is a second channel.')
rep(p, 'and, for a transaction that discloses an attribute range, that range (Section VII-D).',
    'and, for a transaction that discloses an attribute range or a set membership, that predicate (Section VII-D).')
rep(p, 'and an AA that logged its issuances could narrow a transaction to the users issued in that window; the prototype\'s AA keeps no such log, so it would have to observe issuances as they happen.',
    'and the AA\'s session records, which carry chainid, allowAgent, max_height, and a time of issuance, narrow a transaction to the users whose live sessions match. Before session revocation the AA kept no such record and would have had to watch issuances as they happened; holding them is the price of being able to name a session, and what bounds it is that the AA drops a record once the chain head passes its expiry.')

# ---------------------------------------------------------------- VII-D 속성 사용처 비관측
p = find('The attributes occupy slots of the user credential')
rep(p, 'so a service or a contract learns nothing about a slot the user has not disclosed, and about a disclosed slot exactly the range in the public inputs.',
    'so a service or a contract learns nothing about a slot the user has not disclosed, and about a disclosed slot exactly the range in the public inputs, or, under a set predicate, only that the slot holds one of the set\'s members and not which one.')
p = find('A disclosed range shrinks the anonymity set of Section VII-C.')
rep(p, 'The AA holds every user\'s attributes, so from a transaction that discloses ranges it can compute the set of users whose recorded attributes satisfy them;',
    'The AA holds every user\'s attributes, so from a transaction that discloses ranges or a set membership it can compute the set of users whose recorded attributes satisfy them;')
rep(p, 'An exact value (disc_lo = disc_hi) shrinks the set more than a range does, and the disclosed values are public for as long as the chain is, bound to the account and thus to the service.',
    'An exact value (disc_lo = disc_hi) shrinks the set more than a range does, and a set predicate shrinks it as far as the set is small, membership in a set of one being an exact value. The disclosed values are public for as long as the chain is, bound to the account and thus to the service.')

# ---------------------------------------------------------------- VII-E 서비스 인증: 참조 코드 대조
p = find('cert_s binds arid, the origin, and the combined trace key under the AA key,')
rep(p, 'so a credential issued for one service verifies only against that service\'s values.',
    'so a credential issued for one service verifies only against that service\'s values. What the certificate also does not establish is what the service\'s contracts do. A registered service deploys its own factory and verifier, so a service honest at its origin could still deploy a verifier that accepts any argument and then move the assets of the accounts its factory created; approval and the origin check say nothing about this. The wallet closes it by comparing the deployed code of both contracts with its own reference build before it uses a factory, and by requiring the factory\'s two bounds to lie in an accepted band (Section V-I); an attacker then needs a second preimage of a keccak hash. Both checks rest on the wallet\'s reference build being the genuine one, which is a supply-chain assumption rather than a protocol property (Section IX).')

# ---------------------------------------------------------------- VII-F 세션 바인딩과 재생
p = find('Off chain, a login is accepted only if σ is a signature under pk_i')
rep(p, 'and because the AA signature covers C, which commits to arid.', 'and because the AA signature covers Cf_s, which commits to arid.')
rep(p, 'and the AA merely extends its record of C.', 'and the AA merely records the same session a second time.')
rep(p, 'Session requests are signed under sk_i, which the wallet holds and C commits to;',
    'Session requests are signed under sk_i, which the wallet holds and C_s commits to;')
rep(p, 'a statement is meant to be reused for many transactions.',
    'a statement is meant to be reused for many transactions. The same signature covers the expiry, the agent flag, and the tag of the proof it accompanies, so a relayer holding one transaction of a session cannot pair its payload with a different proof of the same user and session key and thereby put on the record an expiry, an agent permission, or an opening tag that the user did not authorize for that transaction (Section V-I).')

# ---------------------------------------------------------------- VII-H 승인 개봉: AA 가 남기는 것
p = find('The tag (c1, c2) is a hashed ElGamal ciphertext under pk_trace')
rep(p, 'here the AA keeps no per-login record at all.',
    'here the AA keeps no record from which a session could be attributed: the per-session record of Section V-D holds neither the service, nor the pseudonym, nor the tag.')

# ---------------------------------------------------------------- VII-G 폐기 건전성: 세션 리프
p = find('Once leaf(Cf_u) is in the tree and the root is published,')
rep(p, 'Every statement of the user, at every service, commits to the same Cf_u, so one leaf ends all of them.',
    'Every statement of the user, at every service, commits to the same Cf_u, so one leaf ends all of them. A session leaf ends one session and only that one, by the same argument applied to condition 4′: Cf_s is recomputed in condition 2 from the public arid and pk_i and is covered by the AA signature together with Cf_u, so a prover cannot answer condition 4′ with a session commitment other than the signed one.')
rep(p, 'A revoked account cannot obtain a new credential or a new statement, because the AA checks the flag and the active credential before signing.',
    'A revoked account cannot obtain a new credential or a new statement, because the AA checks the flag and the active credential before signing; a revoked session does not stop new logins, which is the point of the finer granularity.')

# ---------------------------------------------------------------- V-G 개봉: AA 가 남기는 것(V8 로 "아무것도 안 남긴다" 가 거짓이 됐다)
p = find('Verifying the transcript before opening is what confines the procedure')
rep(p, 'The AA stores nothing per login; the tag lives only in the transcript the service keeps, and the AA\'s share is applied only at an approved opening.',
    'The AA stores no transcript and no tag: its per-session record of Section V-D holds the session commitment and the values it signed, and nothing from which a login or a pseudonym could be reconstructed. The tag lives only in the transcript the service keeps, and the AA\'s share is applied only at an approved opening.')
p = find('Two designs were considered for attribution.')
rep(p, 'but tracing remains possible for as long as the two shares exist and needs no per-login state at the AA.',
    'but tracing remains possible for as long as the two shares exist and needs no record at the AA from which a session could be attributed.')

# ---------------------------------------------------------------- VIII-A 프로토타입
p = find('The prototype consists of four Node.js processes and five contracts:')
rep(p, 'service registration with operator approval, statement issuance, revocation, publication,',
    'service registration with operator approval, statement issuance, revocation of an account and of a single session, publication,')
rep(p, 'verifies service certificates, obtains its credential and requests statements, keeps a checkpoint of the revocation tree, derives addresses, produces and caches proofs, signs session requests, and takes a per-transaction disclosure choice, a lower and an upper bound per slot, which it checks against its own attributes before proving;',
    'verifies service certificates, compares a service\'s deployed factory and verifier with its own reference build and their bounds with an accepted band before it uses them, obtains its credential and requests statements, keeps a checkpoint of the revocation tree, derives addresses, produces and caches proofs, signs session requests and requests to revoke one of its own sessions, and takes a per-transaction disclosure choice, a lower and an upper bound per slot and a set to test one slot against, which it checks against its own attributes before proving;')

# ---------------------------------------------------------------- VIII-B 회로 크기와 비용
p = find('Two earlier versions of the circuit give the cost of each addition.')
rep(p, 'Two earlier versions of the circuit give the cost of each addition.', 'Successive versions of the circuit give the cost of each addition.')
rep(p, 'brought the circuit to the 25,369 constraints of Table 3:', 'brought the circuit to 25,369:')
rep(p, 'The Pedersen commitment remains the largest single block;',
    'Adding the set-membership predicate of condition 8, one depth-8 Poseidon path with the one-hot decomposition of set_sel, brought it to 27,329; adding the second non-membership proof of condition 4′ brought it to the 37,130 constraints of Table 3, an increase of 9,801 that is one more indexed-tree non-membership gadget over a depth-32 path together with the Poseidon hash of the session leaf. That step is the largest in the circuit\'s history, and what it buys is revocation at the granularity of a session; its 9,801 constraints suggest that the two non-membership proofs together account for more than half the circuit, although we did not itemize the first. The Pedersen commitments are therefore no longer the largest single block;')
p = find('Per-login cost is one issuance round trip of about 0.12 s through the wallet,')
rep(p, 'Per-login cost is one issuance round trip of about 0.12 s through the wallet,', 'Per-login cost is one issuance round trip of about 0.13 s through the wallet,')
rep(p, 'is about 30 ms and the rest is two chain reads and HTTP, plus one Groth16 proof of about 0.8 s and one ECDSA signature;',
    'was about 30 ms when last measured on its own, on 2026-09-21 and before the session record was added, and the rest is two chain reads and HTTP, plus one Groth16 proof of about 1.2 s and one ECDSA signature;')
rep(p, 'An on-chain transaction costs the same proof, reused, plus 402,000 gas for verification and a call without value, 451,000 for a value transfer to a fresh address, and 445,000 for a call that discloses two attribute ranges to the demonstration target, nineteen to twenty-one times a plain transfer; this is the cost of verifying at every transaction that Section V-I accepts. Of this, the nine disclosure inputs cost about 62,800 gas in the verifier over the 14-input version, and the disclosure path, the trailing words, the event, and the target\'s own checks, about 42,400 over a call without disclosure.',
    'An on-chain transaction costs the same proof, reused, plus 417,533 gas for verification and a call without value, 422,366 for a call that discloses one attribute range, 422,982 for one that shows a set membership, and 456,439 for one that does both and runs the demonstration target\'s claim, twenty to twenty-two times a plain transfer of 21,000 gas; this is the cost of verifying at every transaction that Section V-I accepts. The three variants that do no work at the target differ by less than six thousand gas, because verification dominates and its cost follows the number of public inputs: the set predicate added two of them and about 13,900 gas, from 402,130 to 415,991 on the same call, which is what the per-input cost measured earlier for the nine disclosure inputs predicts, and session revocation added no public input and 24 gas. What session revocation costs is paid by the prover, about 0.36 s of the proving time above and nothing measurable at any verifier, off chain or on. Deploying an account through the factory consumed 939,808 gas.')
rep(p, 'a login took 1.05 s (Groth16 proof 0.81 s, issuance 0.12 s, tree sync 0.03 s), a re-validation with the cached proof 31 ms, an on-chain transaction with the cached proof 151 ms including mining on the local node, and a transaction that discloses two attribute ranges, and therefore needs a fresh proof, 1.00 s.',
    'a login took 1.50 s (Groth16 proof 1.26 s, issuance 0.13 s, tree sync 0.03 s), a re-validation with the cached proof 29 ms, an on-chain transaction with the cached proof 164 ms including mining on the local node, and a transaction that discloses a range and a set membership, and therefore needs a fresh proof, 1.45 s.')

# ---------------------------------------------------------------- Table 3 (실측)
for t in d.tables:
    heads = [r.cells[0].text.strip() for r in t.rows]
    if 'R1CS constraints' in heads:
        set_row(t, 'R1CS constraints', ['R1CS constraints', '37,130'])
        set_row(t, 'Public inputs / private inputs', ['Public inputs / private inputs', '25 / 155'])
        set_row(t, 'Revocation tree depth', ['Revocation tree depth / attribute set tree depth', '32 / 8'])
        set_row(t, 'Proving key (zkey) / witness generator (wasm)', ['Proving key (zkey) / witness generator (wasm)', '23.4 MB / 4.8 MB'])
        set_row(t, 'Groth16 proving time (median)', ['Groth16 proving time (median)', '1,192 ms'])
        set_row(t, 'Groth16 verification time (off chain)', ['Groth16 verification time (off chain)', '10.0 ms'])
        set_row(t, 'execute() gas on chain (verification + call)', ['execute() gas on chain (verification + call)', '417,533 (no disclosure) / 422,366 (range) / 422,982 (set) / 456,439 (range + set + AttrGate claim)'])
        row_after(t, 'execute() gas on chain (verification + call)', ['Account deployment gas (CREATE2 through the factory)', '939,808'])
        set_row(t, 'Statement issuance (per login, no proof): wallet round trip / AA work', ['Statement issuance (per login, no proof): wallet round trip / AA work', '127 ms / 30 ms (AA work last measured 2026-09-21)'])
        set_row(t, 'Login end to end through the wallet (issuance + proof) / re-validation (cached proof)', ['Login end to end through the wallet (issuance + proof) / re-validation (cached proof)', '1,497 ms / 29 ms'])
        set_row(t, 'Disclosure inputs / verifier gas over the 14-input version', ['Disclosure and set-predicate public inputs / verifier gas per added public input', '11 (mask, 4 lower, 4 upper bounds, set_sel, set_root) / about 7,000 gas (62,809 for the nine disclosure inputs, 14 → 23; 13,861 for set_sel and set_root, 23 → 25)'])
        row_after(t, 'Disclosure and set-predicate public inputs / verifier gas per added public input', ['Spread of execute() over calls with no disclosure, one range, and one set membership (no work at the target)', 'under 5,500 gas'])
        row_after(t, 'Root freshness bound: service (head age, root age) / account contract', ['Factory bounds the wallet accepts: MAX_ROOT_AGE / lifetime bound L', '1..200 blocks / 1..1600 blocks'])

# ---------------------------------------------------------------- VIII-C 기능 검증
p = find('An end-to-end suite runs the AA, wallet agent, service, and chain in isolation')
rep(p, 'a pseudonym computed for another chain, a commitment that does not open to the public arid, a revoked leaf, an agent flag that is not a bit, and a tag whose ciphertext, key, or randomness has been altered;',
    'a pseudonym computed for another chain, a commitment that does not open to the public arid, a revoked user leaf, a revoked session leaf, a set path that leads to another root, an agent flag that is not a bit, and a tag whose ciphertext, key, or randomness has been altered;')
rep(p, 'The contract suite adds the 23-input verifier, the trailing disclosure words and the Disclosure event, a signature whose digest omits the disclosure inputs, a plain transfer to a receive-only contract, and AttrGate\'s own checks: acceptance, a country mismatch, an age mismatch, and a direct call from an externally owned account.',
    'The contract suite adds the 25-input verifier, the eleven trailing words and the Disclosure event, a signature whose digest omits one of the sixteen signed words, a public word outside the mask that the contract refuses, and a plain transfer to a receive-only contract, together with AttrGate\'s own checks: acceptance, a country outside the policy set, an age below the minimum, and a direct call from an externally owned account. Session revocation has its own end-to-end scenarios: one session of a user revoked while another continues, a re-validation of the revoked session refused and the wallet discarding it, a request to revoke a session of another account refused at the signature, a session already expired refused, and a repeated request answered idempotently.')
p = find('The tree checkpoint has its own suite:')
rep(p, 'were fixed, and each fix is held by a regression test that reproduces its scenario.',
    'were fixed, and each fix is held by a regression test that reproduces its scenario. A second review, after set membership and session revocation were added, found five further defects, which the sections above now describe as the design. The service verifier normalized the public inputs for its own policy checks but handed the strings it had received to the proof verifier, whose parser reads a string that is not all digits as hexadecimal, so a login could carry inputs that the service recorded as a disclosure while the argument was verified against other values; and because the service stored the normalized values, such a session could no longer be opened either, its stored transcript failing verification at the AA. The account contract let a prover put arbitrary values in the bounds of masked-off slots and wrote them to the call-data tail and to the event. The transaction signature digest omitted the expiry, the agent flag, and the tag. The off-chain verifier did not compare set_root with its own policy set, so a prover could present a set of its own choosing. And the wallet compared a factory\'s code but not the two constructor arguments that the comparison masks. Each is closed and each fix has a regression test. None of the five changed the circuit, the proving key, or the on-chain verifier, which is why the figures below separate the cost of the circuit from the cost of the contracts.')
p = find('Publication cost was measured on the local Hardhat node')
rep(p, 'An execute() call with a valid proof and no disclosure consumed 402,130 gas for a call without value to an existing account and 419,262 gas for the first transaction of an account; the contract suite\'s value transfer to a fresh address consumed 450,675 gas; and a call disclosing two attribute ranges to AttrGate, including its claim logic, consumed 444,509 gas; in every case the bulk is the pairing check of the Groth16 verifier. Deploying an account through the factory consumed 852,198 gas once; a heartbeat is a publication with no leaves and consumed 40,285 to 40,305 gas.',
    'An execute() call with a valid proof and no disclosure consumed 417,533 gas for a call without value and 434,665 gas for the first transaction of an account; a call disclosing one attribute range consumed 422,366 gas, one showing a set membership 422,982 gas, and one doing both and running AttrGate\'s claim 456,439 gas; in every case the bulk is the pairing check of the Groth16 verifier. Deploying an account through the factory consumed 939,808 gas once; a heartbeat is a publication with no leaves and consumed 40,273 to 40,305 gas. The contract changes of the second review cost between 1,277 and 1,554 gas per execution, which is the wider signature digest together with the check on the words outside the mask, and about 31,600 gas on an account deployment, which is the larger account bytecode; the generated verifier\'s own deployment cost did not change by a single unit, which is the evidence that it was not touched.')
rep(p, 'All figures are medians of ten runs measured on 2026-09-22 (contract suite: one run; individual runs differ by tens of gas with the proof encoding); the tree-sync figures of Table 3 are medians of five runs measured on 2026-09-23 with leaves published fifty per transaction.',
    'The execute(), deployment, and heartbeat figures are medians of ten runs measured on 2026-09-25, a run that also published a single leaf at 43,856 gas; the publication-batch figures were measured on 2026-09-22 against a freshly deployed log and were not repeated, the log contract being unchanged; the circuit figures of Table 3, from the constraint count to the proving and verification times, are from 2026-09-24, when the session leaf entered the circuit, which the contract changes of 2026-09-25 did not touch; the tree-sync figures are medians of five runs measured on 2026-09-23 with leaves published fifty per transaction. The proving time therefore appears twice: as 1,192 ms in Table 3, the prover measured alone on 2026-09-24, and as 1.26 s in the login above, measured through the wallet\'s interface on 2026-09-25. Individual runs differ by tens of gas with the proof encoding.')

# ---------------------------------------------------------------- IX Limitations
p = find('The AA sees the chain identifier and, through the chain, the quantized expiry')
rep(p, 'A transaction that discloses attribute ranges narrows that set further,', 'A transaction that discloses attribute ranges or a set membership narrows that set further,')
rep(p, 'an exact value narrows it more than a range, and the disclosed values remain public on the chain (Section VII-D).',
    'an exact value narrows it more than a range, a set predicate narrows it as far as the set is small, and the disclosed values remain public on the chain (Section VII-D). The rule has no exception: the anonymity set shrinks with whatever a disclosure reveals, and the only mitigation the design offers is to reveal less.')
p = find('Attribute predicates are ranges over single slots,')
rep(p, 'Attribute predicates are ranges over single slots, with equality as the degenerate range, over a 64-bit integer encoding of each attribute. A predicate relative to the current time, such as an age computed from a birth year at the block in which the proof is verified, must be expressed by the wallet as a bound on the stored value and goes stale as time passes; set membership and predicates across slots are not supported.',
    'Attribute predicates are a range over one slot, with equality as the degenerate range, and membership of one slot in a set of at most 256 values, over a 64-bit integer encoding of each attribute; predicates that relate two slots are not supported. A predicate relative to the current time is not in the circuit. The wallet discloses a bound on the stored value and a verifier that wants an age applies its own clock to that bound, which works for a contract, whose clock is the block timestamp, and not for an off-chain verifier that cannot show which clock it read. Comparing set_root binds the set and not the slot: set_sel is a separate public input, so a verifier that needs a particular attribute to be in the set must check set_sel as well, and one that omits that check is satisfied by a membership proof about another slot.')
p = find('Service registration is approved by an operator,')
rep(p, 'Self-revocation has no rate limit on password attempts.',
    'Self-revocation has no rate limit on password attempts. Opening is not session-scoped: the tag\'s plaintext is a hash of the identity and the service identifier and carries nothing of the session, so opening one transcript yields the user behind it and not the session, and there is no way to authorize an opening that reveals less.')
clone_after(p, 'The wallet\'s checks on a service\'s contracts have two soft edges. The comparison with a reference build is only as good as that build: a wallet whose reference artifacts are not the genuine ones compares against the wrong thing, and a legitimate service compiled with other settings is refused. The band in which the factory\'s root-age and lifetime bounds must lie was set at twice and four times the prototype\'s own values; it is an operational choice, and no part of the security argument derives it. A service that needs a value outside the band cannot be used by this wallet, and a value inside it may still be worse for the user than the prototype\'s.', template=BODY)

p = find('Revocation is a timing channel and has one granularity.')
set_text(p, 'Revocation is a timing channel. Revoking an account inserts one leaf, that of its active credential, so every session of that user, at every service, stops at the same root transition; two services that compare which of their sessions a publication ended can link them, and the anonymity set of a revoked user is the number of leaves in that publication. Revoking a single session inserts one leaf as well, and although a published leaf does not say which kind it is and a service cannot compute Cf_s from anything it holds, the moment at which the session ends is still visible to the service whose session it was. Batching revocations over a longer window enlarges the set; the prototype publishes on command and does not batch for this purpose. A session revocation takes effect at the next publication and not before, so it does not stop a transaction mined in the window between the request and the publication, and sessions issued before the AA began to record them cannot be named at all and end only by expiry. The tree remains append-only, so it now grows with revoked sessions as well as with revoked credentials; folding away the leaves of sessions that have expired is future work.')
log(p, 'Revocation is a timing channel and has one granularity …', '… two leaf kinds, session revocation, its window and its append-only cost')

# ---------------------------------------------------------------- X Conclusion
p = find('This paper separated two things that Web3 services conflate:')
rep(p, 'lets a user cut off a lost device from any browser.', 'lets a user cut off a lost device, or a single session of it, from any browser.')
p = find('The prototype shows that the whole flow fits in one Groth16 argument')
rep(p, 'The prototype shows that the whole flow fits in one Groth16 argument of about twenty-five thousand constraints, proved in under a second and verified in milliseconds, with a per-request cost of a single signature off chain and one verification per transaction on chain, at about 400,000 gas; the same argument discloses a chosen range of an attested attribute to a contract for about 42,000 gas more.',
    'The prototype shows that the whole flow fits in one Groth16 argument of about thirty-seven thousand constraints, proved in about 1.2 s and verified in milliseconds, with a per-request cost of a single signature off chain and one verification per transaction on chain, at about 420,000 gas; the same argument discloses a chosen range of an attested attribute to a contract, or its membership in a set the verifier names, and proves that neither the credential nor the session has been revoked, and the cost of verifying it follows the number of public inputs and not the size of the circuit: the two inputs of the set predicate added about three per cent, and session revocation, which added ten thousand constraints and no input, added nothing measurable.')
rep(p, 'Future work is to extend the range predicates over the committed slots to set membership and time-relative predicates, to constrain r ≠ 0 in the circuit rather than in the verifiers, to let a contract cache a verified session',
    'Future work is to extend the predicates over the committed slots to relations between two slots, to fold away the leaves of expired sessions so that the revocation tree grows with what is live rather than with everything ever revoked, to let a contract cache a verified session')

d.save(DST)
print('saved', DST)
print(f'abstract words: {NW}')
for i, kind, before, after in LOG:
    b = (before[:70] + '…') if len(before) > 70 else before
    a = (after[:70] + '…') if len(after) > 70 else after
    print(f'[{i}] {kind}: {b!r} -> {a!r}')
