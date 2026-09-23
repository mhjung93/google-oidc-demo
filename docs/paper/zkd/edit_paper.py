# zk-Delegation 원고에 승인된 개봉(2-of-2 트레이스 태그)을 반영한다. 스펙 2026-09-16 §11 목록.
import copy, sys
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

SRC = 'documents/2026-09-15_v1_zk-Delegation_research_article.docx'
DST = 'documents/2026-09-16_v2_zk-Delegation_research_article.docx'
FIG = '/tmp/zkd_paper/fig1_login.png'

d = Document(SRC)
ps = d.paragraphs
XMLNS = '{http://www.w3.org/XML/1998/namespace}space'

def find(prefix, style=None):
    for p in d.paragraphs:
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise KeyError(prefix)

def set_text(p, text):
    rs = p.runs
    assert len(rs) >= 1, p.text[:40]
    rs[0].text = text
    for r in rs[1:]:
        r._r.getparent().remove(r._r)

def replace_in(p, old, new):
    t = p.text
    assert old in t, (old[:50], t[:60])
    set_text(p, t.replace(old, new, 1))

def clone_after(p, text, template=None):
    el = copy.deepcopy((template or p)._p)
    ts = el.findall('.//' + qn('w:t'))
    assert len(ts) >= 1
    ts[0].text = text
    ts[0].set(XMLNS, 'preserve')
    for t in ts[1:]:
        t.getparent().remove(t)
    p._p.addnext(el)
    return Paragraph(el, p._parent)

PARA_T = find('A Web3 service must recognize a returning user')          # PARA 템플릿
SUB_T = find('A. ENTITIES')                                                # 소제목 템플릿
REF_T = None
for p in d.paragraphs:
    if p.style.name == 'Ref': REF_T = p

# ---------------- 초록·색인어 ----------------
abs_p = find('ABSTRACT ')
replace_in(abs_p, 'Services cannot link a user across services or across chains. ',
           'Services cannot link a user across services or across chains. Each login also carries a trace tag: the user identifier encrypted under a key that the AA and the service hold as two shares, so a disputed session can be opened only when both, and an operator, act together, and never by either alone. ')
replace_in(abs_p, 'about 22,000 constraints', 'about 25,000 constraints')
replace_in(find('INDEX TERMS'), 'revocation.', 'revocation, conditional traceability.')

# ---------------- 서론 ----------------
p = find('zk-Delegation combines address abstraction with delegated authentication')
replace_in(p, 'It learns a stable local name and nothing about the identity behind it. ',
           'It learns a stable local name and nothing about the identity behind it. The proof also carries a trace tag, the user identifier encrypted under a key that the AA and the service hold as two shares; opening it requires a partial decryption from each and an operator\'s approval, so a disputed session can be attributed without giving either party the power to do so alone. ')
set_text(find('2) We construct zk-Delegation'),
         '2) We construct zk-Delegation, a delegated-authentication protocol in which the AA signs a user-chosen commitment and a service-chosen challenge and learns neither the service nor the pseudonym nor the user\'s attributes, while the service authenticates itself to the user through an AA-issued certificate, and in which every login carries a 2-of-2 trace tag that supports authorized opening of a disputed session.')
replace_in(find('3) We give a pseudonym proof'), 'the AA signature, and non-membership', 'the AA signature, the trace tag, and non-membership')
set_text(find('4) We implement a prototype'),
         '4) We implement a prototype covering service registration and approval, statement issuance, pseudonym verification, session re-validation, administrator- and user-initiated revocation, and authorized opening on a local chain, and report its behavior.')
replace_in(find('The remainder of the paper is organized'), 'including service authentication, attribute privacy, and the pseudonym proof.',
           'including service authentication, attribute privacy, the pseudonym proof, and authorized opening.')

# ---------------- III-A ----------------
p = find('zk-Delegation has five parties.')
replace_in(p, 'A service (the relying party) verifies statements and identifies users by their addresses. ',
           'A service (the relying party) verifies statements and identifies users by their addresses; it holds a signing key for opening requests and one share of a trace key, of which the AA holds the other share. ')
p = find('Table 1 lists the values each party holds.')
replace_in(p, 'A service holds an identifier arid assigned by the AA at registration, a certificate cert_s over that identifier and its origin, and its own view of the revocation chain.',
           'A service holds an identifier arid assigned by the AA at registration, a certificate cert_s over that identifier, its origin, and the combined trace key pk_trace, and its own view of the revocation chain.')

# ---------------- 표 1 ----------------
def set_cell(tc, text):
    ps_ = tc.findall(qn('w:p'))
    for q in ps_[1:]: tc.remove(q)
    q = ps_[0]
    rs = q.findall(qn('w:r'))
    for r in rs[1:]: q.remove(r)
    if not rs: rs = [etree.SubElement(q, qn('w:r'))]
    ts = rs[0].findall(qn('w:t'))
    for t in ts[1:]: rs[0].remove(t)
    if not ts: ts = [etree.SubElement(rs[0], qn('w:t'))]
    ts[0].text = text; ts[0].set(XMLNS, 'preserve')

def add_rows(table, rows):
    tbl = table._tbl
    last = tbl.findall(qn('w:tr'))[-1]
    for row in rows:
        tr = copy.deepcopy(last)
        for tc, text in zip(tr.findall(qn('w:tc')), row):
            set_cell(tc, text)
        tbl.append(tr)

t1 = d.tables[0]
for r in t1.rows:
    if r.cells[0].text == 'cert_s':
        set_cell(r.cells[1]._tc, 'AA signature over (arid, origin, pk_trace), issued on approval')
add_rows(t1, [
    ['pk_service, sk_service', 'service signing key for opening requests', 'no', 'pk_service', 'yes'],
    ['x_svc, X_svc', 'service share of the trace key, X_svc = x_svc·B8', 'no', 'X_svc', 'yes'],
    ['x_AA', 'AA share of the trace key for this service', 'no', 'yes', 'no'],
    ['pk_trace', 'combined trace key X_svc + x_AA·B8', 'per login (from cert_s)', 'yes', 'yes'],
    ['tag = (c1, c2)', 'uid encrypted under pk_trace (hashed ElGamal)', 'creates', 'at opening', 'yes'],
])

# ---------------- III-B / III-C ----------------
set_text(find('Collusion between the AA and a service is outside the model'),
         'Collusion between the AA and a service is outside the model, as is active misbehavior by the AA such as publishing a forged root or choosing a trace key whose share it does not honestly hold. What the model does include is authorized opening: the AA and a service may jointly attribute one disputed session to a user identifier, but only through the procedure of Section V-G, which requires a partial decryption from the service, a partial decryption from the AA, and an operator\'s approval; neither party can open a session alone, and a leaked service log yields only ciphertexts. Sybil resistance across several AA accounts is outside the model; the goal is one address per account and service, not one account per person. Clocks of the AA and the services are assumed to agree to within a few minutes.')
clone_after(find('G9 Liveness independence.'),
            'G10 Authorized opening. A disputed session can be attributed to a user identifier only with the cooperation of both the service and the AA and the approval of an operator; neither the service nor the AA alone, nor an outsider holding the service\'s records, can do so, and opening reveals only that session\'s identifier.')

# ---------------- V-A + 그림 ----------------
p = find('Fig. 1 shows one login.')
replace_in(p, 'its identifier, its certificate, and a fresh challenge r_s (step 1)', 'its identifier, its certificate, the combined trace key pk_trace, and a fresh challenge r_s (step 1)')
replace_in(p, 'produces a zero-knowledge proof that binds address, commitment, signature, and non-revocation, and presents address, session key, and proof to the service (step 4)',
           'produces a zero-knowledge proof that binds address, commitment, signature, non-revocation, and a trace tag that encrypts uid under pk_trace, and presents address, session key, tag, and proof to the service (step 4)')
replace_in(p, 'Later requests in the session are signed under the session key (step 6).',
           'Later requests in the session are signed under the session key (step 6). If a session is later disputed, the service asks the AA to open it (step 7): it sends the login transcript and its own partial decryption of the tag, an operator approves, and the AA completes the decryption and returns the identifier (step 8, Section V-G).')
# 그림 교체
for q in d.paragraphs:
    if q._p.findall('.//' + qn('w:drawing')):
        for r in list(q.runs):
            r._r.getparent().remove(r._r)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run().add_picture(FIG, width=Inches(6.5))
        break
set_text(find('FIGURE 1.'),
         'FIGURE 1. One login (steps 1–6) and one authorized opening (steps 7–8) in zk-Delegation. Solid arrows carry what the AA sees at login (C, chainid, r_s, exptime); dashed arrows carry values the AA never sees at login (arid, cert_s, pk_trace, PPID, pk_i, tag, attributes); dotted arrows are chain reads. Steps 7–8 happen only for a disputed session, after an operator approves; the AA holds nothing per login until then.')

# ---------------- V-B / V-C ----------------
set_text(find('A service registers by sending its display name and its origin.'),
         'A service registers by sending its display name, its origin, a secp256k1 signing key pk_service, and a Baby Jubjub point X_svc = x_svc·B8 whose scalar it keeps. The AA checks that X_svc is a valid subgroup point, records the request as pending, and an operator approves or denies it. On approval the AA draws its own share x_AA, computes the combined trace key pk_trace = X_svc + x_AA·B8, assigns a random 250-bit identifier arid, and signs cert_s = Sign_AA(Poseidon(D_cert, arid, F(origin), pk_trace.x, pk_trace.y)) with domain tag D_cert and F the origin string hashed with SHA-256 and truncated to 250 bits. The service persists arid, pk_trace, and cert_s and re-verifies the certificate under the AA public key at every start; until approval it runs in a waiting state and refuses logins. The certificate covers pk_trace so that a service cannot hand the wallet a key whose full secret it knows: the wallet encrypts only under the key the AA has certified, and the AA adds its share before certifying.')
p = find('To begin a login the service sends the wallet')
replace_in(p, '{arid, origin, cert_s, r_s}', '{arid, origin, cert_s, pk_trace, r_s}')
replace_in(p, 'The wallet verifies cert_s under the AA public key and compares', 'The wallet verifies cert_s under the AA public key over exactly (arid, origin, pk_trace) and compares')

# ---------------- V-E ----------------
set_text(find('The wallet derives PPID = Poseidon(uid, s_u, chainid, arid) and produces'),
         'The wallet derives PPID = Poseidon(uid, s_u, chainid, arid), draws a fresh 250-bit randomness r, and produces a Groth16 argument for the relation in Table 2. The public inputs are, in this order, PPID, arid, pk_i, exptime, chainid, r_s, root, pk_AA.x, pk_AA.y, pk_trace.x, pk_trace.y, c1.x, c1.y, c2. The witness is uid, s_u, blind, a_1..a_4, the signature σ_AA, r, and the non-membership path of Section VI. The circuit computes C_pt and C from the witness, so a prover cannot substitute a commitment the AA did not sign.')
add_rows(d.tables[1], [['5', 'r < 2^250,  c1 = r·B8,  K = r·pk_trace,  c2 = uid + Poseidon(K.x, K.y)', 'the tag encrypts this statement\'s uid under the certified combined key (verifiable encryption)']])
p = find('Two range checks belong to the relation.')
clone_after(p, 'Condition 5 is a hashed ElGamal encryption [19] of uid under pk_trace, computed inside the circuit from the same uid signal that opens the commitment and derives the address. The service therefore knows that the ciphertext it stores, if ever opened, yields exactly the identity behind this statement, without learning anything about it now; and since r is fresh per proof, tags of one user are unlinkable except through the address itself. Decryption needs K = x_svc·c1 + x_AA·c1, one partial product from each share, which is a 2-of-2 threshold key in the sense of Desmedt and Frankel [18] and the basis of authorized opening in Section V-G.')
p = find('The service compares pk_AA.x and pk_AA.y')
replace_in(p, 'a proof under a self-chosen key is valid as an argument and worthless as a statement.',
           'a proof under a self-chosen key is valid as an argument and worthless as a statement. It likewise compares pk_trace.x and pk_trace.y with the key in its own registration file, and rejects a tag whose c1 is the identity point, since r = 0 would leave uid unmasked.')

# ---------------- V-F + 새 V-G ----------------
p = find('The wallet sends the service the proof, its public inputs, and a signature')
replace_in(p, 'and that arid is its own identifier.', 'that arid is its own identifier, and that pk_trace is its combined key.')
replace_in(p, 'and identifies the user by PPID.', 'and identifies the user by PPID. It also appends the full transcript, the public inputs and the proof, to a login log that it keeps private; this log is the material for opening.')
p = find('Requests within a session carry')
h = clone_after(p, 'G. AUTHORIZED OPENING', template=SUB_T)
g1 = clone_after(h, 'Opening a disputed session takes three parties and three steps. The service selects a transcript of the pseudonym in question from its login log, computes its partial decryption D_svc = x_svc·c1, and sends the AA the transcript, D_svc, a timestamp, and an EIP-191 signature under pk_service over (arid, r_s, PPID, D_svc, timestamp). The AA checks, in this order, that the service is registered and approved, that the timestamp is within five minutes, that the signature verifies under the registered pk_service, that the transcript\'s arid and pk_trace are those of the requesting service, that its pk_AA is the AA\'s own key, that D_svc is a valid subgroup point, and finally that the Groth16 argument verifies. It then records the request as pending, without computing anything about the identity.', template=PARA_T)
g2 = clone_after(g1, 'An operator approves or denies the pending request. On approval the AA computes K = D_svc + x_AA·c1 and uid = c2 − Poseidon(K.x, K.y), checks that uid is a registered account, and records the result in a permanent audit log; on denial nothing is decrypted. The service fetches the outcome with a second signed request. A request for a session that is already pending or approved returns the existing record, while a failed or denied one may be requested again, since a wrong partial decryption must be correctable.', template=PARA_T)
clone_after(g2, 'Verifying the transcript before opening is what confines the procedure to the requesting service\'s own sessions: a transcript carries arid and pk_trace as public inputs bound by the proof, so a service cannot open a session of another service even if it obtains that service\'s log, and the AA cannot be used as a decryption oracle for ciphertexts that no valid proof accompanies. The AA stores nothing per login; the tag lives only in the transcript the service keeps, and the AA\'s share is applied only at an approved opening.', template=PARA_T)

# ---------------- VII ----------------
p = find('The AA receives uid, C_pt, chainid, r_s, π_issue, and sig_u.')
replace_in(p, 'This leak is accepted and stated.', 'This leak is accepted and stated. The AA does not see the tag: it travels only from the wallet to the service inside the proof, and the AA\'s share is applied only at an approved opening.')
p = find('cert_s binds arid to an origin under the AA key.')
replace_in(p, 'cert_s binds arid to an origin under the AA key.', 'cert_s binds arid, the origin, and the combined trace key under the AA key, and it is issued only after an operator has approved the registration.')
replace_in(p, 'in the prototype anyone can register.', 'the approval criteria are operational policy, not something the protocol checks.')
set_text(find('H. WHAT AA–SERVICE COLLUSION WOULD YIELD'), 'H. AUTHORIZED OPENING (G10)')
set_text(find('The AA records (uid, r_s) and the service records (PPID, r_s).'),
         'The tag (c1, c2) is a hashed ElGamal ciphertext under pk_trace = (x_svc + x_AA)·B8. Recovering uid requires K = r·pk_trace, and r·pk_trace = x_svc·c1 + x_AA·c1; whoever holds one share can compute one term but not the other, and computing r from c1 is a discrete logarithm. A service alone, an AA alone, or an outsider who obtains a service\'s log therefore learns nothing from the tags. This is what distinguishes the design from opening by joining plaintext records, where a leaked service log lets the AA link identities by itself; here the AA keeps no per-login record at all. Condition 5 of the relation guarantees that what is opened is the identity behind that very statement, so a service cannot be misled by a wallet that encrypts a different value, and the transcript check at the AA ties every opening to the service that verified the login. What the tag does not protect against is the AA and a service combining their shares outside the procedure, which would open every tag of that service without an operator; nor does it protect against an AA that chooses a combined key whose full secret it knows, since the service cannot verify that x_AA is honestly held. Both lie outside the honest-but-curious model. The second can be closed by having the AA prove knowledge of x_AA at registration, and the first by splitting the AA share with a third party, in both cases without changing the tag format. Compared with escrow under a single auditor key, the 2-of-2 arrangement has no key that opens everything by itself; compared with record joining, the opening capability follows the shares rather than the logs.')

# ---------------- VIII ----------------
p = find('The prototype consists of four Node.js processes and one contract.')
replace_in(p, 'The AA exposes registration, service registration, issuance, revocation, publication, an administrator page, and a user account page with self-revocation.',
           'The AA exposes registration, service registration with operator approval, issuance, revocation, publication, the opening endpoints, an administrator page with approval controls, and a user account page with self-revocation.')
replace_in(p, 'The service exposes service information, challenge issuance, login, re-validation, and a session request endpoint,',
           'The service exposes service information, challenge issuance, login, re-validation, a session request endpoint, and an opening request that reads its private login log,')
t3 = d.tables[2]
vals = {'R1CS constraints': '25,505', 'Public inputs / private inputs': '14 / 78', 'Proving key (zkey) / witness generator (wasm)': '15.4 MB / 3.8 MB',
        'Groth16 proving time (min – max)': '865 ms (832 – 1,298 ms)', 'Groth16 verification time': '13.4 ms', 'Proof size (JSON encoding)': '722 B'}
for r in t3.rows:
    k = r.cells[0].text
    if k in vals: set_cell(r.cells[1]._tc, vals[k])
set_text(find('The earlier five-slot version of the circuit'),
         'Two earlier versions of the circuit give the cost of each addition. The five-slot commitment without attribute slots had 18,786 constraints; adding four attribute slots and making r_s public brought it to 22,224; adding the trace tag, one fixed-base and one variable-base scalar multiplication plus one Poseidon call, brought it to 25,505, about 15 percent more, with proving time within the run-to-run variation of the previous version. The Pedersen commitment remains the largest single block; a Poseidon commitment would be cheaper in the circuit but would force the issuance proof from a sigma protocol into a second SNARK, which is why Pedersen was kept.')
replace_in(find('Per-login cost is one issuance round trip'), 'one Groth16 proof of about 0.8 s', 'one Groth16 proof of about 0.9 s')
set_text(find('An end-to-end suite runs the AA, wallet agent, service, and chain in isolation'),
         'An end-to-end suite runs the AA, wallet agent, service, and chain in isolation and exercises the scenarios of Sections V and VI: service registration held pending until an operator approves it; user registration; login with issuance and session creation; re-validation under an unchanged root, which reuses the cached proof and yields the identical session signature; account revocation and publication; a re-validation with a stale proof, refused with stale_root, and a session request refused with revalidate_required; a synchronized re-validation refused because the statement is revoked, and a new login refused because the account is disabled; administrator restoration followed by re-issuance and login under the same address; a second login with the same r_s refused as a used challenge; a wallet request whose origin or trace key differs from the certificate refused; a challenge past its lifetime refused; and an opening in which the service requests attribution of a pseudonym, the pending record carries no identifier, the operator approves, and the service receives the account behind the session. A separate suite covers the opening checks one by one: a request signed by another service, a transcript of another service, a tag under a foreign combined key, a stale timestamp, an unapproved service, a wrong partial decryption that fails at approval and can be retried, and a denied request. The circuit suite checks the relation against forged inputs, among them a pseudonym computed for another chain, a commitment that does not open to the public arid, a revoked leaf, and a tag whose ciphertext, key, or randomness has been altered.')

# ---------------- IX ----------------
set_text(find('The AA sees the chain identifier and the challenge.'),
         'The AA sees the chain identifier. Where its allow list has several entries it learns which chain the user is using; this is the one value outside the commitment that the AA must check itself.')
set_text(find('Service registration is unauthenticated.'),
         'Service registration is approved by an operator, but the criteria for approval are operational policy; the certificate attests that the origin was approved and that the AA holds a share of its trace key, not that the service is trustworthy. Opening rests on the same operator: the AA and a service that pool their shares outside the procedure, or an AA that constructs a combined key it fully knows, can open tags without approval. The first requires splitting the AA share with a third party, the second requires the AA to prove knowledge of its share at registration; the tag format accommodates both. Loss of the AA\'s share for a service seals that service\'s past tags permanently, so the AA state file is also the opening capability, and a service that discards its own share loses the same. Self-revocation has no rate limit on password attempts.')

# ---------------- X ----------------
p = find('This paper separated two things that Web3 services conflate')
replace_in(p, 'and lets a user cut off a lost device from any browser.',
           'and lets a user cut off a lost device from any browser. Encrypting the identity under a key split between the authority and the service, and proving in the same argument that the ciphertext holds the identity behind the statement, gives disputed sessions an attribution path that needs both parties and an operator, without giving any of them the power to trace on their own.')
p = find('The prototype shows that the whole flow fits in one Groth16 argument')
replace_in(p, 'about twenty-two thousand constraints', 'about twenty-five thousand constraints')
replace_in(p, 'to decide whether conditional traceability through the challenge should be formalized as an audited opening or closed off,',
           'to let the AA prove knowledge of its trace share at registration and to split that share with a third party so that opening survives an authority that colludes with a service,')

# ---------------- 참고문헌 ----------------
last_ref = REF_T
r18 = clone_after(last_ref, '[18] Y. Desmedt and Y. Frankel, "Threshold cryptosystems," in Advances in Cryptology - CRYPTO \'89, LNCS, vol. 435, Santa Barbara, CA, USA, 1989, pp. 307-315, doi: 10.1007/0-387-34805-0_28.')
clone_after(r18, '[19] T. ElGamal, "A public key cryptosystem and a signature scheme based on discrete logarithms," IEEE Trans. Inf. Theory, vol. 31, no. 4, pp. 469-472, Jul. 1985, doi: 10.1109/TIT.1985.1057074.')

d.save(DST)
print('saved', DST)
