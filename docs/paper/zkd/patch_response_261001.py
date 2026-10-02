#!/usr/bin/env python3
"""documents/260930_meeting_MHJ(with feedback).pptx → documents/261001_meeting_MHJ_response.pptx (2026-10-01).

9/30 회의 피드백(파란 상자 6개)에 대한 대응 덱. 교수님 용어(IdP·RP·Token·auid_i·salt·π_IdP/π_RP/π_i)는 유지하고,
피드백 상자는 각 장 아래 띠로 줄여 남긴다(원문 그대로). 대응 내용은 BLUE. 2장은 설계 메모
docs/superpowers/specs/2026-10-01-mode3-credential-registry-memo.md §7 을 따른다.
새 3장(auid_i 의 원래 역할 → 대체)을 끼워 넣어 전체 12장이 된다.
실행: python3 docs/paper/zkd/patch_response_261001.py
"""
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC = 'documents/260930_meeting_MHJ(with feedback).pptx'
OUT = 'documents/261001_meeting_MHJ_response.pptx'
BLUE = '1F4EB0'
GREY_FILL = 'F2F2F2'
A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
MC = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
FILLS = (A + 'noFill', A + 'solidFill', A + 'gradFill', A + 'blipFill', A + 'pattFill', A + 'grpFill')

prs = Presentation(SRC)
S = list(prs.slides)
assert len(S) == 11, len(S)

# ---------- 조회 ----------
def ptext(pel):
    return ''.join(t.text or '' for t in pel.iter() if t.tag in (A + 't', M + 't'))
def sptree(slide):
    return slide._element.find('.//' + P + 'spTree')
def name_of(el):
    return el.find('.//' + P + 'cNvPr').get('name')
def plain_sp(slide, name):
    hits = [el for el in sptree(slide) if el.tag == P + 'sp' and name_of(el) == name]
    assert len(hits) == 1, (name, len(hits)); return hits[0]
def alt(slide, name):
    hits = [el for el in sptree(slide) if el.tag == MC + 'AlternateContent' and name_of(el.find(MC + 'Choice/' + P + 'sp')) == name]
    assert len(hits) == 1, (name, len(hits)); return hits[0]
def de_alt(slide, name):
    ac = alt(slide, name); sp = copy.deepcopy(ac.find(MC + 'Choice/' + P + 'sp'))
    ac.addprevious(sp); ac.getparent().remove(ac); return sp
def paras(sp):
    return sp.find(P + 'txBody').findall(A + 'p')

# ---------- 글 ----------
def style(rpr, color=None, size=None, bold=None):
    rpr.attrib.pop('strike', None); rpr.attrib.pop('sym', None)
    for c in list(rpr):
        if c.tag == A + 'sym': rpr.remove(c)
    if size is not None: rpr.set('sz', str(int(size * 100)))
    if bold is not None: rpr.set('b', '1' if bold else '0')
    if color is not None:
        for c in list(rpr):
            if c.tag in FILLS: rpr.remove(c)
        sf = etree.Element(A + 'solidFill'); etree.SubElement(sf, A + 'srgbClr').set('val', color)
        rpr.insert(1 if len(rpr) and rpr[0].tag == A + 'ln' else 0, sf)
    return rpr
def run_template(pel):
    r = pel.find(A + 'r')
    if r is not None and r.find(A + 'rPr') is not None: t = copy.deepcopy(r.find(A + 'rPr'))
    else:
        e = pel.find(A + 'endParaRPr'); t = copy.deepcopy(e) if e is not None else etree.Element(A + 'rPr'); t.tag = A + 'rPr'
    return style(t)
def segs_of(x):
    if isinstance(x, str): return [(x, None)]
    if isinstance(x, tuple): return [x]
    return list(x)
def set_para(pel, content, color=None, size=None, bold=None, lvl=None):
    tmpl = run_template(pel)
    for c in list(pel):
        if c.tag != A + 'pPr': pel.remove(c)
    if lvl is not None:
        pp = pel.find(A + 'pPr')
        if pp is None: pp = etree.Element(A + 'pPr'); pel.insert(0, pp)
        if lvl: pp.set('lvl', str(lvl))
        else: pp.attrib.pop('lvl', None)
    for text, col in segs_of(content):
        r = etree.SubElement(pel, A + 'r'); r.append(style(copy.deepcopy(tmpl), col if col is not None else color, size, bold))
        etree.SubElement(r, A + 't').text = text
    return pel
def rebuild(sp, lines):
    """txBody 를 새로 쓴다. lines = [dict(content, lvl, color, size, bold)]; pPr 은 원본의 같은 lvl 것을 복사."""
    tx = sp.find(P + 'txBody'); old = tx.findall(A + 'p'); ppr = {}
    for pel in old:
        pp = pel.find(A + 'pPr'); ppr.setdefault(int(pp.get('lvl', '0')) if pp is not None else 0, pp)
    tmpl = run_template(old[0])
    for pel in old: tx.remove(pel)
    for ln in lines:
        lvl = ln.get('lvl', 0); pel = etree.SubElement(tx, A + 'p'); src = ppr.get(lvl, ppr.get(0))
        if src is not None:
            pp = copy.deepcopy(src)
            if lvl: pp.set('lvl', str(lvl))
            else: pp.attrib.pop('lvl', None)
            pel.append(pp)
        elif lvl: etree.SubElement(pel, A + 'pPr').set('lvl', str(lvl))
        for text, col in segs_of(ln['content']):
            r = etree.SubElement(pel, A + 'r'); r.append(style(copy.deepcopy(tmpl), col if col is not None else ln.get('color'), ln.get('size'), ln.get('bold')))
            etree.SubElement(r, A + 't').text = text
    return sp
def place(sp, left=None, top=None, width=None, height=None):
    xfrm = sp.find(P + 'spPr/' + A + 'xfrm'); off, ext = xfrm.find(A + 'off'), xfrm.find(A + 'ext')
    if left is not None: off.set('x', str(int(Inches(left))))
    if top is not None: off.set('y', str(int(Inches(top))))
    if width is not None: ext.set('cx', str(int(Inches(width))))
    if height is not None: ext.set('cy', str(int(Inches(height))))
def textbox(slide, left, top, width, height, lines, size=12, color=BLUE, fill=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)); tf = tb.text_frame; tf.word_wrap = True
    if fill: tb.fill.solid(); tb.fill.fore_color.rgb = RGBColor.from_string(fill)
    for i, ln in enumerate(lines):
        ln = (ln,) if isinstance(ln, str) else tuple(ln)
        text, sz, col = ln[0], (ln[1] if len(ln) > 1 else size), (ln[2] if len(ln) > 2 else color)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); r = p.add_run(); r.text = text; r.font.size = Pt(sz)
        if col: r.font.color.rgb = RGBColor.from_string(col)
    return tb
def strip(slide, name, left, top, width, height, size=10.5):
    """교수님 피드백 상자 → 원문 그대로 한 문단으로 합쳐 띠로 줄인다(채움·흰 글자 유지)."""
    sp = plain_sp(slide, name); ps = paras(sp)
    text = ' / '.join(t for t in (ptext(q).strip() for q in ps) if t)
    for q in ps[1:]: q.getparent().remove(q)
    set_para(ps[0], '지난 회의 피드백 — ' + text, size=size)
    sp.find(P + 'txBody/' + A + 'bodyPr').set('anchor', 'ctr')
    place(sp, left, top, width, height)
    return sp

BODY_TMPL = copy.deepcopy(plain_sp(S[7], '내용 개체 틀 2'))   # 8장 본문(평문, lvl0·lvl1) — AC 본문을 갈아끼울 틀
def fresh_body(slide, name='내용 개체 틀 2'):
    ac = alt(slide, name); body = copy.deepcopy(BODY_TMPL); ac.addprevious(body); ac.getparent().remove(ac); return body
L0, L1 = 16, 13

# ---------- 2장 Revocation ----------
s = S[1]
rebuild(fresh_body(s), [
    dict(content='User attribute 이외의 session attribute를 위한 revocation', size=L0),
    dict(content='현재: 같은 Merkle tree, 리프 종류 둘(사용자 H(4, H(C_u)), 세션 H(5, H(C_s))); π_RP 가 비멤버십 2건', lvl=1, size=L1),
    dict(content='VCL 이 왜 필요한가 — 책임성', color=BLUE, size=L0),
    dict(content='AA 서명은 비공개라 사용자가 발급을 볼 수 없다. 쓸 수 있는 자격증명이 공개 구조에 있어야 "내 uid 로 발급된 것이 내 것뿐인가"를 탐지할 수 있다 (AA 는 sk_u·비밀번호를 쥐므로 내 uid 로 C′ 을 만들 수 있음 → Sybil·모함)', lvl=1, color=BLUE, size=L1),
    dict(content='뺐던 이유도 유효: 검증에는 서명 + 만료로 충분(발급 때 쓰기 0). 집합 VCL 은 탐지에 무력(어느 항목이 내 것인지 모름) → 사용자별 슬롯이어야 함. 세션은 제외(AA 가 C_u 위에 세션을 위조 못 함)', lvl=1, color=BLUE, size=L1),
    dict(content='RCL 두 종류 = 등록부 + 폐기 트리 (제안 O2)', color=BLUE, size=L0),
    dict(content='사용자 자격증명 → 등록부: 키 = 등록 커밋먼트 cm_u, 값 = 활성 H(C_u) 한 칸 (멤버십, 발급·교체 때 갱신)', lvl=1, color=BLUE, size=L1),
    dict(content='세션 → 폐기 트리: append-only 비멤버십, 현행 그대로. 사용자 리프 종류는 등록부로 흡수', lvl=1, color=BLUE, size=L1),
    dict(content='latest active = 슬롯 자체', color=BLUE, size=L0),
    dict(content='π_RP 가 슬롯(cm_u) = H(C_u) 멤버십과 C_u 의 salt = cm_u 의 salt 등식을 증명 → AA 가 C′ 을 쓰려면 내 칸을 바꿔야 하고 그러면 내 자격증명이 죽어 바로 드러남. 폐기해서 확인할 필요 없음', lvl=1, color=BLUE, size=L1),
    dict(content='한계: 내 uid 의 두 번째 등록은 못 잡음(키가 uid 와 무관) — VRF 키는 과제. 비용: 제약 +2천 안팎, 공개 입력 +1(등록부 root), 쓰기는 자격증명 발급 때만', lvl=1, color=BLUE, size=L1),
])
strip(s, '직사각형 5', 0.37, 6.3, 12.6, 0.6)

# ---------- 새 3장: auid_i 의 원래 역할 → 대체 (3장 앞에 끼운다) ----------
layout = next(l for l in prs.slide_layouts if l.name == '제목 및 내용')
ns = prs.slides.add_slide(layout)
ns.shapes.title.text = 'auid_i — 원래 설계에서의 역할, 그리고 무엇으로 대체했나'
for ph in list(ns.placeholders):
    if ph.placeholder_format.idx != 0: ph._element.getparent().remove(ph._element)
rows = [
    ['auid_i 가 하던 일 (원래 설계)', '어떻게', '지금 맡는 것', '어떻게'],
    ['토큰의 주체를 uid 없이 표시', 'Token = Sign(auid_i, …), auid_i = H(auid, i), auid = H(uid, salt)', 'H(C_u)', '토큰이 H(C_u) 를 서명; C_u = Commit(uid, salt, a) — RP 는 못 보고 π_RP 가 열어 PPID 를 계산'],
    ['이 토큰이 이 세션 것', 'i = 세션 번호', 'H(C_s) 안의 pk_i', 'C_s = Commit(arid, pk_i; blind), 세션마다 새 키; sk_i 로 RP nonce 서명'],
    ['IdP 가 세션끼리·RP 를 못 잇게', '세션마다 다른 auid_i', 'H(C_s) 가 매번 새 값', 'blind·pk_i 가 새로; arid 는 C_s 안이라 IdP 는 못 봄'],
    ['추적 손잡이', 'RP 가 auid_i → IdP → uid', 'Trace Tag', 'Enc(pk_trace, H(uid, arid)); RP 조각 + IdP 조각 + 운영자 승인'],
    ['세션 폐기 식별자', 'IdP 가 auid_i 로 세션 지정', 'IdP 세션 기록의 H(C_s)', '발급 때 기록, 폐기 때 리프 H(5, H(C_s))'],
]
tbl = ns.shapes.add_table(len(rows), 4, Inches(0.5), Inches(1.6), Inches(12.3), Inches(3.6)).table
for w, col in zip((2.4, 3.2, 2.2, 4.5), tbl.columns): col.width = Inches(w)
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j); cell.text = val
        for q in cell.text_frame.paragraphs:
            for r in q.runs:
                r.font.size = Pt(11); r.font.bold = (i == 0)
                if i and j >= 2: r.font.color.rgb = RGBColor.from_string(BLUE)
textbox(ns, 0.5, 5.4, 12.3, 1.3, [
    ('요지: auid_i 하나가 겸하던 다섯 역할을 H(C_u) · H(C_s)/pk_i · Trace Tag · IdP 기록 넷으로 나눴다. 연결은 π_RP 안에서만 성립한다.', 13),
    ('빼서 좋아진 것: IdP 와 RP 가 같은 값을 보지 않는다(옛 auid_i 는 둘 다 봐서 맞대면 uid ↔ PPID 가 바로 이어짐) → 공모 연결은 설계된 개봉 절차로만. 세션 토큰마다 IdP 쪽 ZKP 가 없어진다(C_u 발급 때 1회).', 12),
    ('H(C_u) 로의 "대체"는 주체 결속 한 역할뿐이고, 세션별로 바뀌어야 하는 역할은 H(C_s) 가 맡는다 — 그래서 H(C_u) 는 고정이어도 된다(한쪽에만 보이는 값).', 12),
])
lst = prs.slides._sldIdLst; el = lst[-1]; lst.remove(el); lst.insert(2, el)

# ---------- (옛 3장) ID Transformation — 지금 설계 ----------
s = S[2]
tok = de_alt(s, 'Rounded Rectangle 4')
rebuild(tok, [
    dict(content='(* Token 관련 정보 *)', size=14),
    dict(content='Token = Sign_IdP( H(C_u), H(C_s), expire time, chainid, allowAgent )', lvl=1, color=BLUE, size=12.5),
    dict(content='auid_i 자리는 H(C_s)(세션마다 새 blind·새 pk_i); 주체 결속은 H(C_u). RP 가 보는 고정 이름 PPID 는 지갑이 계산, π_RP 로 증명', lvl=1, color=BLUE, size=12.5),
    dict(content=[('(* User attributes *)', None), (' → C_u', BLUE)], size=14),
    dict(content='uid, salt', lvl=1, size=12.5),
    dict(content=[('other attributes (e.g., age)', None), (' → 슬롯 4개(출생연도·국가·등급·예비) — 6슬롯 + 매핑으로 바꿀 예정(6장)', BLUE)], lvl=1, size=12.5),
    dict(content=[('(* Session attributes *)', None), (' → C_s', BLUE)], size=14),
    dict(content=[('arid, pk_i', None), (' (session key; blind 로 은닉 — IdP 는 못 봄)', BLUE)], lvl=1, size=12.5),
    dict(content=[('expire time (block height)', None), (' — C_s 밖, 서명 본문에', BLUE)], lvl=1, size=12.5),
])
idp = de_alt(s, 'Rounded Rectangle 7')
rebuild(idp, [
    dict(content='IdP', size=18),
    dict(content='uid (로그인으로 확인), H(C_u), H(C_s) 만 봄 — auid 없음', lvl=1, color=BLUE, size=12.5),
    dict(content='기록: (uid, cm_u, H(C_u), 속성) + 세션 (H(C_s), expire, chainid)', lvl=1, color=BLUE, size=12.5),
    dict(content='제안 O2: 활성 H(C_u) 를 등록부 슬롯(cm_u)에 게시', lvl=1, color=BLUE, size=12.5),
])
rp = de_alt(s, 'Rounded Rectangle 9')
rebuild(rp, [
    dict(content='RP', size=18),
    dict(content='사용자 구분은 PPID = H(uid, salt, chainid, arid) — H(C_u)·auid 는 못 봄', lvl=1, color=BLUE, size=12.5),
    dict(content=[('rid → arid', None), (' : IdP 가 등록 승인 때 뽑는 난수(cert_s 로 origin 에 결속)', BLUE)], lvl=1, size=12.5),
    dict(content='세션마다 새 pk_i 와 Trace Tag 를 봄', lvl=1, color=BLUE, size=12.5),
])
usr = de_alt(s, 'Rounded Rectangle 10')
rebuild(usr, [
    dict(content='User (지갑)', size=18),
    dict(content='uid, salt (salt 는 지갑만 앎 → IdP 도 PPID 를 못 만듦)', lvl=1, color=BLUE, size=12.5),
    dict(content='PPID = H(uid, salt, chainid, arid) — RP 별 고정', lvl=1, color=BLUE, size=12.5),
    dict(content='세션마다 C_s = Commit(arid, pk_i; blind), pk_i 새로', lvl=1, color=BLUE, size=12.5),
])
plain_sp(s, '제목 1').find('.//' + A + 't').text = 'ID Transformation — 지금 설계'
strip(s, '직사각형 8', 0.37, 6.52, 6.6, 0.45, size=10)

# ---------- (옛 4장) π_IdP ----------
s = S[3]
rebuild(fresh_body(s), [
    dict(content=[('사용자 → IdP 에게 ', None), ('직접 제출 (RP FE 경유 없음 → IdP 는 어느 RP 인지 모름); C_u 발급 시 1회 — 세션 토큰마다는 ZKP 없음', BLUE)], size=15),
    dict(content='증명 내용', size=15),
    dict(content=[('(1) uid, auid 관계 증명', None), (' → C_u = Commit(uid, salt, a; blind) 의 열림: uid 는 로그인 계정, a 는 IdP 기록, salt 는 등록 때 낸 cm_u 의 salt 와 같음 (시그마 프로토콜, 회로 없음, 80 ms)', BLUE)], lvl=1, size=14),
    dict(content=[('(2) auid, C_u 관계 증명', None), (' → 불필요 — 맞음. auid 자체가 없고, (1) 이 uid·salt·a 결속을 한 번에 함', BLUE)], lvl=1, size=14),
])
set_para(paras(plain_sp(s, 'Rounded Rectangle 38'))[0], '(uid, cm_u, H(C_u), 속성) 기록', color=BLUE, size=14)
set_para(paras(plain_sp(s, 'Rounded Rectangle 50'))[0], '세션 토큰 발행 — ZKP 없이 sk_u 서명만', color=BLUE, size=13)
strip(s, '직사각형 11', 0.37, 6.2, 7.9, 0.5)

# ---------- (옛 5장) π_RP ----------
s = S[4]
rebuild(fresh_body(s), [
    dict(content=[('사용자 → RP 에게 제출', None), (' — 로그인마다 π_RP 1건 (Groth16, 제약 37,130; 증명 1.19 s·검증 10 ms; 공개 입력 25개)', BLUE)], size=15),
    dict(content='증명 내용', size=15),
    dict(content=[('(1) Valid Token (IdP 서명 && not expired)', None), (' — IdP 서명(EdDSA)은 회로 안에서; expire time 은 공개 입력, 만료 비교는 RP·컨트랙트가 현재 블록 높이로', BLUE)], lvl=1, size=14),
    dict(content=[('(2) Valid Token 내 IDs (auid, C_u, arid) 정보 정확성 증명', None), (' → C_u = Commit(uid, salt, a; blind_u), C_s = Commit(arid, pk_i; blind_s) 의 열림 증명 — auid 는 없음; arid·pk_i 는 공개 입력', BLUE)], lvl=1, size=14),
    dict(content=[('(3) Valid Token 내 IDs 기반하여 생성된 PPID 검증', None), (' — PPID = H(uid, salt, chainid, arid) 를 (2) 와 같은 uid·salt 로 계산, 공개 입력과 일치', BLUE)], lvl=1, size=14),
    dict(content=[('(4) (if any) Vector Commitment 기반 사용자 속성값 Selective Disclosure 증명', None), (' — 슬롯별 범위(lo ≤ a ≤ hi)·집합 소속 조건을 마스크로 선택; 값은 비공개', BLUE)], lvl=1, size=14),
    dict(content="(5) 폐기 확인 — 현재: 사용자 리프·세션 리프 비멤버십 2건 (리프는 IdP 발급 기록과 바로 이어져 공개 불가 → 증명 안에서만). 제안 O2 뒤: 세션 비멤버십 + 등록부 멤버십(슬롯(cm_u) = H(C_u)) + salt 등식", lvl=1, color=BLUE, size=14),
    dict(content='(6) Trace Tag = Enc(pk_trace, H(uid, arid)) 가 같은 uid·arid 로 만들어졌음', lvl=1, color=BLUE, size=14),
])

# ---------- (옛 6장) 속성 ----------
s = S[5]
set_para(paras(plain_sp(s, '제목 1'))[0], [('(4)', None), ('의 데모 화면 — 속성 슬롯과 매핑 (대응)', BLUE)])
pic = next(sh for sh in s.shapes if sh.shape_id == 5); pic.height = Inches(3.9); pic.width = Inches(9.36 * 3.9 / 4.6)
textbox(s, 8.75, 2.1, 4.3, 4.0, [
    ('현재', 13),
    '슬롯 4개 정수(출생연도·국가·등급·예비), 관리자가 수정 가능 → 활성 자격증명 교체, 지갑이 다시 받음. 화면은 정수 그대로',
    ('대응', 13),
    '(i) 범용 슬롯 6개 attribute1~6',
    '(ii) 매핑 계층: 속성 이름·인코딩 표(국가 KR ↔ 410, 등급 gold ↔ 2, 생년 ↔ 정수) — IdP 가 정의, 지갑·RP 화면은 이름과 값으로',
    '(iii) 정수는 회로 안에만. 조건 입력도 "나이 ≥ 19", "국가 ∈ {KR, JP}" 처럼',
    ('비용', 13),
    '회로: 제너레이터 2 + 범위 술어 2 → 공개 입력 25 → 29, zkey·검증자 재생성, 인덱스 전면 수정, 벤치 재측정. 등록부(O2)와 같은 V9 에서',
], size=11.5)
strip(s, '직사각형 5', 0.37, 6.3, 12.6, 0.55)

# ---------- (옛 7장) π_i ----------
s = S[6]
rebuild(fresh_body(s), [
    dict(content=[('사용자 혹은 RP → 트랜잭션 생성시 첨부 → 수신자에게 전달 (TX + π_i + Trace Tag', None), (' + σ_tx)', BLUE)], size=15),
    dict(content='π_i 는 별도 증명이 아님 — 로그인 때의 π_RP(같은 회로·공개 입력)를 세션 안에서 재사용; 트랜잭션마다 새 ZKP 없음, 검증은 컨트랙트가 매번', color=BLUE, size=15),
    dict(content='증명 내용', size=15),
    dict(content=[('(1) π_RP 증명 내용 (Valid Token && IDs 검증 && PPID 검증', None), (' && 폐기 확인 && 선택 공개)', BLUE)], lvl=1, size=14),
    dict(content=[('(2) session public key pk_i 증명', None), (' — pk_i 는 π_RP 공개 입력(C_s 안의 값); 트랜잭션은 sk_i 서명 σ_tx 로 결속 (TX 본문 11워드 + expire + allowAgent + Trace Tag = 16워드)', BLUE)], lvl=1, size=14),
    dict(content=[('(3) Trace Tag 가 Valid Token 내 ', None), ('uid·arid', BLUE), (' 기반하여 생성되었음', None), (' — 서명이 아니라 암호화: Enc(pk_trace, H(uid, arid)); 회로 안에서 같은 uid·arid 임을 검증', BLUE)], lvl=1, size=14),
])
ac = alt(s, 'TextBox 7'); ac.getparent().remove(ac)
textbox(s, 1.0, 4.9, 11.4, 1.3, [
    ('Trace Tag = Enc(pk_trace, H(uid, arid)),   pk_trace = pk_RP + pk_IdP   (Sign_IdP(auid) 아님)', 15),
    ('열기: RP 조각 + IdP 조각 + 운영자 승인 → H(uid, arid) → IdP 기록에서 uid — 2-of-2 가산 복호, threshold 서명 없음. 개봉 범위는 "이 RP 에서의 이 사용자"까지', 12),
    ('컨트랙트 검사 순서: nonce → σ_tx → 성명(계정·IdP 키·allowAgent·Trace Tag 결속·root·root 나이·만료) → Groth16 — 실행 gas 약 417k–456k', 12),
], fill=GREY_FILL)

# ---------- (옛 8장) Address Abstraction ----------
s = S[7]
rebuild(plain_sp(s, '내용 개체 틀 2'), [
    dict(content='주소 추상화의 특성 넷 — 셋은 이미 목표에 있고 chain-agnostic 만 새로 정의', size=17),
    dict(content='Uniqueness = zkAA Injectiveness → G1: PPID = H(uid, salt, chainid, arid), 같은 문맥에서 사용자가 다르면 주소도 다름 (정리 1)', lvl=1, color=BLUE, size=13.5),
    dict(content='Immutability = zkAA Tamper resistance → G2: 계정 = CREATE2(factory, PPID); 지갑·IdP·RP 누구도 한 계정의 주소를 바꾸지 못함 (정리 2)', lvl=1, color=BLUE, size=13.5),
    dict(content='Unlinkability = zkAA Privacy-preservation → G3: 주소에서 uid 를 얻지 못하고 RP·체인 간 주소를 잇지 못함 (정리 3)', lvl=1, color=BLUE, size=13.5),
    dict(content='Chain-agnostic consistency → G11(신설): 어느 문맥 (chainid, arid) 에서든 계정의 주소는 등록 때 커밋한 신원 (uid, salt) 하나에서 파생한 값이고, 받아들여진 어떤 증명도 두 번째 신원으로 열리지 않음', lvl=1, color=BLUE, size=13.5),
    dict(content='근거: π_IdP 가 salt 를 등록 cm_u 에 결속(Pedersen 결합성) + 등록은 uid 당 1회 → 정리 2 의 증명에서 "같은 문맥" 전제만 풀면 그대로 (상한 동일). 뜻: 등록 한 번이 모든 체인·RP 의 주소를 정하되, 주소는 문맥마다 다르고 서로 이어지지 않음(G3)', lvl=1, color=BLUE, size=13.5),
    dict(content='나머지 zkAA 특성도 그대로: Unforgeability → G6·G7(IdP 서명 + σ_tx), Correctness → 완전성, Chronicle → G10 + 온체인 개봉 기록', lvl=1, color=BLUE, size=13.5),
    dict(content='Address abstraction 자체는 특성이 아니라 정의: 주소 함수 Addr(id, ctx) = H(uid, salt, chainid, arid) 와 그것을 검증하는 성명(π_RP 의 PPID 조건). 제안 O2 가 들어가면 G12 "발급 책임성" 추가', lvl=1, color=BLUE, size=13.5),
])
strip(s, '직사각형 4', 0.37, 6.3, 12.6, 0.5)

# ---------- (옛 9장) Todo ----------
s = S[8]
rebuild(plain_sp(s, '내용 개체 틀 2'), [
    dict(content='프로토타입 구성 및 평가', size=17),
    dict(content='데모 시나리오', lvl=1, size=15),
    dict(content=[('salt 변경 → ', None), ('자격증명 발급 거절', BLUE), (' 시나리오 (세션 인증이 아니라 발급 단계: π_IdP 가 salt 를 cm_u 에 결속; G2·G11 시연)', BLUE)], lvl=2, size=13.5),
    dict(content=[('서로 다른 RP 들에 대해 세션 인증 시나리오', None), (' — 지갑 origin 목록·RP #2 구현 완료(9/30); 같은 사용자가 RP 마다 다른 PPID·다른 주소', BLUE)], lvl=2, size=13.5),
    dict(content=[('가상자산거래소가 RP 인 시나리오', None), (' — RP #2 를 거래소로: 나이·국가 조건 공개, Trace Tag 개봉', BLUE)], lvl=2, size=13.5),
    dict(content='평가: V9(등록부 + 속성 6슬롯) 뒤 재측정 — 제약·증명 시간·실행 gas·게시 gas', lvl=1, color=BLUE, size=14),
    dict(content='이번 회의에서 정할 것', color=BLUE, size=17),
    dict(content='O2 등록부 채택 (키 cm_u → 활성 H(C_u); π_RP 멤버십 + salt 등식) — 세션은 폐기 트리 유지', lvl=1, color=BLUE, size=14),
    dict(content='등록 키를 지갑이 생성(pk_u 만 제출) — 지금은 IdP 가 sk_u 를 만들어 줘 "내 요청 없이 바꿨다"는 증거력이 약함', lvl=1, color=BLUE, size=14),
    dict(content='속성 6슬롯 + 매핑 계층 (6장)', lvl=1, color=BLUE, size=14),
    dict(content='한계로 적을 것: 내 uid 의 두 번째 등록(VRF 키 과제)', lvl=1, color=BLUE, size=14),
    dict(content='논문: §V 구조(등록부 + 폐기 트리), G11·G12 추가, (A10) 완화, 특성 표', lvl=1, color=BLUE, size=14),
])
strip(s, '직사각형 4', 0.37, 6.35, 12.6, 0.5)

# ---------- (옛 11장) Overall Flow ----------
plain_sp(S[10], '제목 1').find('.//' + A + 't').text = 'Overall Flow — 논문 Fig. 1 (등록부 O2 반영 시 수정)'

prs.save(OUT)
out = Presentation(OUT)
left = [(i + 1, name_of(el.find(MC + 'Choice/' + P + 'sp'))) for i, s in enumerate(out.slides) for el in sptree(s) if el.tag == MC + 'AlternateContent']
print('saved', OUT, '| slides', len(out.slides), '| 남은 수식 도형:', left)
