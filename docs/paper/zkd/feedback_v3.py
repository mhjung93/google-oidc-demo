# v3 셀프 피드백 라운드(2026-09-18): 모순·잔존 표현 정정. 제자리 수정(백업 .bak-0918).
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
P = 'documents/2026-09-18_v3_zk-Delegation_research_article.docx'
d = Document(P)
def find(prefix):
    for p in d.paragraphs:
        if p.text.startswith(prefix): return p
    raise KeyError(prefix)
def set_text(p, text):
    rs = p.runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
def rep(p, old, new):
    assert old in p.text, (old[:60], p.text[:80]); set_text(p, p.text.replace(old, new, 1))
def set_cell(tc, text):
    ps_ = tc.findall(qn('w:p'))
    for q in ps_[1:]: tc.remove(q)
    q = ps_[0]; rs = q.findall(qn('w:r'))
    for r in rs[1:]: q.remove(r)
    ts = rs[0].findall(qn('w:t'))
    for t in ts[1:]: rs[0].remove(t)
    ts[0].text = text; ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')

# F1 서론: "and nothing more" 가 다음 문장과 모순
rep(find('We take a different view of what an address is.'),
    'In this paper an address is the name under which a service identifies a user, and nothing more. It need not be',
    'In this paper an address is first the name under which a service identifies a user. It need not be')
# F2 서론: 챌린지는 AA 가 아예 보지 않는다
p = find('zk-Delegation combines address abstraction with delegated authentication')
rep(p, 'and a fresh random challenge, which is the only thing the AA will ever see of the service.',
    'and a fresh random challenge, which never reaches the AA.')
rep(p, 'a session public key, and any user-managed attributes, together with the challenge.',
    'a session public key, and any user-managed attributes, together with a per-session agent-permission flag.')
# F27 zkLogin 대비
rep(find('The second question is how such a name is attached'),
    'but the address remains an on-chain account and the identity provider still sees the client.',
    'but the address is an on-chain account tied to one chain, and the identity provider still sees the client.')
# F4 II-B zkAA 비교: 온체인에서는 우리도 트랜잭션마다 증명을 검증한다
p = find('OIDC is an identity layer on OAuth 2.0')
rep(p, 'the price is a proof per transaction. zk-Delegation introduces sessions so that one proof is reused, and therefore must handle revocation during a session,',
    'the price is a fresh statement, with its issuance, per transaction. zk-Delegation introduces sessions so that one statement is reused, off chain with one proof and on chain with the same proof re-verified per transaction but never re-issued, and therefore must handle revocation during a session,')
# F5 G7
set_text(find('G7 Session binding.'),
         'G7 Session binding. A login is accepted by a service only once, for the challenge that service issued and only under the session key the statement covers; later requests in the session are bound to that session key, and transactions on chain to that key and the account\'s nonce.')
# F7 서비스 등록: 팩토리 배포
rep(find('A service registers by sending its display name'),
    'until approval it runs in a waiting state and refuses logins.',
    'until approval it runs in a waiting state and refuses logins. A service that will act on a chain then deploys its account factory (Section V-I) and keeps its address next to the certificate.')
# F9 표 1 태그 행
for r in d.tables[0].rows:
    if r.cells[0].text.strip() == 'tag = (c1, c2)':
        set_cell(r.cells[1]._tc, 'Poseidon(uid, arid) encrypted under pk_trace (hashed ElGamal)')
# F14 gas 배수
rep(find('Per-login cost is one issuance round trip'), 'roughly fifteen times a plain transfer', 'roughly eighteen times a plain transfer')
# F15 하트비트는 자동
rep(find('The newly inserted leaves travel as calldata'),
    'The prototype publishes on an administrator\'s command, and a publication with no leaves is allowed: it re-signs the unchanged root under a new epoch and serves as the heartbeat that the account contracts of Section V-I require.',
    'The prototype publishes revocations on an administrator\'s command. A publication with no leaves is also allowed: it re-signs the unchanged root under a new epoch, and the AA sends one automatically every 50 blocks as the heartbeat that the account contracts of Section V-I require.')
# F16 VI-C: 온체인에서는 가스가 든다
rep(find('Services read the root and the block head directly from the chain'),
    'costs the wallet less than a second to regenerate and no gas,', 'costs the wallet less than a second to regenerate and, off chain, no gas,')
# F19 VIII-A 서명 형식
rep(find('The prototype consists of four Node.js processes and four contracts'),
    'session signatures are secp256k1 ECDSA in the EIP-191 message format through ethers 6.',
    'session signatures are secp256k1 ECDSA in the EIP-191 message format through ethers 6, and transaction signatures are over the raw payload digest so that the account contract recovers them with ecrecover.')
# F20 측정 횟수 표기 통일
rep(find('Table 3 gives the size of the pseudonym circuit'), 'Times are medians of five runs', 'Times are medians of five runs (ten for the issuance proof)')
rep(find('TABLE 3.'), '(median of 10 runs, single machine)', '(medians, single machine)')
# F22 한계: 성명 안의 챌린지는 이제 없다
rep(find('The AA is on the login path.'),
    'at the cost of giving up per-login challenges inside the statement.',
    'at the cost of a longer-lived statement, and therefore a longer exposure between a revocation and the expiry that bounds it.')
d.save(P); print('patched')
