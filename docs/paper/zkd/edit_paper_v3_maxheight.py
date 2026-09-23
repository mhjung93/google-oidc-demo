# v3 갱신(2026-09-18 오후): max_height 를 지갑이 정하고 AA 는 그대로 서명, 검증자가 상한 L 강제. 실측치 갱신. 제자리 수정(백업 .bak-0918b).
import shutil
from docx import Document
from docx.oxml.ns import qn
P = 'documents/2026-09-18_v3_zk-Delegation_research_article.docx'
shutil.copy(P, P + '.bak-0918b')
d = Document(P)
def find(prefix):
    for p in d.paragraphs:
        if p.text.startswith(prefix): return p
    raise KeyError(prefix)
def set_text(p, text):
    rs = p.runs; rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
def rep(p, old, new):
    assert old in p.text, (old[:60], p.text[:80]); set_text(p, p.text.replace(old, new, 1))
def set_cell(tc, text):
    ps_ = tc.findall(qn('w:p'))
    for q in ps_[1:]: tc.remove(q)
    q = ps_[0]; rs = q.findall(qn('w:r'))
    for r in rs[1:]: q.remove(r)
    ts = rs[0].findall(qn('w:t'))
    for t in ts[1:]: rs[0].remove(t)
    ts[0].text = text; ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
def set_row(t, first, vals):
    for r in t.rows:
        if r.cells[0].text.strip() == first:
            for c, v in zip(r.cells, vals): set_cell(c._tc, v)
            return
    raise KeyError(first)

# V-C 발급: 지갑이 max_height 를 정한다
p = find('All scalars lie in [0, 2^250); unused attribute slots are zero.')
rep(p, 'The wallet then sends the AA the tuple (uid, C_pt, chainid, allowAgent), a proof π_issue, and a signature sig_u = Sign_sk_u(Poseidon(D_req, C_pt.x, C_pt.y, chainid, allowAgent)) with D_req a domain tag.',
    'The wallet chooses the expiry of the statement itself: it reads the chain head and sets max_height = ⌈(head + T) / G⌉·G, with a lifetime T of 300 blocks and a grid G of 100 blocks in the prototype, so that every statement requested within one grid window carries the same expiry (Section VII-C). It then sends the AA the tuple (uid, C_pt, chainid, allowAgent, max_height), a proof π_issue, and a signature sig_u = Sign_sk_u(Poseidon(D_req, C_pt.x, C_pt.y, chainid, allowAgent, max_height)) with D_req a domain tag, so that no intermediary can alter the requested expiry.')
p = find('with uid, C_pt, and cm_u public.')
rep(p, 'The AA checks, in this order, the request format, that the account is not disabled, that chainid is in its allow list, sig_u under the stored pk_u, and π_issue. It then reads the head height of the chain named by chainid and sets max_height = ⌈(head + T) / G⌉·G, with a lifetime T of 300 blocks and a grid G of 100 blocks in the prototype, so that every statement issued within one grid window carries the same expiry (Section VII-C); it computes C itself from C_pt and signs',
    'The AA checks, in this order, the request format, that the account is not disabled, that chainid is in its allow list and that the chain answers, that max_height is below 2^64, sig_u under the stored pk_u, and π_issue. It neither chooses nor adjusts the expiry: as with the max_epoch of zkLogin, the expiry is the user\'s choice, the issuer binds it, and the verifier bounds it (Sections V-F and V-I). It computes C itself from C_pt and signs')
p = find('The AA never sees the challenge. A replayed issuance request')
rep(p, 'the AA then merely extends the expiry of its record for that C, and the statement remains bound',
    'the AA then holds one more copy of a record it already has, and the statement remains bound')
# V-F 서비스 검사: 상한 L
p = find('The wallet sends the service the proof, its public inputs, and a signature σ = Sign_sk_i(r_s)')
rep(p, 'It checks that the block head it just read is at most max_height, that chainid is the chain it reads,',
    'It checks that the block head it just read is at most max_height and that max_height is at most that head plus a lifetime bound L, 400 blocks in the prototype, so that a wallet cannot request a statement that outlives the bound; that chainid is the chain it reads,')
# V-I 컨트랙트 검사
p = find('A transaction is execute(payload, σ, π, public inputs), submitted by any relayer')
rep(p, 'that the block number is at most max_height; and finally the Groth16 argument.',
    'that the block number is at most max_height and max_height at most the block number plus the lifetime bound L fixed in the factory; and finally the Groth16 argument.')
# VII-C 체인 채널: 양자화는 지갑의 몫
p = find('The chain is a second channel.')
rep(p, 'max_height is quantized to a grid, so every statement issued within one window of G blocks carries the same value, and the AA\'s issuance log narrows a transaction to the users issued in that window.',
    'max_height is quantized to a grid by the wallet, so every statement requested within one window of G blocks carries the same value, and the AA\'s issuance log narrows a transaction to the users issued in that window. The grid protects the user, not the AA, so it is applied by the wallet; a wallet that ignores it exposes only its own issuance time, and the verifiers\' bound L caps what any wallet can request.')
# VIII-B 비용
p = find('Table 3 gives the size of the pseudonym circuit')
rep(p, 'Times are medians of five runs (ten for the issuance proof)', 'Times are medians of ten runs')
p = find('Per-login cost is one issuance round trip')
rep(p, 'An on-chain transaction costs the same proof, reused, plus about 390,000 gas for verification and execution, roughly eighteen times a plain transfer; this is the cost of verifying at every transaction that Section V-I accepts.',
    'An on-chain transaction costs the same proof, reused, plus 340,000 to 390,000 gas for verification and execution depending on the inner call, sixteen to eighteen times a plain transfer; this is the cost of verifying at every transaction that Section V-I accepts. Measured end to end through the wallet agent\'s HTTP interface against the isolated stack, a login with issuance took 1.39 s (issuance 0.42 s, proof 0.90 s), a re-validation with the cached proof 33 ms, and an on-chain transaction with the cached proof 157 ms including mining on the local node.')
p = find('Publication cost was measured on the local Hardhat node')
rep(p, 'An execute() call with a valid proof consumed 387,682 gas, the bulk of it the pairing check of the Groth16 verifier; a heartbeat is a publication with no leaves and costs the base of roughly 40,000 gas.',
    'An execute() call with a valid proof consumed 387,905 gas for a value transfer to a fresh address and 339,341 gas for a call without value to an existing account, the bulk of it the pairing check of the Groth16 verifier; deploying an account through the factory consumed 713,542 gas once; a heartbeat is a publication with no leaves and consumed 40,261 gas.')
# 표
for t in d.tables:
    heads = [r.cells[0].text.strip() for r in t.rows]
    if 'max_height' in heads:
        for r in t.rows:
            if r.cells[0].text.strip() == 'max_height': set_cell(r.cells[1]._tc, 'statement expiry as a block height, chosen by the wallet on a grid'); break
    if 'Groth16 proving time (median)' in heads:
        set_row(t, 'Groth16 proving time (median)', ['Groth16 proving time (median)', '857 ms'])
        set_row(t, 'Groth16 verification time (off chain)', ['Groth16 verification time (off chain)', '14.4 ms'])
        set_row(t, 'execute() gas on chain (verification + call)', ['execute() gas on chain (verification + call)', '387,905 (value transfer) / 339,341 (call)'])
        set_row(t, 'Statement lifetime (T, grid G) / challenge lifetime', ['Statement lifetime (T, grid G, bound L) / challenge lifetime', '300 blocks, G = 100, L = 400 / 120 s'])
d.save(P); print('paper patched')
