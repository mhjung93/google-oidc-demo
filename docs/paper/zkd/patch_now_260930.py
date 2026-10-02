#!/usr/bin/env python3
"""documents/260930_now.pptx → documents/260930_now_revised.pptx (2026-09-30).

교수님 용어(IdP·RP·User·Token·auid·auid_i·rid/arid·salt·π_IdP·π_RP·π_i·PPID·pk_IdP)는 그대로 두고,
내용만 현재 구현(feat/mode3-cia)에 맞춘다. 원본은 건드리지 않는다. 바꾼 글은 BLUE.
근거: results/mode3_session_revocation_20260924.md(제약·증명 시간), results/mode3_review_fixes_20260925.md(gas·지연),
      results/mode3_disclosure_bench_20260922.md(π_u 80 ms), lib/mode3_issuance.js(시그마 증명 문장),
      docs/paper/zkd/security_formal.md(G1–G3 = 정리 1–3), 2026-09-30 대화의 G11 초안.
실행: python3 docs/paper/zkd/patch_now_260930.py
"""
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC = 'documents/260930_now.pptx'
OUT = 'documents/260930_now_revised.pptx'
BLUE = '1F4EB0'      # 바꾼 글. 발표용으로 검정을 원하면 '000000'
GREY_FILL = 'F2F2F2'

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
MC = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'
FILLS = (A + 'noFill', A + 'solidFill', A + 'gradFill', A + 'blipFill', A + 'pattFill', A + 'grpFill')

prs = Presentation(SRC)
S = list(prs.slides)
assert len(S) == 10, len(S)


# ---------- 조회 ----------
def ptext(pel):
    return ''.join(t.text or '' for t in pel.iter() if t.tag in (A + 't', M + 't'))

def sptree(slide):
    return slide._element.find('.//' + P + 'spTree')

def name_of(el):
    return el.find('.//' + P + 'cNvPr').get('name')

def plain_sp(slide, name):
    """spTree 바로 아래의 p:sp — 정확히 하나."""
    hits = [el for el in sptree(slide) if el.tag == P + 'sp' and name_of(el) == name]
    assert len(hits) == 1, (name, len(hits))
    return hits[0]

def alt(slide, name):
    """수식(oMath) 때문에 mc:AlternateContent 로 싸인 도형 — Choice 쪽 이름으로 찾는다."""
    hits = [el for el in sptree(slide) if el.tag == MC + 'AlternateContent'
            and name_of(el.find(MC + 'Choice/' + P + 'sp')) == name]
    assert len(hits) == 1, (name, len(hits))
    return hits[0]

def de_alt(slide, name):
    """AlternateContent 를 Choice 의 p:sp(사본)로 바꿔 넣는다. 수식은 rebuild 로 지워진다."""
    ac = alt(slide, name)
    sp = copy.deepcopy(ac.find(MC + 'Choice/' + P + 'sp'))
    ac.addprevious(sp)
    ac.getparent().remove(ac)
    return sp

def paras(sp):
    return sp.find(P + 'txBody').findall(A + 'p')

def para_with(sp, needle):
    hits = [pel for pel in paras(sp) if needle in ptext(pel)]
    assert len(hits) == 1, (needle, len(hits))
    return hits[0]


# ---------- 글 ----------
def style(rpr, color=None, size=None, bold=None):
    rpr.attrib.pop('strike', None)
    rpr.attrib.pop('sym', None)
    for c in list(rpr):
        if c.tag == A + 'sym':
            rpr.remove(c)
    if size is not None:
        rpr.set('sz', str(int(size * 100)))
    if bold is not None:
        rpr.set('b', '1' if bold else '0')
    if color is not None:
        for c in list(rpr):
            if c.tag in FILLS:
                rpr.remove(c)
        sf = etree.Element(A + 'solidFill')
        etree.SubElement(sf, A + 'srgbClr').set('val', color)
        rpr.insert(1 if len(rpr) and rpr[0].tag == A + 'ln' else 0, sf)
    return rpr

def run_template(pel):
    r = pel.find(A + 'r')
    if r is not None and r.find(A + 'rPr') is not None:
        t = copy.deepcopy(r.find(A + 'rPr'))
    else:
        e = pel.find(A + 'endParaRPr')
        t = copy.deepcopy(e) if e is not None else etree.Element(A + 'rPr')
        t.tag = A + 'rPr'
    return style(t)

def segs_of(x):
    """'글' | ('글', 색) | [('글', 색), ...] → [(글, 색)]; 색 None 은 원본 색 유지."""
    if isinstance(x, str):
        return [(x, None)]
    if isinstance(x, tuple):
        return [x]
    return list(x)

def set_para(pel, content, color=None, size=None, bold=None, lvl=None):
    """문단의 런을 전부 새 글로 바꾼다(첫 런의 서식 유지, 취소선 제거)."""
    tmpl = run_template(pel)
    for c in list(pel):
        if c.tag != A + 'pPr':
            pel.remove(c)
    if lvl is not None:
        pp = pel.find(A + 'pPr')
        if pp is None:
            pp = etree.Element(A + 'pPr')
            pel.insert(0, pp)
        if lvl:
            pp.set('lvl', str(lvl))
        else:
            pp.attrib.pop('lvl', None)
    for text, col in segs_of(content):
        r = etree.SubElement(pel, A + 'r')
        r.append(style(copy.deepcopy(tmpl), col if col is not None else color, size, bold))
        etree.SubElement(r, A + 't').text = text
    return pel

def add_after(pel, content, lvl=None, color=None, size=None, bold=None):
    new = copy.deepcopy(pel)
    pel.addnext(new)
    return set_para(new, content, color=color, size=size, bold=bold, lvl=lvl)

def append_run(pel, text, color=BLUE, size=None, bold=None):
    runs = pel.findall(A + 'r')
    tmpl = copy.deepcopy(runs[-1].find(A + 'rPr')) if runs and runs[-1].find(A + 'rPr') is not None else run_template(pel)
    r = etree.Element(A + 'r')
    r.append(style(tmpl, color, size, bold))
    etree.SubElement(r, A + 't').text = text
    if runs:
        runs[-1].addnext(r)
    else:
        e = pel.find(A + 'endParaRPr')
        (e.addprevious if e is not None else pel.append)(r)
    return pel

def rebuild(sp, lines):
    """txBody 를 통째로 새로 쓴다. lines = [dict(content, lvl, color, size, bold)]. pPr 은 원본의 같은 lvl 것을 복사."""
    tx = sp.find(P + 'txBody')
    old = tx.findall(A + 'p')
    ppr = {}
    for pel in old:
        pp = pel.find(A + 'pPr')
        ppr.setdefault(int(pp.get('lvl', '0')) if pp is not None else 0, pp)
    tmpl = run_template(old[0])
    for pel in old:
        tx.remove(pel)
    for ln in lines:
        lvl = ln.get('lvl', 0)
        pel = etree.SubElement(tx, A + 'p')
        src = ppr.get(lvl, ppr.get(0))
        if src is not None:
            pp = copy.deepcopy(src)
            if lvl:
                pp.set('lvl', str(lvl))
            else:
                pp.attrib.pop('lvl', None)
            pel.append(pp)
        elif lvl:
            etree.SubElement(pel, A + 'pPr').set('lvl', str(lvl))
        for text, col in segs_of(ln['content']):
            r = etree.SubElement(pel, A + 'r')
            r.append(style(copy.deepcopy(tmpl), col if col is not None else ln.get('color'), ln.get('size'), ln.get('bold')))
            etree.SubElement(r, A + 't').text = text
    return sp

def place(sp, left=None, top=None, width=None, height=None):
    xfrm = sp.find(P + 'spPr/' + A + 'xfrm')
    off, ext = xfrm.find(A + 'off'), xfrm.find(A + 'ext')
    if left is not None: off.set('x', str(int(Inches(left))))
    if top is not None: off.set('y', str(int(Inches(top))))
    if width is not None: ext.set('cx', str(int(Inches(width))))
    if height is not None: ext.set('cy', str(int(Inches(height))))

def textbox(slide, left, top, width, height, lines, size=12, color=BLUE, fill=None):
    """lines: 'text' | (text, size) | (text, size, color)."""
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = RGBColor.from_string(fill)
    for i, ln in enumerate(lines):
        ln = (ln,) if isinstance(ln, str) else tuple(ln)
        text, sz, col = ln[0], (ln[1] if len(ln) > 1 else size), (ln[2] if len(ln) > 2 else color)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = text
        r.font.size = Pt(sz)
        if col:
            r.font.color.rgb = RGBColor.from_string(col)
    return tb


# 7장 본문(수식 때문에 AlternateContent)을 갈아끼울 틀 — 5장 본문을 손대기 전에 복사해 둔다.
BODY_TMPL = copy.deepcopy(plain_sp(S[4], '내용 개체 틀 2'))
L0, L1, L2 = 20, 16, 15


# ---------- 2장 Revocation of Session Attributes ----------
body = plain_sp(S[1], '내용 개체 틀 2')
p = para_with(body, '별도')
set_para(p, '같은 Merkle tree 를 사용 — 리프 종류만 둘: 사용자 리프 H(4, H(C_u)), 세션 리프 H(5, H(C_s)) (도메인 태그로 구분)', color=BLUE, size=L1)
p = add_after(p, 'π_RP 가 같은 root 에 대해 비멤버십 2건(사용자 리프·세션 리프)을 증명 — 트리 하나, root 하나', lvl=1, color=BLUE, size=L1)
p = add_after(p, '세션 폐기 = 폐기된 세션마다 리프 하나 추가(개수 제한 없음); 폐기되지 않은 세션·토큰은 그대로', lvl=1, color=BLUE, size=L1)
p = add_after(p, '검증자용 유효 목록(VCL)은 없음 — IdP 내부 발급 기록은 있음', lvl=0, color=BLUE, size=L0)
p = add_after(p, "RP·컨트랙트는 목록 없이 IdP 서명 + 만료 전 + 두 리프 비멤버십만 봄 — 트리는 '발급 뒤 만료 전에 폐기됐는가'만 답함 → 리프 쓰기는 폐기 때만(발급 때 0)", lvl=1, color=BLUE, size=L1)
p = add_after(p, 'IdP 는 발급 때 uid → {H(C_s), expire time, chainid, allowAgent} 를 기록해 두고, 폐기 요청 때 그 기록으로 권한(sk_u 서명)·존재·만료·중복을 확인한 뒤 리프를 만듦; 만료된 기록은 정리되어 살아 있는 세션 수만큼만 유지', lvl=1, color=BLUE, size=L1)
p = add_after(p, '비용 (V7 → V8, 실측)', lvl=0, color=BLUE, size=L0)
p = add_after(p, '제약 27,329 → 37,130 (+9,801), 증명 0.83 → 1.19 s (+0.36 s), 검증 10 ms·온체인 검증 gas 불변(공개 입력 25개 그대로)', lvl=1, color=BLUE, size=L1)
p = add_after(p, 'root 게시: 리프 1개 43,856 gas; 하트비트(50블록마다) 약 40,300 gas', lvl=1, color=BLUE, size=L1)


# ---------- 3장 ID Transformation ----------
s3 = S[2]
tok = de_alt(s3, 'Rounded Rectangle 4')            # 토큰 상자 (수식 C_s·C_u·pk_i 는 글로)
place(tok, left=0.9, top=1.85, width=5.9, height=4.0)
rebuild(tok, [
    dict(content='(* Token 관련 정보 *)', size=14),
    dict(content='Token = Sign_IdP( H(C_u), H(C_s), expire time, chainid, allowAgent )', lvl=1, color=BLUE, size=13),
    dict(content='auid_i(세션별 가명) 자리는 H(C_s) — 세션마다 새 blind·새 pk_i; RP 가 보는 고정 이름 PPID 는 지갑이 계산하고 π_RP 로 증명', lvl=1, color=BLUE, size=13),
    dict(content=[('(* User attributes *)', None), (' → C_u', BLUE)], size=14),
    dict(content='uid, salt', lvl=1, size=13),
    dict(content=[('other attributes (e.g., age)', None), (' → 출생연도·국가·등급·예비 4개', BLUE)], lvl=1, size=13),
    dict(content=[('(* Session attributes *)', None), (' → C_s', BLUE)], size=14),
    dict(content=[('arid, pk_i', None), (' (session key; blind 로 은닉 — IdP 는 못 봄)', BLUE)], lvl=1, size=13),
    dict(content=[('expire time (block height)', None), (' — C_s 밖, 서명 본문에', BLUE)], lvl=1, size=13),
])
textbox(s3, 0.9, 6.42, 5.9, 0.4, ['세션 토큰 발급마다 ZKP 없음 — IdP 는 sk_u 서명만 검사 (π_IdP 는 C_u 발급 시 1회)'], size=12)

idp = plain_sp(s3, 'Rounded Rectangle 7'); place(idp, left=7.1, top=1.05, width=3.3, height=1.9)
rebuild(idp, [
    dict(content='IdP', size=20),
    dict(content='uid', lvl=1, size=13),
    dict(content='H(C_u), H(C_s) 만 봄 (커밋먼트 해시)', lvl=1, color=BLUE, size=13),
    dict(content='auid 없음 → 기록: (uid, H(C_u)) + 세션 (H(C_s), expire, chainid)', lvl=1, color=BLUE, size=13),
])
rp = plain_sp(s3, 'Rounded Rectangle 9'); place(rp, left=7.1, top=3.08, width=3.3, height=1.92)
rebuild(rp, [
    dict(content='RP', size=20),
    dict(content=[('auid → auid_i', None), (' 대신: 고정 PPID + 세션마다 새 pk_i 를 봄 (auid 는 없음)', BLUE)], lvl=1, size=13),
    dict(content=[('rid → arid', None), (' (IdP 등록 시 부여)', BLUE)], lvl=1, size=13),
    dict(content='Trace Tag (열기 전엔 난수)', lvl=1, color=BLUE, size=13),
])
usr = plain_sp(s3, 'Rounded Rectangle 10'); place(usr, left=7.1, top=5.15, width=3.3, height=1.85)
rebuild(usr, [
    dict(content='User (지갑)', size=20),
    dict(content='uid, salt (salt 는 지갑만 앎)', lvl=1, color=BLUE, size=13),
    dict(content='PPID = H(uid, salt, chainid, arid) — RP 별 고정', lvl=1, color=BLUE, size=13),
    dict(content='세션마다 C_s = Commit(arid, pk_i; blind), pk_i 새로', lvl=1, color=BLUE, size=13),
])
# 오른쪽 설명 셋 — 수식 둘은 글로 새로, Salt 는 원본 상자에 덧붙임
for nm in ('TextBox 13', 'TextBox 14'):
    ac = alt(s3, nm); ac.getparent().remove(ac)
textbox(s3, 10.5, 3.15, 2.65, 0.7, ['r_RP (= C_s 의 blind) = RP privacy: arid 를 IdP 에 숨김'], size=11)
textbox(s3, 10.5, 3.95, 2.65, 0.7, ['r_S = session nonce: RP 가 발급, sk_i 로 서명(증명 밖)'], size=11)
salt = plain_sp(s3, 'TextBox 11'); place(salt, left=10.5, top=5.4, width=2.65, height=0.9)
salt.find(P + 'txBody/' + A + 'bodyPr').set('wrap', 'square')
set_para(paras(salt)[0], [('Salt = user privacy', None), (': 지갑만 앎 → IdP 도 PPID 를 못 만듦(공모 차단)', BLUE)], size=11)
s3.notes_slide.notes_text_frame.text = (
    '용어 대응(슬라이드 → 코드·논문): IdP → AA/CIA(cia.js) · RP → 서비스(mode3_rp.js) · Token → σ_AA = Sign_AA(Cf_u, Cf_s, max_height, chainid, allowAgent) · '
    'salt → s_u · auid_i(세션별 가명) → 세션마다 새 C_s(H(C_s) = Cf_s)·pk_i · PPID(RP 별 고정) = Poseidon(uid, s_u, chainid, arid) · auid → 없음 · C_u/C_s → Pedersen 커밋먼트, H(C_u) = Cf_u, H(C_s) = Cf_s (Poseidon) · '
    'expire time → max_height · π_IdP → π_u(lib/mode3_issuance.js, 시그마 프로토콜) · π_RP → pi_cred V8(circuits) · π_i → 같은 증명 · '
    'pk_IdP(태그) → pk_trace = X_svc + X_AA · r_RP → blind_s · r_S → 로그인 nonce r_s · Trace Tag → (c1, c2).')


# ---------- 4장 π_IdP ----------
s4 = S[3]
body = plain_sp(s4, '내용 개체 틀 2')
set_para(para_with(body, 'RP FE'), [('사용자 → IdP 에게 ', None),
        ('직접 제출 (RP FE 경유 없음 → IdP 는 어느 RP 인지 모름); C_u 발급 시 1회 — 세션 토큰마다는 ZKP 없음(sk_u 서명만 검사)', BLUE)], size=L2)
set_para(para_with(body, '증명 내용'), '증명 내용', size=L2)
set_para(para_with(body, 'uid-auid'), [('(1) uid-auid 관계 증명', None),
        (' → C_u 열림 증명: C_u = Commit(uid, salt, a; blind) — uid 는 로그인 계정, a 는 IdP 기록, salt 는 등록 때 낸 cm_u = Commit(salt; r_u) 와 같음 (auid = H(Salt, uid) 대신)', BLUE)], size=14)
set_para(para_with(body, '랜덤값'), [('(2) user attribute commitment를 user attribute a와 랜덤값 r로 만들었는지', None),
        (' — (1) 에 포함: 같은 시그마 증명(회로 없음; 증명 80 ms·검증 100 ms)이 보임', BLUE)], size=14)
# 그림 라벨
lab = plain_sp(s4, 'TextBox 57'); place(lab, left=1.6, width=1.0)
set_para(paras(lab)[0], 'C_u, cm_u', color=BLUE, size=12)
set_para(paras(plain_sp(s4, 'TextBox 15'))[0], 'salt, r', color=BLUE, size=14)
set_para(paras(plain_sp(s4, '직사각형 42'))[0], 'Commitment Opening', color=BLUE, size=11)
set_para(paras(plain_sp(s4, 'TextBox 51'))[0], 'Commit(…)', color=BLUE, size=12)
rebuild(de_alt(s4, 'TextBox 60'), [dict(content='C_u′', color=BLUE, size=12)])
set_para(paras(plain_sp(s4, 'Rounded Rectangle 38'))[0], '(uid, H(C_u), 속성) 발급 기록', color=BLUE, size=14)
set_para(paras(plain_sp(s4, 'Rounded Rectangle 50'))[0], '세션 토큰 발행 (이후 ZKP 없음)', color=BLUE, size=14)


# ---------- 5장 π_RP ----------
s5 = S[4]
body = plain_sp(s5, '내용 개체 틀 2')
set_para(para_with(body, 'RP 에게 제출'), [('사용자 → RP 에게 제출', None),
        (' — 로그인마다 π_RP 1건 (Groth16, 제약 37,130; 증명 1.19 s·검증 10 ms; 공개 입력 25개)', BLUE)], size=L1)
set_para(para_with(body, 'not expired'), [('(1) Valid Token (IdP 서명 && not expired)', None),
        (' — IdP 서명(EdDSA)은 회로 안에서 검증; expire time 은 공개 입력이고 만료 비교는 RP·컨트랙트가 현재 블록 높이로', BLUE)], size=L2)
set_para(para_with(body, '정보 정확성'), [('(2) Valid Token 내 IDs (auid, auid_i, arid) 정보 정확성 증명', None),
        (' → 실제로는 C_u = Commit(uid, salt, a; blind_u), C_s = Commit(arid, pk_i; blind_s) 의 열림 증명 — auid 는 없고 auid_i 자리는 C_s(세션마다 새로); arid·pk_i 는 공개 입력', BLUE)], size=L2)
set_para(para_with(body, 'PPID 검증'), [('(3) Valid Token 내 IDs 기반하여 생성된 PPID 검증', None),
        (' — PPID = H(uid, salt, chainid, arid) 를 (2) 와 같은 uid·salt 로 계산해 공개 입력과 일치 (RP 별 고정)', BLUE)], size=L2)
p = para_with(body, 'Selective Disclosure')
set_para(p, [('(4) (if any) Valid Token 내 Vector Commitment 기반 사용자 속성값 Selective Disclosure 증명', None),
        (' — 속성 4개(출생연도·국가·등급·예비)에 범위(lo ≤ a ≤ hi)·집합 소속(깊이 8 Merkle) 조건을 마스크로 선택; 값 자체는 비공개', BLUE)], size=L2)
p = add_after(p, "(5) 폐기 트리 비멤버십 2건 — 사용자 리프·세션 리프 ∉ Tree(root), root 는 공개 입력. 리프 H(C_u)·H(C_s) 는 IdP 발급 기록과 바로 이어져 공개할 수 없으므로 증명 안에서만 가능 → 'merkle proof 빼기' 대신 유지 제안", lvl=1, color=BLUE, size=L2)
p = add_after(p, '(6) Trace Tag = Enc(pk_trace, H(uid, arid)) 가 같은 uid·arid 로 만들어졌음 (다음 장)', lvl=1, color=BLUE, size=L2)


# ---------- 6장 데모 화면 ----------
s6 = S[5]
set_para(paras(plain_sp(s6, '제목 1'))[0], [('(4)', None), ('의 데모 화면 — 구현됨 (지갑 UI: 공개할 속성 조건 선택)', BLUE)])
textbox(s6, 10.15, 2.3, 3.0, 4.3, [
    '지갑이 고른 조건이 π_RP 의 공개 입력(마스크·lo/hi·집합 root)이 된다',
    'RP 는 값이 아니라 조건 충족 여부만 본다',
    '예) 나이 ≥ 19 → 출생연도 ≤ 2007',
    '예) 국가 ∈ {410, 392, 840, 276, 250}',
    "'내 실제 값으로 채우기' 는 데모 편의 버튼",
    "다음: RP 화면에 '조건 충족' 만 보이는 장면 추가",
], size=12)


# ---------- 7장 π_i ----------
s7 = S[6]
ac = alt(s7, '내용 개체 틀 2')
body = copy.deepcopy(BODY_TMPL)
ac.addprevious(body); ac.getparent().remove(ac)
rebuild(body, [
    dict(content=[('사용자 혹은 RP → 트랜잭션 생성시 첨부 → 수신자에게 전달 (TX + π_i + Trace Tag', None), (' + σ_tx)', BLUE)], size=L1),
    dict(content='지갑이 (TX, π_RP, Trace Tag, σ_tx) 를 만들고 RP 가 relayer 로 계정 컨트랙트에 제출', lvl=1, color=BLUE, size=L2),
    dict(content='π_i 는 별도 증명이 아님 — 로그인 때의 π_RP(같은 회로·같은 공개 입력 25개)를 세션 안에서 재사용; 트랜잭션마다 새 ZKP 없음, 검증은 컨트랙트가 매번', color=BLUE, size=L1),
    dict(content='증명 내용', size=L1),
    dict(content=[('(1) π_RP 증명 내용 (Valid Token && IDs 검증 && PPID 검증', None), (' && 비멤버십 && 선택 공개)', BLUE)], lvl=1, size=L2),
    dict(content=[('(2) session public key pk_i(= pk_s) 증명', None),
                  (' — pk_i 는 π_RP 의 공개 입력(C_s 안의 값); 트랜잭션은 sk_i 서명 σ_tx 로 결속: σ_tx 가 TX 본문 11워드 + expire time + allowAgent + Trace Tag(3워드) = 16워드를 덮음', BLUE)], lvl=1, size=L2),
    dict(content=[('(3) Trace Tag 가 Valid Token 내 ', None), ('uid·arid', BLUE), (' 기반하여 생성되었음', None),
                  (' — auid 가 아니라 H(uid, arid) 를 암호화; 회로 안에서 같은 uid·arid 임을 검증', BLUE)], lvl=1, size=L2),
])
ac = alt(s7, 'TextBox 7'); ac.getparent().remove(ac)   # 'Trace Tag = Enc(pk_IdP, auid)' 수식 상자
textbox(s7, 1.0, 4.95, 11.4, 1.55, [
    ('Trace Tag = Enc(pk_trace, H(uid, arid)),   pk_trace = pk_RP + pk_IdP   (auid 아님)', 16),
    ('열기: RP 조각 + IdP 조각 + 운영자 승인 → H(uid, arid) → IdP 기록에서 uid — 2-of-2 가산 복호; threshold 서명은 없음(IdP 조각에만 Schnorr PoK)', 13),
    ('컨트랙트 검사 순서: nonce → σ_tx → 성명(계정·IdP 키·allowAgent·Trace Tag 결속·root·root 나이·만료) → Groth16 검증 — 실행 gas 약 417k–456k', 13),
], fill=GREY_FILL)


# ---------- 8장 Address Abstraction ----------
rebuild(plain_sp(S[7], '내용 개체 틀 2'), [
    dict(content='Chain-agnostic 특성을 제외하고 다른 특성들은 이미 포함되어 있음', size=18),
    dict(content=[('Uniqueness, Immutability, Unlinkability', None), (' = zkAA 의 Injectiveness · Tamper resistance · Privacy-preservation = 우리 목표 G1 · G2 · G3', BLUE)], lvl=1, size=14),
    dict(content='G1 Injectiveness: PPID = H(uid, salt, chainid, arid) — 같은 문맥에서 사용자가 다르면 주소도 다름 (정리 1)', lvl=1, color=BLUE, size=14),
    dict(content='G2 Tamper resistance: 계정 = CREATE2(factory, PPID) — 지갑·IdP·RP 누구도 그 계정의 주소를 바꾸지 못함 (정리 2)', lvl=1, color=BLUE, size=14),
    dict(content='G3 Privacy-preservation: 주소에서 uid 를 얻지 못하고 RP·체인 간 주소를 잇지 못함 (정리 3)', lvl=1, color=BLUE, size=14),
    dict(content='나머지 zkAA 특성: Unforgeability → G6·G7(IdP 서명 + σ_tx), Correctness → 완전성, Chronicle → G10 + 온체인 개봉 기록', lvl=1, color=BLUE, size=14),
    dict(content=[('따라서 chain-agnostic ', None), ('만 별도 특성으로 정의 — 문맥 간 신원 결속 (cross-context consistency, G11)', BLUE)], size=18),
    dict(content='정의: 어느 문맥 (chainid, arid) 에서든 계정의 주소는 등록 때 커밋한 신원 (uid, salt) 하나에서 파생한 H(uid, salt, chainid, arid) 이고, 받아들여진 어떤 증명도 두 번째 신원으로 열리지 않는다', lvl=1, color=BLUE, size=14),
    dict(content="근거: π_IdP 가 salt 를 등록 커밋먼트 cm_u 에 결속(Pedersen 결합성) + IdP 는 uid 당 등록 1회 → 정리 2 의 증명에서 '같은 문맥' 전제만 풀면 그대로 성립(상한 동일)", lvl=1, color=BLUE, size=14),
    dict(content='뜻: 등록 한 번이 모든 체인·RP 의 주소를 정한다 — 주소가 같은 것이 아니라 규칙과 비밀이 하나이고, 주소는 문맥마다 다르되 서로 이어지지 않음(G3)', lvl=1, color=BLUE, size=14),
    dict(content='Address abstraction 자체는 특성이 아니라 정의: 주소 함수 Addr(id, ctx) = H(uid, salt, chainid, arid) 와 그것을 검증하는 성명(π_RP 의 PPID 조건)', lvl=1, color=BLUE, size=14),
])


# ---------- 9장 Todo ----------
body = plain_sp(S[8], '내용 개체 틀 2')
set_para(para_with(body, 'salt'), [('salt 변경에 따른 ', None), ('토큰(C_u) 발급 실패', BLUE), (' 시나리오 추가', None),
        (' — 세션 인증이 아니라 발급 단계에서 거절: π_IdP 가 salt 를 등록 커밋먼트 cm_u 에 결속 (G2·G11 시연)', BLUE)])
append_run(para_with(body, '서로 다른'), ' — 구현 완료(9/30): 지갑이 RP origin 목록을 받음, RP #2(:3101) 실행 스크립트; 같은 사용자가 RP 마다 다른 PPID·다른 주소')
p = para_with(body, '가상자산거래소')
append_run(p, ' — RP #2 를 거래소로: 나이 ≥ 19 · 국가 집합 조건 공개, Trace Tag 개봉(RP 조각 + IdP 조각 + 운영자 승인)')
p = add_after(p, '폐기 트리(RCL) 그래픽 표시 — 리프 목록·비멤버십 구간·root 게시 시각 (관리 페이지, 다음 작업)', lvl=2, color=BLUE)
p = add_after(p, '회의에서 정할 것', lvl=1, color=BLUE)
p = add_after(p, '세션 폐기 구조: 트리 하나 + 세션 리프(현재 구현) vs 사용자별 슬롯 리프', lvl=2, color=BLUE)
p = add_after(p, "π_RP 의 폐기 비멤버십 유지 여부 ('merkle proof 빼기' 의견)", lvl=2, color=BLUE)
p = add_after(p, 'π_IdP 를 C_u 발급 시 1회로 두는 근거', lvl=2, color=BLUE)
p = add_after(p, '공개 속성 예시(출생연도·국가·등급·예비) 확정', lvl=2, color=BLUE)

prs.save(OUT)
# 남은 수식 확인 — 제목 π_IdP·π_RP·π_i 와 4장 빨간 π_IdP 라벨만 남아야 한다
left = [(i + 1, name_of(el.find(MC + 'Choice/' + P + 'sp'))) for i, s in enumerate(Presentation(OUT).slides)
        for el in sptree(s) if el.tag == MC + 'AlternateContent']
print('saved', OUT, '| 남은 수식 도형:', left)
