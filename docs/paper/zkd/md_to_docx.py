# Markdown(형식 문서) → docx. 사용: python3 docs/paper/zkd/md_to_docx.py <src.md> <dst.docx>
# (2026-09-18 형식 문서 두 편을 만들 때 쓴 인라인 렌더러를 파일로 옮긴 것 — 표·제목·목록·굵게만 처리한다.)
import re, sys
from docx import Document
from docx.shared import Pt
src_path, dst = sys.argv[1], sys.argv[2]
s = open(src_path, encoding='utf-8').read()
src = s.split('\n'); d = Document(); st = d.styles['Normal']; st.font.name = 'DejaVu Sans'; st.font.size = Pt(10)
i = 0
while i < len(src):
    line = src[i]
    if line.startswith('|') and i+1 < len(src) and src[i+1].startswith('|---'):
        rows = []
        while i < len(src) and src[i].startswith('|'):
            if not src[i].startswith('|---'): rows.append([c.strip() for c in src[i].strip('|').split('|')])
            i += 1
        t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c); cell.text = ''
                run = cell.paragraphs[0].add_run(val.replace('**','').replace('`','')); run.font.size = Pt(8); run.bold = (r == 0)
        continue
    if line.startswith('# '): d.add_heading(line[2:], 0)
    elif line.startswith('## '): d.add_heading(line[3:], 1)
    elif line.startswith('### '): d.add_heading(line[4:], 2)
    elif re.match(r'^\d+\. ', line): d.add_paragraph(re.sub(r'^\d+\. ', '', line).replace('**','').replace('`',''), style='List Number')
    elif line.startswith('- '):
        p_ = d.add_paragraph(style='List Bullet')
        for k, seg in enumerate(re.split(r'\*\*', line[2:])): r = p_.add_run(seg.replace('`','')); r.bold = (k % 2 == 1)
    elif line.strip() == '': pass
    else:
        p_ = d.add_paragraph()
        for k, seg in enumerate(re.split(r'\*\*', line)): r = p_.add_run(seg.replace('`','')); r.bold = (k % 2 == 1)
    i += 1
d.save(dst); print('docx ok', dst)
