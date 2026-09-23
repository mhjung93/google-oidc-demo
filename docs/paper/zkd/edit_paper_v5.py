# zk-Delegation 원고 v4 → v5 (2026-09-23): 자격증명 이중 구조(C_u/C_s, 2026-09-21 설계) + 오프체인 root 나이
# fail-closed(2026-09-23 점검 C-1 결정) + 지갑 증분 동기화(2026-09-23 설계) + 실측 갱신. v4 는 건드리지 않고 v5 사본에만 적용한다.
#
# 표기 대응(코드 → 논문): userCommit C_u / sessionCommit C_s, Cf = Poseidon(C.x, C.y), π_u(코드) = 사용자 자격증명 증명 π_u,
# 세션 발급에는 ZKP 없음(sig_u), 리프 = mask_252(Poseidon(4, Cf_u)), σ_AA = Sign(D_cred, Cf_u, Cf_s, max_height, chainid, allowAgent).
# 속성 슬롯 번호는 논문 관례 a_1..a_4(코드는 a₀..a₃). 생성원은 논문 관례 G_1(uid), G_2(arid), G_3(s_u), G_4(pk_i), G_5..G_8(a), H(blind).
# 실측: results/mode3_onchain_bench_20260921.md(V5), results/mode3_disclosure_bench_20260922.md(V6), results/mode3_rcl_sync_20260923.md.
import copy, shutil, sys
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIG = 'docs/paper/zkd/fig1_login_v5.png'   # draw_fig1_v5.py 산출물

SRC = 'documents/2026-09-22_v4_zk-Delegation_research_article.docx'
DST = 'documents/2026-09-23_v5_zk-Delegation_research_article.docx'
shutil.copy(SRC, DST)
d = Document(DST)
XMLNS = '{http://www.w3.org/XML/1998/namespace}space'
PARAS = list(d.paragraphs)
IDX = {id(p._p): i for i, p in enumerate(PARAS)}
LOG = []

def pidx(p): return IDX.get(id(p._p), 'new')
def log(p, before, after, kind='edit'): LOG.append((pidx(p), kind, before, after))

def find(prefix, style=None):
    for p in PARAS:
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise KeyError(prefix)

def set_text(p, text):
    rs = p.runs
    assert len(rs) >= 1, p.text[:40]
    rs[0].text = text
    for r in rs[1:]:
        r._r.getparent().remove(r._r)

def rep(p, old, new):
    t = p.text
    assert old in t, (old[:80], t[:120])
    set_text(p, t.replace(old, new, 1))
    log(p, old, new)

def clone_after(p, text, template=None):
    el = copy.deepcopy((template or p)._p)
    ts = el.findall('.//' + qn('w:t'))
    assert len(ts) >= 1
    ts[0].text = text
    ts[0].set(XMLNS, 'preserve')
    for t in ts[1:]:
        t.getparent().remove(t)
    p._p.addnext(el)
    q = Paragraph(el, p._parent)
    LOG.append((f'after {pidx(p)}', 'insert', '', text))
    return q

def remove(p):
    LOG.append((pidx(p), 'remove', p.text, ''))
    p._p.getparent().remove(p._p)

def set_cell(tc, text):
    ps_ = tc.findall(qn('w:p'))
    for q in ps_[1:]: tc.remove(q)
    q = ps_[0]
    rs = q.findall(qn('w:r'))
    for r in rs[1:]: q.remove(r)
    if not rs: rs = [etree.SubElement(q, qn('w:r'))]
    ts = rs[0].findall(qn('w:t'))
    for t in ts[1:]: rs[0].remove(t)
    if not ts: ts = [etree.SubElement(rs[0], qn('w:t'))]
    ts[0].text = text; ts[0].set(XMLNS, 'preserve')

def set_row(table, key, row):
    for r in table.rows:
        if r.cells[0].text.strip() == key:
            before = ' | '.join(c.text.strip() for c in r.cells)
            for tc, text in zip(r.cells, row): set_cell(tc._tc, text)
            LOG.append(('table', 'row', before, ' | '.join(row)))
            return
    raise KeyError(key)

def row_after(table, key, row):
    tbl = table._tbl
    for tr in tbl.findall(qn('w:tr')):
        if tr.findall(qn('w:tc'))[0].xpath('string(.)').strip() == key:
            new = copy.deepcopy(tr)
            for tc, text in zip(new.findall(qn('w:tc')), row): set_cell(tc, text)
            tr.addnext(new)
            LOG.append(('table', 'row-insert', f'after "{key}"', ' | '.join(row)))
            return
    raise KeyError(key)

def row_remove(table, key):
    tbl = table._tbl
    for tr in tbl.findall(qn('w:tr')):
        if tr.findall(qn('w:tc'))[0].xpath('string(.)').strip() == key:
            LOG.append(('table', 'row-remove', key, ''))
            tbl.remove(tr); return
    raise KeyError(key)

BODY = find('A Web3 service must recognize a returning user')          # 본문 문단 템플릿(PARA)
HEAD = find('B. REGISTRATION')                                          # 소절 제목 템플릿(Normal + spacing)
EQ = find('C_pt = uid·G_1 + arid·G_2')                                  # 수식 문단 템플릿

# ---------------------------------------------------------------- Abstract (IEEE Access 상한 250 단어 — v4 는 291 단어였다)
p = find('ABSTRACT ')
ABSTRACT = ("ABSTRACT A Web3 service identifies its users by a blockchain address: reuse across services and chains makes a user linkable, and a login says nothing about who holds the key. "
  "Delegating authentication to an existing account tells the authority which service the user visits. "
  "zk-Delegation redefines the address as the name under which a service knows a user, derived from one user secret for every chain and service, and realized on chain as a contract account whose transactions carry the login's proof. "
  "The authentication authority (AA) verifies once, in zero knowledge, a commitment to the user's identity, secret, and recorded attributes; per login it signs, without a proof, a statement over that credential and a session commitment with a block-height expiry, and learns neither pseudonym nor service, not even from the chain. "
  "The wallet can disclose a chosen range of an attested attribute inside the same proof. "
  "Each service receives a pairwise pseudonym with a zero-knowledge proof of its derivation and non-revocation and cannot link a user across services or chains. "
  "A trace tag, the identity encrypted under a key shared by AA and service, lets a disputed session be opened only by both with an operator. "
  "One proof serves a session. "
  "Revocation is an append-only tree of credentials whose root the AA publishes on a public chain; users can revoke their own accounts. "
  "A prototype proof has about 25,400 constraints, proves in under a second, and verifies in milliseconds off chain or about 400,000 gas on chain.")
set_text(p, ABSTRACT); log(p, 'ABSTRACT (291 words)', 'ABSTRACT (rewritten, ≤ 250 words)')

# ---------------------------------------------------------------- Introduction
p = find('zk-Delegation combines address abstraction with delegated authentication')
rep(p, 'The wallet asks the AA for a statement over a Pedersen commitment to the user identity, the service identifier, the user secret, a session public key, and the user\'s attributes as the AA holds them on record, together with a per-session agent-permiss',
    'Once per user, the wallet has the AA verify in zero knowledge a Pedersen commitment to the user identity, the user secret, and the user\'s attributes as the AA holds them on record, the user credential; per login it asks the AA for a statement over that credential and a second commitment to the service identifier and a session public key, together with a per-session agent-permiss')
p = find('2) We construct zk-Delegation')
rep(p, 'in which the AA signs a user-chosen commitment with a block-height expiry and learns neither the service nor the pseudonym, not even by reading the chain, and attests the user\'s attributes without learning where they are disclosed,',
    'in which the AA verifies a user credential once and signs per-session statements over it with a block-height expiry and no further proof, learns neither the service nor the pseudonym, not even by reading the chain, and attests the user\'s attributes without learning where they are disclosed,')
p = find('3) We give a pseudonym proof')
rep(p, 'binds the pseudonym, the commitment, the AA signature, the trace tag, and non-membership in a chain-published revocation tree into one statement,',
    'binds the pseudonym, the two commitments, the AA signature, the trace tag, and non-membership of the user credential in a chain-published revocation tree into one statement,')
p = find('4) We implement a prototype')
rep(p, 'covering service registration and approval, statement issuance, pseudonym verification,',
    'covering service registration and approval, credential and statement issuance, pseudonym verification,')
p = find('The remainder of the paper is organized as follows.')
rep(p, 'including service authentication, attribute attestation and selective disclosure, the pseudonym proof,',
    'including service authentication, the user credential and attribute attestation, statement issuance, selective disclosure, the pseudonym proof,')

# ---------------------------------------------------------------- V-A Overview + Fig. 1 caption
p = find('Fig. 1 shows one login.')
rep(p, 'The wallet verifies the certificate against the origin it is talking to, builds a commitment C over the user identity, the service identifier, the user secret, a fresh session key, and the user\'s attributes, chooses whether an AI agent may act in this session (allowAgent), and asks the AA for a statement over C (step 2). The AA authenticates the user, verifies that the commitment is well formed and carries the registered secret, and returns a signature over C, an expiry given as a block height, the chain identifier, and the flag (step 3).',
    'The wallet verifies the certificate against the origin it is talking to, builds a session commitment C_s over the service identifier and a fresh session key, chooses whether an AI agent may act in this session (allowAgent), and asks the AA, under a signature with its long-term key, for a statement over its user credential C_u, a commitment to the user identity, the user secret, and the user\'s attributes that the AA verified once in zero knowledge (Section V-B), and over C_s (step 2). The AA authenticates the user, checks that C_u is the user\'s active credential, and returns a signature over C_u, C_s, an expiry given as a block height, the chain identifier, and the flag (step 3); no proof is made at this step.')
rep(p, 'produces a zero-knowledge proof that binds address, commitment, signature, non-revocation, and a trace tag',
    'produces a zero-knowledge proof that binds address, the two commitments, signature, non-revocation of the credential, and a trace tag')
p = find('FIGURE 1. One login')
rep(p, 'FIGURE 1. One login (steps 1–6), one on-chain transaction (step 6\'), and one authorized opening (steps 7–8) in zk-Delegation. Solid black arrows carry what the AA sees at login (C, chainid, max_height, allowAgent);',
    'FIGURE 1. One user credential (step 0, once per user), one login (steps 1–6), one on-chain transaction (step 6\'), and one authorized opening (steps 7–8) in zk-Delegation. Solid black arrows carry what the AA sees: C_u and π_u once, and Cf_u, C_s, chainid, max_height, allowAgent per login;')

# ---------------------------------------------------------------- Fig. 1 그림 교체 (v3 그림 → v5 그림)
for q in d.paragraphs:
    if q._p.findall('.//' + qn('w:drawing')):
        for r in list(q.runs): r._r.getparent().remove(r._r)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.add_run().add_picture(FIG, width=Inches(6.5))
        LOG.append((pidx(q), 'figure', 'fig1_login_v3.png', FIG))
        break
else:
    raise SystemExit('그림 문단을 찾지 못했다')

# ---------------------------------------------------------------- V-B Registration: 등록 문단 정정 + 새 소절 "user credential"
p = find('Registration happens once per user and once per service.')
rep(p, 'The AA also returns the user\'s attribute vector a_1..a_4 from its account record (Section V-D), and the wallet stores s_u, r_u, sk_u, and the attributes; the AA sees neither s_u nor r_u. The generators G_1, …, G_8, H are the nothing-up-my-sleeve base points of the circomlib Pedersen implementation on the Baby Jubjub curve [16], and G_3 and H are the same points that carry s_u and the blinding in the statement commitment below, so that equality of s_u across the two commitments can be proved with a sigma protocol.',
    'The AA also returns the user\'s attribute vector a_1..a_4 from its account record, and the wallet stores s_u, r_u, sk_u, and the attributes; the AA sees neither s_u nor r_u. The generators G_1, …, G_8, H are the nothing-up-my-sleeve base points of the circomlib Pedersen implementation on the Baby Jubjub curve [16], and G_3 and H are the same points that carry s_u and the blinding in the user credential below, so that equality of s_u across the two commitments can be proved with a sigma protocol.')
svc = find('A service registers by sending its display name')
# 새 소절을 서비스 등록 문단 뒤, "C. SERVICE AUTHENTICATION" 앞에 넣는다(제목 → 본문 → 수식 → 본문 → 수식 → 본문).
h = clone_after(svc, 'B\'. THE USER CREDENTIAL', template=HEAD)
q1 = clone_after(h, 'Before its first login the wallet has the AA verify, once, a commitment to what the AA knows about the user. It draws a blinding factor blind_u and computes the user credential', template=BODY)
e1 = clone_after(q1, 'C_u = uid·G_1 + s_u·G_3 + a_1·G_5 + a_2·G_6 + a_3·G_7 + a_4·G_8 + blind_u·H,   Cf_u = Poseidon(C_u.x, C_u.y),', template=EQ)
q2 = clone_after(e1, 'where uid, s_u, and blind_u lie in [0, 2^250) and the attributes a_1..a_4 in [0, 2^64); unused attribute slots are zero. The attributes are not chosen by the user. They are the AA\'s account record for uid, in the prototype a birth year, an ISO 3166 numeric country code, a level, and a spare slot, handed to the wallet at registration and re-fetched from the AA, under a request signed with sk_u, whenever the record changes; only an AA administrator changes the record. The wallet sends the AA (uid, C_u) with a Fiat-Shamir sigma protocol π_u for the relation', template=BODY)
e2 = clone_after(q2, 'PoK{ (s_u, blind_u, r_u) : C_u − uid·G_1 − Σ a_k·G_{4+k} = s_u·G_3 + blind_u·H  ∧  cm_u = s_u·G_3 + r_u·H },', template=EQ)
q3 = clone_after(e2, 'with uid, a_1..a_4, C_u, and cm_u public, where the AA takes uid from the authenticated session and a_1..a_4 from its own record, never from the request. The verifier subtracts uid·G_1 + Σ a_k·G_{4+k} from C_u and checks two Schnorr equations that share the response for s_u; a commitment built over any other identity, secret, or attribute values fails the check, since Pedersen commitments are binding. The proof shows that the credential contains the authenticated identity, the registered secret, and the recorded attributes, and nothing else about its contents. The AA computes Cf_u itself from the point it verified, records it as the account\'s active credential together with its revocation leaf leaf(Cf_u) of Section VI-A, and returns nothing but success; it keeps no copy of C_u\'s opening and never sees s_u or blind_u. A user has one active credential at a time. When an administrator changes the attribute record, the AA retires the active credential, that is, it enters leaf(Cf_u) into the next publication (Section VI), and the wallet obtains a new credential over the new values at its next login; the address is unaffected, since it depends on neither C_u nor the attributes. A revoked account is disabled at the AA and its credential leaf is published; a disabled account obtains no new credential.', template=BODY)

# ---------------------------------------------------------------- V-D Statement issuance → session statement
p = find('D. STATEMENT ISSUANCE WITH ATTESTED ATTRIBUTES')
set_text(p, 'D. STATEMENT ISSUANCE'); log(p, 'D. STATEMENT ISSUANCE WITH ATTESTED ATTRIBUTES', 'D. STATEMENT ISSUANCE')
p = find('The wallet generates a session key pair (pk_i, sk_i) on secp256k1')
rep(p, 'draws a blinding factor blind, and computes the commitment', 'draws a blinding factor blind_s, and computes the session commitment')
p = find('C_pt = uid·G_1 + arid·G_2 + s_u·G_3 + pk_i·G_4')
set_text(p, 'C_s = arid·G_2 + pk_i·G_4 + blind_s·H,   Cf_s = Poseidon(C_s.x, C_s.y),'); log(p, 'C_pt = …', 'C_s = …')
p = find('uid, arid, s_u, and blind lie in [0, 2^250), pk_i in [0, 2^160)')
set_text(p, 'where arid and blind_s lie in [0, 2^250) and pk_i in [0, 2^160). The session commitment carries what is new in this login, the service and the session key; the identity, the secret, and the attributes are already in the user credential C_u, which the AA verified once. The wallet chooses the expiry of the statement itself: it reads the chain head and sets max_height = ⌈(head + T) / G⌉·G, with a lifetime T of 300 blocks and a grid G of 100 blocks in the prototype, so that every statement requested within one grid window carries the same expiry (Section VII-C). It then sends the AA the tuple (uid, Cf_u, C_s, chainid, allowAgent, max_height) and a signature sig_u = Sign_sk_u(Poseidon(D_req, Cf_u, C_s.x, C_s.y, chainid, allowAgent, max_height)) with D_req a domain tag, so that no intermediary can alter the requested expiry or substitute a session commitment. No zero-knowledge proof accompanies the request: the identity and the secret were proved once when C_u was issued, and C_s need only be a well-formed point, since nothing about its opening is attested. Any arid or pk_i could be committed here, but the pseudonym proof recomputes C_s from the public arid and pk_i (Table 2, condition 2), so a statement over a session commitment for one service or key cannot be presented for another.')
log(p, 'uid, arid, s_u, and blind lie …', 'where arid and blind_s lie …')
p = find('PoK{ (arid, s_u, pk_i, blind, r_u) :')
remove(p)
p = find('with uid, a_1..a_4, C_pt, and cm_u public, where the AA takes uid')
set_text(p, 'The AA checks, in this order, the request format, that max_height is below 2^64, that the account exists and is not disabled, that chainid is in its allow list, sig_u under the stored pk_u, that Cf_u is the account\'s active credential, that C_s is a valid point of the prime-order subgroup, and that the chain answers; it re-reads the disabled flag and the active credential after the chain call, since that call is the one point at which the request yields, and signs. It neither chooses nor adjusts the expiry: as with the max_epoch of zkLogin, the expiry is the user\'s choice, the issuer binds it, and the verifier bounds it (Sections V-F and V-I). It computes Cf_s itself from the point it received and signs')
log(p, 'with uid, a_1..a_4, C_pt, and cm_u public …', 'The AA checks, in this order …')
p = find('σ_AA = Sign_AA(Poseidon(D_cred, C, max_height, chainid, allowAgent)),')
set_text(p, 'σ_AA = Sign_AA(Poseidon(D_cred, Cf_u, Cf_s, max_height, chainid, allowAgent)),'); log(p, 'σ_AA over C', 'σ_AA over Cf_u, Cf_s')
p = find('an EdDSA signature over Baby Jubjub with Poseidon as the message hash and D_cred a domain tag')
set_text(p, 'an EdDSA signature over Baby Jubjub with Poseidon as the message hash and D_cred a domain tag that separates statements from every other object the AA signs, chosen because it verifies cheaply inside the circuit. The AA returns (max_height, chainid, allowAgent, σ_AA) and records nothing: the only per-user state it holds is the active credential leaf of Section V-B\', which is what an account revocation publishes, and it holds no record of when or how often a user logged in. The three values outside the commitments are exactly those the AA has to check or set; the service identifier and the session key are inside C_s, and the identity, the secret, and the attributes inside C_u. We call the former AA attributes and the latter user attributes: AA attributes (the expiry, the chain, the agent flag) are set or verified by the AA and appear in the signed message in the clear; user attributes (in the prototype a birth year, a country code, and a level, together with the session key and the service) appear only inside the commitments. The AA never sees the service identifier or the session key; the attributes it verified against its record through π_u once, so every statement over that credential attests them, but no verifier sees them unless the user discloses a range of one (Section V-E), and the AA does not see where that happens.')
log(p, 'an EdDSA signature … records (uid, leaf(C), …)', 'an EdDSA signature … records nothing')
p = find('The AA never sees the challenge. A replayed issuance request obtains a fresh signature over the same commitment;')
set_text(p, 'The AA never sees the challenge. A replayed issuance request obtains a fresh signature over the same two commitments; the statement remains bound to the same session key, which the replayer does not hold, so no service and no contract accepts anything made from it, and the AA, which records nothing per issuance, is not even burdened with a duplicate. The only per-login state in the system is the service\'s own challenge record, which is what allows the AA to hold nothing that links a user to a login at all.')
log(p, 'The AA never sees the challenge …', '… records nothing per issuance')

# ---------------------------------------------------------------- V-E proof + Table 2
p = find('The wallet derives PPID = Poseidon(uid, s_u, chainid, arid), draws a fresh 250-bit randomness r')
rep(p, 'The witness is uid, s_u, blind, a_1..a_4, the signature σ_AA, r, and the non-membership path of Section VI. The circuit computes C_pt and C from the witness, so a prover cannot substitute a commitment the AA did not sign.',
    'The witness is uid, s_u, blind_u, blind_s, a_1..a_4, the signature σ_AA, r, and the non-membership path of Section VI. The circuit computes C_u from the witness and C_s from the public arid and pk_i, hashes both to Cf_u and Cf_s, and verifies σ_AA over them, so a prover cannot substitute a credential or a session commitment the AA did not sign, nor present a statement issued for one service or session key under another.')
p = find('Two range checks belong to the relation. Every committed scalar is constrained to 250 bits')
rep(p, 'which is what makes the commitment binding to one representative (Section IV-D)', 'which is what makes the commitments binding to one representative (Section IV-D)')
p = find('Condition 7 is selective disclosure.')
rep(p, 'Because the attributes inside C are those the AA verified against its record at issuance (Section V-D),', 'Because the attributes inside C_u are those the AA verified against its record when the credential was issued (Section V-B\'),')

for t in d.tables:
    heads = [r.cells[0].text.strip() for r in t.rows]
    if heads[0] == 'Symbol':
        set_row(t, 'a_1..a_4, blind', ['a_1..a_4, blind_u', 'user attributes (the AA\'s account record) and credential blinding', 'yes', 'a_1..a_4', 'disclosed ranges only'])
        set_row(t, 'C', ['C_u, Cf_u', 'user credential over (uid, s_u, a_1..a_4) and its field encoding Poseidon(C_u.x, C_u.y)', 'yes', 'Cf_u (verified once)', 'in proof only'])
        row_after(t, 'C_u, Cf_u', ['C_s, Cf_s, blind_s', 'session commitment over (arid, pk_i) and its field encoding; blinding', 'yes', 'C_s (per login)', 'in proof only'])
    if heads[0] == '#':
        set_row(t, '1', ['1', 'EdDSA.Verify(pk_AA, Poseidon(D_cred, Cf_u, Cf_s, max_height, chainid, allowAgent), σ_AA) = 1', 'the AA issued a statement over this credential and this session commitment for this chain, expiry, and agent flag'])
        set_row(t, '2', ['2', 'C_u = uid·G_1 + s_u·G_3 + Σ a_k·G_{4+k} + blind_u·H,  Cf_u = Poseidon(C_u);  C_s = arid·G_2 + pk_i·G_4 + blind_s·H,  Cf_s = Poseidon(C_s)', 'the credential opens to the hidden uid, s_u, attributes; the session commitment to the public arid and pk_i'])
        set_row(t, '4', ['4', 'leaf(Cf_u) = mask_252(Poseidon(4, Cf_u)), the low 252 bits, is not in the indexed Merkle tree with the given root', 'the user credential has not been revoked as of this root'])
        set_row(t, '5', ['5', 'r < 2^250,  c1 = r·B8,  K = r·pk_trace,  c2 = Poseidon(uid, arid) + Poseidon(K.x, K.y)', 'the tag encrypts a service-specific hash of this credential\'s uid under the certified combined key (verifiable encryption)'])
    if 'R1CS constraints' in heads:
        set_row(t, 'Issuance proof π_issue: prove / verify', ['User-credential proof π_u (once per user): prove / verify', '80 ms / 99 ms'])
        row_after(t, 'User-credential proof π_u (once per user): prove / verify', ['Statement issuance (per login, no proof): wallet round trip / AA work', '124 ms / 30 ms'])
        row_after(t, 'Statement issuance (per login, no proof): wallet round trip / AA work', ['Login end to end through the wallet (issuance + proof) / re-validation (cached proof)', '1,054 ms / 31 ms'])
        row_after(t, 'Login end to end through the wallet (issuance + proof) / re-validation (cached proof)', ['Wallet revocation-tree sync at 1,000 leaves: full replay / incremental, one new leaf / no change', '141 ms / 40 ms / 23 ms'])
        set_row(t, 'Root freshness bound: service / account contract', ['Root freshness bound: service (head age, root age) / account contract', '600 s, MAX_ROOT_AGE = 100 blocks / MAX_ROOT_AGE = 100 blocks, heartbeat every 50'])

# ---------------------------------------------------------------- V-F
p = find('The wallet sends the service the proof, its public inputs, and a signature σ = Sign_sk_i(r_s)')
rep(p, 'It reads the revocation root and block head from the chain, refusing if the head is older than ten minutes, and requires the root in the proof to equal the root it just read.',
    'It reads the revocation root, the block of its last publication, and the block head from the chain, refusing if the head is older than ten minutes, requires the root in the proof to equal the root it just read, and refuses if the root\'s last publication is older than MAX_ROOT_AGE blocks, the same bound the account contract applies (Section VI-C).')
p = find('A statement is scoped to the session that r_s names.')
rep(p, 'the wallet then computes a new proof for the same statement against the new root and re-validates.',
    'the wallet then applies the new leaves to its copy of the tree (Section VI-C), computes a new proof for the same statement against the new root, and re-validates.')

# ---------------------------------------------------------------- VI-A, VI-C, VI-D
p = find('The revocation set, which we also call the revoked-credential list (RCL)')
rep(p, 'Leaves are the 252-bit values leaf(C) = mask_252(Poseidon(3, C)). A leaf is inserted only when a statement is revoked, never at issuance, so the tree size is the number of revocations, not the number of logins, and the tree carries no record of who logged in where.',
    'Leaves are the 252-bit values leaf(Cf_u) = mask_252(Poseidon(4, Cf_u)), one per user credential, never per statement: a leaf is inserted only when a credential is revoked or retired, never at issuance, so the tree size is the number of revocations, not the number of logins or of users, and the tree carries no record of who logged in where. A service cannot recognize a leaf, since Cf_u appears in no transcript, and the AA cannot recognize a session from a leaf, since the leaf names a credential and not a login.')
p = find('The tree is append-only. Revoking one statement inserts its leaf.')
set_text(p, 'The tree is append-only. Revoking an account inserts one leaf, that of its active credential, and sets disabled = true at the AA; every statement issued over that credential, at every service, fails condition 4 against the next root, since all of them commit to the same Cf_u. Retiring a credential because its attribute record changed inserts the same leaf without disabling the account, and the next login issues a new credential. The two mechanisms are complementary: the tree kills statements already in the user\'s hands, which cannot be recalled, and the flag prevents new ones. Recovery from a compromise never removes a leaf, because removal would also revive the attacker; recovery issues a new credential, and the address is unchanged because it depends only on uid, s_u, chainid, and arid.')
log(p, 'The tree is append-only. Revoking one statement …', '… Revoking an account inserts one leaf …')
p = find('Services read the root and the block head directly from the chain and accept only proofs against the root they just read.')
rep(p, 'The account contract applies the same rule with the block number in place of the clock: it accepts only the log\'s current root and refuses if the last publication is older than MAX_ROOT_AGE blocks.',
    'The root has an age as well as an identity. The log records the block of its last publication and the AA republishes an unchanged root every H blocks (Section V-I); a service refuses a root whose last publication is older than MAX_ROOT_AGE blocks, exactly as the account contract does with the block number in place of the clock, so that an AA that stops publishing, and therefore stops publishing revocations, stops off-chain logins and on-chain transactions alike after MAX_ROOT_AGE blocks rather than never. An earlier version of the prototype applied the age bound only on chain and let off-chain logins continue against a stale root; the review that found this is reported in Section VIII-C.')
p = find('This is where the liveness goal G9 is met.')
set_text(p, 'This is where the liveness goal G9 is met. Nothing on the verification path involves the AA: the service reads the chain, the wallet rebuilds the tree from chain events, and re-validation after a root change needs only the statement the wallet already holds. If the AA is unreachable, new logins stop, because each login is a new issuance, but every session in progress continues off chain until the heartbeat lapses, and on chain accounts continue until then and stop MAX_ROOT_AGE blocks later, which is the price of a verifier that cannot read a clock and of a revocation list that is only as live as its publisher.')
log(p, 'This is where the liveness goal G9 is met …', '… until the heartbeat lapses …')
q = clone_after(p, 'The wallet does not replay the log from genesis at every login. It keeps the tree in memory, together with the block up to which it has applied the log, and persists the ordered list of leaves to a checkpoint file; a login fetches only the Revoked events since that block, inserts their leaves in the order they were published, and compares the resulting root with the root the contract holds at the same block. If they differ, whether because the checkpoint is corrupt, the chain was reset, or the log was redeployed, the wallet discards its copy and replays from genesis once, and if the replayed root still differs it refuses, as before. The checkpoint holds only leaves, which are public in the log\'s calldata, so it carries no secret and can be deleted at any time at the cost of one replay. Because the indexed tree is deterministic in the insertion order, the checkpoint needs no Merkle nodes, and because the root is checked after every application, the checkpoint is a cache and not a trust assumption. At a thousand leaves a login\'s sync falls from about 140 ms for a replay to about 40 ms when one leaf was published since the last login and about 25 ms when none was (Table 3).', template=BODY)
p = find('An administrator revokes a single statement, which the AA accepts only for leaves in its own issuance record, or an entire account.')
rep(p, 'An administrator revokes a single statement, which the AA accepts only for leaves in its own issuance record, or an entire account.',
    'An administrator revokes an account, or retires its credential by changing its attribute record; there is no revocation of a single statement, since the AA holds no record of statements (Section IX).')
rep(p, 'sets the flag immediately, and enters the leaves into the next publication.', 'sets the flag immediately, and enters the credential leaf into the next publication.')
p = find('No separate contract is needed for user-initiated revocation.')
rep(p, 'The two things a contract would have to check, the password and the set of leaves that belong to the account, cannot leave the AA: the password cannot be verified from public calldata, and the leaf set exists only in the AA\'s issuance record.',
    'The two things a contract would have to check, the password and the leaf that belongs to the account, cannot leave the AA: the password cannot be verified from public calldata, and the mapping from an account to its credential leaf exists only in the AA\'s record.')

# ---------------------------------------------------------------- VII
p = find('PPID is a collision-resistant hash of (uid, s_u, chainid, arid);')
rep(p, 'condition 2 of the relation opens C to some s_u, the issuance proof shows that this s_u is the one in cm_u,',
    'condition 2 of the relation opens C_u to some s_u, the credential proof π_u showed at issuance that this s_u is the one in cm_u,')
p = find('A service sees PPID, arid, pk_i, max_height, chainid, allowAgent, the root, pk_trace, the tag (c1, c2), and a Groth16 argument.')
rep(p, 'so it reveals nothing about uid, s_u, the attributes, or C beyond what the public inputs imply.', 'so it reveals nothing about uid, s_u, the attributes, blind_u, blind_s, or the commitments beyond what the public inputs imply.')
p = find('The AA receives (uid, C_pt, chainid, allowAgent), π_issue, and sig_u.')
set_text(p, 'At credential issuance the AA receives (uid, C_u) and π_u, and per statement (uid, Cf_u, C_s, chainid, allowAgent, max_height) and sig_u. The service identifier is inside C_s, whose hiding property rests on the blinding factor blind_s; the blinding is 250 bits rather than a full scalar, so hiding is computational rather than perfect and rests on the hardness of a bounded-interval discrete logarithm, which is as hard as the unbounded problem for this interval size. π_u is a zero-knowledge sigma protocol over C_u alone and is made before any service is involved, so it can reveal nothing about a service; sig_u is a signature over C_s and gives nothing beyond C_s. The certificate, the origin, and the address are never sent to the AA. The AA does learn chainid, since it must check it against its allow list; where the list has one entry this is no information, and where it has several the AA learns which chain the user is using. This leak is accepted and stated. The AA does not see the tag: it travels only from the wallet to the service inside the proof, and the AA\'s share is applied only at an approved opening. Because the AA keeps no record of issuances, what it can later join is only what it observes at the moment of issuance and what the chain shows (next paragraph).')
log(p, 'The AA receives (uid, C_pt, …), π_issue …', 'At credential issuance the AA receives (uid, C_u) and π_u …')
p = find('The chain is a second channel.')
rep(p, 'the commitment stays inside the circuit, the challenge is not signed,', 'both commitments stay inside the circuit, the challenge is not signed,')
rep(p, 'and the AA\'s issuance log narrows a transaction to the users issued in that window.', 'and an AA that logged its issuances could narrow a transaction to the users issued in that window; the prototype\'s AA keeps no such log, so it would have to observe issuances as they happen.')
p = find('The attributes and the session key occupy slots of the same commitment,')
rep(p, 'The attributes and the session key occupy slots of the same commitment, so the argument of the previous paragraph covers what a verifier sees of them:',
    'The attributes occupy slots of the user credential and the session key a slot of the session commitment, so the argument of the previous paragraph covers what a verifier sees of them:')
rep(p, 'since they are its own account record and π_issue verifies the commitment against that record;', 'since they are its own account record and π_u verified the credential against that record;')
rep(p, 'At issuance it sees neither the service nor the pseudonym (G4) and at login nothing at all,', 'At credential issuance it sees no service, at statement issuance neither the service nor the pseudonym (G4), and at login nothing at all,')
p = find('Once leaf(C) is in the tree and the root is published, condition 4 cannot be satisfied for that C against the new root:')
set_text(p, 'Once leaf(Cf_u) is in the tree and the root is published, condition 4 cannot be satisfied for any statement over that credential against the new root: the indexed tree requires a leaf (v, v_next) with v < leaf(Cf_u) < v_next, and no such gap exists once leaf(Cf_u) is itself a leaf, given collision resistance of Poseidon in the path and the 252-bit bound on the compared values. Every statement of the user, at every service, commits to the same Cf_u, so one leaf ends all of them. A service accepts only the root it just read and only while its last publication is younger than MAX_ROOT_AGE blocks, so the revoked credential is dead from the block in which the root is published, plus the service\'s read latency, which the ten-minute bound caps; on chain it is dead from that block, since the account contract compares against the log in the same state. A revoked account cannot obtain a new credential or a new statement, because the AA checks the flag and the active credential before signing. The AA cannot present different lists to different services, because the root is on a public chain and the leaves are in its calldata; a service or user who suspects misbehavior can rebuild the tree from the log. What the design does not protect against is an AA that publishes a forged root, which is outside the model, and a chain reorganization deeper than the service\'s read, which is a property of the chosen chain.')
log(p, 'Once leaf(C) …', 'Once leaf(Cf_u) …')

# ---------------------------------------------------------------- VIII
p = find('The prototype consists of four Node.js processes and five contracts:')
rep(p, 'The AA exposes registration, which returns the account\'s attributes, an endpoint from which the wallet re-fetches them, service registration with operator approval, issuance,',
    'The AA exposes registration, which returns the account\'s attributes, credential issuance, an endpoint from which the wallet re-fetches the attributes, service registration with operator approval, statement issuance,')
rep(p, 'The wallet agent holds the user\'s secrets, verifies service certificates, requests statements, derives addresses, produces and caches proofs,',
    'The wallet agent holds the user\'s secrets, or in a browser deployment receives them per request from a MetaMask Snap that keeps them, verifies service certificates, obtains its credential and requests statements, keeps a checkpoint of the revocation tree, derives addresses, produces and caches proofs,')
p = find('Two earlier versions of the circuit give the cost of each addition.')
rep(p, 'Narrowing the four attribute slots from 250 to 64 bits and adding the disclosure predicates of Section V-E brought the circuit to the 25,369 constraints of Table 3:',
    'Splitting the commitment into a user credential and a session commitment added a second Pedersen commitment and a second Poseidon hash and moved the AA signature to the pair, bringing it to 26,601; narrowing the four attribute slots from 250 to 64 bits and adding the disclosure predicates of Section V-E brought the circuit to the 25,369 constraints of Table 3:')
rep(p, 'a Poseidon commitment would be cheaper in the circuit but would force the issuance proof from a sigma protocol into a second SNARK, which is why Pedersen was kept.',
    'a Poseidon commitment would be cheaper in the circuit but would force the credential proof from a sigma protocol into a second SNARK, which is why Pedersen was kept.')
p = find('Per-login cost is one issuance round trip, about 0.18 s of sigma-protocol work split between wallet and AA,')
rep(p, 'Per-login cost is one issuance round trip, about 0.18 s of sigma-protocol work split between wallet and AA, plus one Groth16 proof of about 0.8 s and one ECDSA signature.',
    'Per-login cost is one issuance round trip of about 0.12 s through the wallet, of which the AA\'s own work, a signature verification, a subgroup check, and an EdDSA signature, is about 30 ms and the rest is two chain reads and HTTP, plus one Groth16 proof of about 0.8 s and one ECDSA signature; the credential proof π_u, about 0.18 s of sigma-protocol work split between wallet and AA, is paid once per user and again only after an attribute change. Before the split into two commitments the issuance round trip was 0.42 s, since every login carried π_u.')
rep(p, 'a login took 1.05 s (Groth16 proof 0.86 s), a re-validation with the cached proof 31 ms,', 'a login took 1.05 s (Groth16 proof 0.81 s, issuance 0.12 s, tree sync 0.03 s), a re-validation with the cached proof 31 ms,')
p = find('An end-to-end suite runs the AA, wallet agent, service, and chain in isolation')
rep(p, 'user registration; login with issuance and session creation;', 'user registration and credential issuance; login with statement issuance and session creation; a second login that reuses the credential and pays no credential proof;')
rep(p, 'a synchronized re-validation refused because the statement is revoked,', 'a synchronized re-validation refused because the credential is revoked,')
rep(p, 'administrator restoration followed by re-issuance and login under the same address;', 'administrator restoration followed by a new credential and login under the same address; an attribute change that retires the credential and a login that obtains a new one with the new values;')
q = clone_after(p, 'The tree checkpoint has its own suite: a delta applied to a checkpoint yields the root of a full replay; a second sync fetches events only from the block after the last one applied and none at all when the head has not moved; a heartbeat publication applies no leaf; a restarted wallet restores from the checkpoint; a checkpoint for another log, a corrupt file, an unknown version, or a leaf the tree rejects is discarded with a warning and replaced by a replay; a checkpoint whose leaves were altered so that its recorded root matches passes restoration but fails the comparison with the contract root and is replaced by a replay; a contract that lies about its root makes the replay fail closed; concurrent syncs share one in-flight update and observe no intermediate root, which an earlier version exposed to a concurrent login as a root the chain had never published; and a chain that was reset below the checkpoint\'s block forces a replay. A structured review of the whole prototype after these features were complete, one pass per subsystem against the scenarios of Sections V to VII, found no critical defect; its eight important findings, among them the off-chain root-age bound of Section VI-C, a duplicate-request key at the opening endpoint that omitted c2 and PPID and could have attributed one user\'s session to another when two transcripts shared the tag randomness, and a re-authorization dialog in the browser wallet whose agent-permission text came from the service rather than from the session, were fixed, and each fix is held by a regression test that reproduces its scenario.', template=BODY)
p = find('Publication cost was measured on the local Hardhat node with a freshly deployed log contract.')
rep(p, 'All figures are medians of ten runs measured on 2026-09-22 (contract suite: one run; individual runs differ by tens of gas with the proof encoding).',
    'All figures are medians of ten runs measured on 2026-09-22 (contract suite: one run; individual runs differ by tens of gas with the proof encoding); the tree-sync figures of Table 3 are medians of five runs measured on 2026-09-23 with leaves published fifty per transaction.')

# ---------------------------------------------------------------- IX Limitations
p = find('The AA is on the login path. Because each login is an issuance, an unreachable AA stops new logins;')
set_text(p, 'The AA is on the login path. Because each login is an issuance, an unreachable AA stops new logins; sessions in progress continue, off chain and on, only until the heartbeat lapses and MAX_ROOT_AGE blocks more, since a root that is no longer republished is one whose revocations are no longer being published either, and both verifiers refuse it (Section VI-C). A design that issued one statement per service and reused it across logins would remove the AA from the path at the cost of a longer-lived statement, and therefore a longer exposure between a revocation and the expiry that bounds it. The user credential already goes some way here: it is issued once and its proof is paid once, and what remains per login is a signature round trip.')
log(p, 'The AA is on the login path …', '… until the heartbeat lapses …')
p = find('Attribute predicates are ranges over single slots, with equality as the degenerate range,')
rep(p, 'and a change to the record retires the user\'s current statements and takes effect at the next issuance, which the same publication window exposes as a revocation.',
    'and a change to the record retires the user\'s credential and takes effect at the next login, which the same publication window exposes as a revocation.')
p = find('Revocation is a timing channel. Revoking an account inserts the leaves of all its unexpired statements in one publication,')
set_text(p, 'Revocation is a timing channel and has one granularity. Revoking an account inserts one leaf, its active credential, so every session of that user, at every service, stops at the same root transition; two services that compare which of their sessions a publication ended can link them, and the anonymity set of a revoked user is the number of leaves in that publication. Batching revocations over a longer window enlarges the set; the prototype publishes on command and does not batch for this purpose. The same granularity means that there is no revocation of a single session: an administrator who wants to end one service\'s session of a user ends all of them, and an attribute change ends them too. Per-session revocation would need the AA to keep a per-statement record, which is exactly what the credential removed.')
log(p, 'Revocation is a timing channel …', '… and has one granularity …')
p = find('Expiry is a block height. A service or contract that reads a lagging head accepts a statement for that lag')
rep(p, 'the AA keeps issuance records for fifty blocks past max_height so that account revocation still covers such statements, and a larger lag reopens the gap.',
    'since account revocation publishes the credential leaf and not per-statement leaves, it covers such statements regardless of the lag.')

# ---------------------------------------------------------------- X Conclusion
p = find('This paper separated two things that Web3 services conflate:')
rep(p, 'Delegating authentication to an authority that signs a commitment rather than an identity attaches those names to authenticated persons',
    'Delegating authentication to an authority that verifies a commitment to the identity once and thereafter signs commitments rather than identities attaches those names to authenticated persons')
p = find('The prototype shows that the whole flow fits in one Groth16 argument of about twenty-five thousand constraints,')
rep(p, 'reusing one statement across a session moves that cost to a revocation root the chain publishes.',
    'reusing one statement across a session moves that cost to a revocation root the chain publishes, and verifying the user\'s credential once rather than at every login moves the only zero-knowledge work the authority ever does out of the login path.')

# 2026-09-23 추가(사용자 질문 "지갑이 Cred_s 를 받으면 서비스 행세가 가능한가"에 대한 답을 VII-E 에 한 문장으로)
p = find('cert_s binds arid, the origin, and the combined trace key under the AA key,')
rep(p, 'the approval criteria are operational policy, not something the protocol checks.',
    'the approval criteria are operational policy, not something the protocol checks. The same certificate also keeps a wallet from posing as a service: holding a session credential gives the wallet nothing a service holds, neither sk_service, nor x_svc, nor a certificate for an origin of its own, and every verifier takes arid and pk_trace from its own registration rather than from the proof, so a credential issued for one service verifies only against that service\'s values.')

d.save(DST)
print('saved', DST)
for i, kind, before, after in LOG:
    b = (before[:70] + '…') if len(before) > 70 else before
    a = (after[:70] + '…') if len(after) > 70 else after
    print(f'[{i}] {kind}: {b!r} -> {a!r}')
