# zk-Delegation 원고 v3 → v4 (2026-09-22): 속성 선택 공개(V6) 반영.
# 스펙 docs/superpowers/specs/2026-09-22-mode3-selective-disclosure-design.md,
# 실측 results/mode3_disclosure_bench_20260922.md. v3 는 건드리지 않고 v4 사본에만 적용한다.
# 논문은 단일 커밋 C_pt / π_issue 표기를 쓰므로(이중 구조 C_u/C_s 미반영) 그 표기 안에서 V6 를 서술한다:
# 코드의 π_u V2 는 여기서 π_issue 의 개정판이고, 속성 슬롯 번호는 논문 관례(a_1..a_4)를 따른다(코드의 a₀..a₃).
import copy, shutil
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

SRC = 'documents/2026-09-18_v3_zk-Delegation_research_article.docx'
DST = 'documents/2026-09-22_v4_zk-Delegation_research_article.docx'
shutil.copy(SRC, DST)
d = Document(DST)
XMLNS = '{http://www.w3.org/XML/1998/namespace}space'
PARAS = list(d.paragraphs)                      # 프록시를 붙들어 두어야 id 가 안정된다
IDX = {id(p._p): i for i, p in enumerate(PARAS)}
DUMP = '/tmp/claude-1000/-home-node1phi-Desktop-google-oidc-demo/92bb2d9b-b323-4423-a306-f93d036ea1fc/scratchpad/paper_v3_dump.txt'
DIDX = {}
try:
    import re
    for line in open(DUMP, encoding='utf-8'):
        m = re.match(r'\[(\d+)\] \([^)]*\) (.*)$', line.rstrip('\n'))
        if m: DIDX.setdefault(m.group(2)[:60], m.group(1))
except OSError:
    pass
ORIG_TEXT = {id(p._p): p.text for p in PARAS}
LOG = []

def pidx(p):
    i = IDX.get(id(p._p), 'new')
    t = ORIG_TEXT.get(id(p._p), '')
    return DIDX.get(t[:60], i)   # 덤프 번호가 있으면 그것을 쓴다
def log(p, before, after, kind='edit'):
    LOG.append((pidx(p), kind, before, after))

def find(prefix, style=None):
    for p in PARAS:
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return p
    raise KeyError(prefix)

def set_text(p, text):
    rs = p.runs
    assert len(rs) >= 1, p.text[:40]
    rs[0].text = text
    for r in rs[1:]:
        r._r.getparent().remove(r._r)

def rep(p, old, new):
    t = p.text
    assert old in t, (old[:60], t[:80])
    set_text(p, t.replace(old, new, 1))
    log(p, old, new)

def clone_after(p, text, template=None):
    el = copy.deepcopy((template or p)._p)
    ts = el.findall('.//' + qn('w:t'))
    assert len(ts) >= 1
    ts[0].text = text
    ts[0].set(XMLNS, 'preserve')
    for t in ts[1:]:
        t.getparent().remove(t)
    p._p.addnext(el)
    q = Paragraph(el, p._parent)
    LOG.append((f'after {pidx(p)}', 'insert', '', text))
    return q

def set_cell(tc, text):
    ps_ = tc.findall(qn('w:p'))
    for q in ps_[1:]: tc.remove(q)
    q = ps_[0]
    rs = q.findall(qn('w:r'))
    for r in rs[1:]: q.remove(r)
    if not rs: rs = [etree.SubElement(q, qn('w:r'))]
    ts = rs[0].findall(qn('w:t'))
    for t in ts[1:]: rs[0].remove(t)
    if not ts: ts = [etree.SubElement(rs[0], qn('w:t'))]
    ts[0].text = text; ts[0].set(XMLNS, 'preserve')

def set_row(table, key, row):
    for r in table.rows:
        if r.cells[0].text.strip() == key:
            before = ' | '.join(c.text.strip() for c in r.cells)
            for tc, text in zip(r.cells, row): set_cell(tc._tc, text)
            LOG.append(('table', 'row', before, ' | '.join(row)))
            return
    raise KeyError(key)

def row_after(table, key, row):
    tbl = table._tbl
    for tr in tbl.findall(qn('w:tr')):
        if tr.findall(qn('w:tc'))[0].xpath('string(.)').strip() == key:
            new = copy.deepcopy(tr)
            for tc, text in zip(new.findall(qn('w:tc')), row): set_cell(tc, text)
            tr.addnext(new)
            LOG.append(('table', 'row-insert', f'after "{key}"', ' | '.join(row)))
            return
    raise KeyError(key)

PARA_T = find('A Web3 service must recognize a returning user')

# ---------------------------------------------------------------- Abstract
p = find('ABSTRACT ')
rep(p, 'it learns neither pseudonym, service, nor attributes and cannot recognize the user\'s transactions on the chain.',
    'it learns neither pseudonym nor service and cannot recognize the user\'s transactions on the chain. The commitment also carries attributes that the AA attests from its own account record, and the wallet can disclose a chosen range of any of them to a service or a contract inside the same proof, without the AA learning where.')
rep(p, 'A prototype\'s proof has about 25,600 constraints, proves in under a second, and verifies in milliseconds off chain or 390,000 gas on chain.',
    'A prototype\'s proof has about 25,400 constraints, proves in under a second, and verifies in milliseconds off chain or about 400,000 gas on chain.')

# ---------------------------------------------------------------- Introduction
p = find('zk-Delegation combines address abstraction with delegated authentication')
rep(p, 'the user secret, a session public key, and any user-managed attributes, together with',
    'the user secret, a session public key, and the user\'s attributes as the AA holds them on record, together with')
rep(p, 'checks that the committed identity and secret are the registered ones without opening the commitment',
    'checks that the committed identity, secret, and attributes are the registered ones without opening the commitment')
rep(p, 'The AA therefore sees neither the service nor the pseudonym nor the attributes; the pseudonym does not exist on the AA side at all, since the wallet derives it.',
    'The AA therefore sees neither the service nor the pseudonym; the pseudonym does not exist on the AA side at all, since the wallet derives it. The attributes it does know, since it attests them, but it does not learn where they are used: the wallet can disclose a chosen range of an attribute to a service or a contract inside the proof, and the AA takes no part in that disclosure.')

p = find('2) We construct zk-Delegation')
rep(p, 'learns neither the service nor the pseudonym nor the user\'s attributes, not even by reading the chain,',
    'learns neither the service nor the pseudonym, not even by reading the chain, and attests the user\'s attributes without learning where they are disclosed,')
p = find('3) We give a pseudonym proof')
rep(p, 'and that the same proof authorizes transactions from a contract account located by the pseudonym.',
    'that the same proof authorizes transactions from a contract account located by the pseudonym, and that it can disclose a chosen range of an attested attribute to a service or a contract.')
p = find('4) We implement a prototype')
rep(p, 'on-chain execution from pseudonym accounts, and authorized opening',
    'on-chain execution from pseudonym accounts, selective disclosure of an attested attribute to a contract, and authorized opening')
p = find('The remainder of the paper is organized as follows.')
rep(p, 'including service authentication, attribute privacy, the pseudonym proof,',
    'including service authentication, attribute attestation and selective disclosure, the pseudonym proof,')

# ---------------------------------------------------------------- Background (PairCT comparison)
p = find('Our earlier protocol PairCT applies pairwise pseudonyms')
rep(p, 'so it learns neither the user\'s attributes nor the service.',
    'so it learns neither the pseudonym nor the service.')

# ---------------------------------------------------------------- Goals G5
p = find('G5 Attribute privacy.')
rep(p, 'G5 Attribute privacy. The AA does not learn the user\'s attributes or the session key inside the commitment it signs.',
    'G5 Attribute-use unobservability. The AA, which records the user\'s attributes and attests them, does not learn to which service or in which transaction an attribute is disclosed, nor the session key inside the commitment it signs; a service or a contract learns only the ranges the user chooses to disclose and nothing about the other slots.')

# ---------------------------------------------------------------- Fig. 1 caption
p = find('FIGURE 1. One login')
rep(p, '(arid, cert_s, pk_trace, PPID, pk_i, tag, attributes, r_s)', '(arid, cert_s, pk_trace, PPID, pk_i, tag, r_s)')

# ---------------------------------------------------------------- V-B registration
p = find('Registration happens once per user and once per service.')
rep(p, 'The wallet stores s_u, r_u, sk_u, and any attributes the user enters; the AA sees none of these.',
    'The AA also returns the user\'s attribute vector a_1..a_4 from its account record (Section V-D), and the wallet stores s_u, r_u, sk_u, and the attributes; the AA sees neither s_u nor r_u.')

# ---------------------------------------------------------------- V-D issuance
p = find('D. STATEMENT ISSUANCE WITH ATTRIBUTE PRIVACY')
rep(p, 'D. STATEMENT ISSUANCE WITH ATTRIBUTE PRIVACY', 'D. STATEMENT ISSUANCE WITH ATTESTED ATTRIBUTES')

p = find('All scalars lie in [0, 2^250); unused attribute slots are zero.')
rep(p, 'All scalars lie in [0, 2^250); unused attribute slots are zero.',
    'uid, arid, s_u, and blind lie in [0, 2^250), pk_i in [0, 2^160), and the attributes a_1..a_4 in [0, 2^64); unused attribute slots are zero. The attributes are not chosen by the user. They are the AA\'s account record for uid, in the prototype a birth year, an ISO 3166 numeric country code, a level, and a spare slot, handed to the wallet at registration and re-fetched from the AA, under a request signed with sk_u, whenever the record changes; only an AA administrator changes the record, and a change retires the user\'s unexpired statements so that the next issuance carries the new values.')

p = find('PoK{ (arid, s_u, pk_i, a_1..a_4, blind, r_u) :')
rep(p, 'PoK{ (arid, s_u, pk_i, a_1..a_4, blind, r_u) : C_pt = uid·G_1 + arid·G_2 + s_u·G_3 + pk_i·G_4 + Σ a_k·G_{4+k} + blind·H  ∧  cm_u = s_u·G_3 + r_u·H },',
    'PoK{ (arid, s_u, pk_i, blind, r_u) : C_pt − uid·G_1 − Σ a_k·G_{4+k} = arid·G_2 + s_u·G_3 + pk_i·G_4 + blind·H  ∧  cm_u = s_u·G_3 + r_u·H },')

p = find('with uid, C_pt, and cm_u public.')
rep(p, 'with uid, C_pt, and cm_u public. It shows that the commitment contains the authenticated identity and the registered secret, and nothing else about its contents.',
    'with uid, a_1..a_4, C_pt, and cm_u public, where the AA takes uid from the authenticated session and a_1..a_4 from its own record, never from the request. The verifier subtracts uid·G_1 + Σ a_k·G_{4+k} from C_pt and checks two Schnorr equations that share the response for s_u; a commitment built over any other attribute values fails the check, since Pedersen commitments are binding. The proof shows that the commitment contains the authenticated identity, the registered secret, and the recorded attributes, and nothing else about its contents.')

p = find('σ_AA = Sign_AA(Poseidon(D_cred, C, max_height, chainid, allowAgent)),')
p = find('an EdDSA signature over Baby Jubjub with Poseidon as the message hash')
rep(p, 'The three values outside C are exactly those the AA has to check or set; everything the AA has no business seeing is inside C. We call the former AA attributes and the latter user attributes: AA attributes (the expiry, the chain, the agent flag) are set or verified by the AA and appear in the signed message in the clear; user attributes (for example a nickname, an e-mail address, an income level, or an age, together with the session key) are chosen by the user and appear only inside the commitment.',
    'The three values outside C are exactly those the AA has to check or set; the service identifier, the session key, and the attributes are inside C. We call the former AA attributes and the latter user attributes: AA attributes (the expiry, the chain, the agent flag) are set or verified by the AA and appear in the signed message in the clear; user attributes (in the prototype a birth year, a country code, and a level, together with the session key) appear only inside the commitment. The AA never sees the service identifier or the session key; the attributes it knows and verifies against its record through π_issue, so the signature attests them, but no verifier sees them unless the user discloses a range of one (Section V-E), and the AA does not see where that happens.')

# ---------------------------------------------------------------- V-E pseudonym proof
p = find('The wallet derives PPID = Poseidon(uid, s_u, chainid, arid), draws a fresh 250-bit randomness r')
rep(p, 'The public inputs are, in this order, PPID, arid, pk_i, max_height, chainid, allowAgent, root, pk_AA.x, pk_AA.y, pk_trace.x, pk_trace.y, c1.x, c1.y, c2.',
    'The public inputs are, in this order, PPID, arid, pk_i, max_height, chainid, allowAgent, root, pk_AA.x, pk_AA.y, pk_trace.x, pk_trace.y, c1.x, c1.y, c2, followed by the nine disclosure inputs disc_mask, disc_lo_1..disc_lo_4, disc_hi_1..disc_hi_4 of condition 7, 23 in all.')

p = find('Two range checks belong to the relation.')
rep(p, 'Every committed scalar is constrained to 250 bits, which is what makes the commitment binding to one representative (Section IV-D).',
    'Every committed scalar is constrained to 250 bits, and each attribute slot to 64 bits, which is what makes the commitment binding to one representative (Section IV-D) and what the comparators of condition 7 assume.')

p = find('The service compares pk_AA.x and pk_AA.y with the AA key it is configured with.')
q = clone_after(p,
    'Condition 7 is selective disclosure. The public input disc_mask ∈ [0, 16) names the attribute slots being disclosed, and for each slot k the public inputs disc_lo_k and disc_hi_k, both below 2^64, give the range claimed for it. For every slot whose mask bit is set the circuit enforces disc_lo_k ≤ a_k ≤ disc_hi_k; an exact value is disclosed by setting disc_lo_k = disc_hi_k. Slots whose bit is clear are unconstrained, and the wallet sets their bounds to zero. A login uses disc_mask = 0, and the proof is then exactly the statement of the preceding conditions; a transaction, or a login at a service that asks for it, carries a fresh proof with the mask and bounds the user chooses. Because the attributes inside C are those the AA verified against its record at issuance (Section V-D), a disclosed range is attested by the AA although the AA takes no part in the disclosure, and the undisclosed slots stay hidden by the hiding of the commitment and the zero knowledge of the argument. The predicates cost eight 64-bit range checks on the bounds and eight comparators, about 1,100 constraints, and the narrower attribute slots remove more than that from the commitment, so the circuit is slightly smaller than without them (Section VIII-B). The prototype supports ranges over single slots only; set membership and predicates that relate two slots or a slot to the current block are not supported (Section IX).',
    template=PARA_T)

# ---------------------------------------------------------------- V-F login checks
p = find('The wallet sends the service the proof, its public inputs, and a signature σ = Sign_sk_i(r_s)')
rep(p, 'It verifies the Groth16 argument, then σ under pk_i.',
    'It checks that disc_mask is below 16 and records the disclosure inputs with the session; the prototype\'s service asks for no disclosure at login, and a service policy that does is a matter of configuration. It verifies the Groth16 argument, then σ under pk_i.')

# ---------------------------------------------------------------- V-I on-chain execution
p = find('A transaction is execute(payload, σ, π, public inputs), submitted by any relayer')
rep(p, 'The wallet signs the payload under sk_i together with the chain identifier, the account address, and the account\'s nonce, so that neither another chain nor another account can reuse the signature;',
    'The wallet signs the payload under sk_i together with the chain identifier, the account address, the account\'s nonce, and the nine disclosure inputs of the proof it attaches, as keccak256(abi.encode(chainid, wallet, to, value, data, nonce, disc_mask, disc_lo, disc_hi)), so that neither another chain nor another account can reuse the signature and a relayer cannot pair the payload with another proof of the same session that discloses something else;')
rep(p, 'that allowAgent is a bit and c1 is not the identity point;',
    'that allowAgent is a bit, disc_mask is below 16, and c1 is not the identity point;')
rep(p, 'It then increments the nonce, performs the call, and emits the outcome together with pk_i, max_height, allowAgent, and the tag, so that a transaction is a transcript in the sense of Section V-G.',
    'It then increments the nonce, appends the nine disclosure inputs to the call data as trailing words in the manner of ERC-2771, performs the call, and emits the outcome together with pk_i, max_height, allowAgent, and the tag, so that a transaction is a transcript in the sense of Section V-G, and a Disclosure event when disc_mask is not zero.')
q = clone_after(p,
    'A target contract that acts on a disclosed attribute reads the last 288 bytes of its call data and requires, through the factory\'s isWallet mapping, that msg.sender is an account the factory deployed, since only such an account has verified the proof whose inputs it appends; a Solidity ABI decoder ignores trailing call data, so a target that knows nothing of disclosure is unaffected. The words are appended whether or not disc_mask is zero. Appending them only for a non-zero mask was considered and rejected: payload.data is chosen by the user, so a user could reuse a cached proof with mask 0 and place a forged (mask, lo, hi) tail at the end of payload.data, which the target would then read as attested; this was found in review and reproduced against the demonstration target. With the tail always present, the last 288 bytes of the call are the inputs of the proof the account has just verified, and a forged tail inside payload.data lies before them. The one exception is a call with empty payload.data and mask 0, a plain transfer: no tail is appended, so that a receive-only contract, whose receive function runs only on empty call data, still accepts the transfer; a forged tail would make the data non-empty, which restores the rule. The prototype\'s target, AttrGate, accepts one claim per account if slots 1 and 2 are disclosed with a_2 equal to a configured country code and the upper bound of a_1 at most a configured birth year, that is, a residency and an age check.',
    template=PARA_T)

p = find('The proof is the session\'s proof. The wallet re-submits it with every transaction')
rep(p, 'verification, however, is paid at every transaction, about 390,000 gas (Section VIII-C)',
    'verification, however, is paid at every transaction, about 400,000 gas (Section VIII-C)')
p = find('Off-chain login and on-chain execution share one circuit and one statement.')
rep(p, 'Section VII-C shows that this gives the AA nothing beyond the chain identifier it already knows and a coarse expiry.',
    'Section VII-C shows that this gives the AA nothing beyond the chain identifier it already knows and a coarse expiry, and Section VII-D what a disclosed range adds.')

# ---------------------------------------------------------------- VII-C chain channel
p = find('The chain is a second channel.')
rep(p, 'What remains matchable is (chainid, allowAgent, max_height).',
    'What remains matchable is (chainid, allowAgent, max_height) and, for a transaction that discloses an attribute range, that range (Section VII-D).')

# ---------------------------------------------------------------- VII-D G5
p = find('D. ATTRIBUTE PRIVACY (G5)')
rep(p, 'D. ATTRIBUTE PRIVACY (G5)', 'D. ATTRIBUTE-USE UNOBSERVABILITY (G5)')
p = find('Attributes and the session key occupy slots of the same commitment')
old = p.text
set_text(p, 'The attributes and the session key occupy slots of the same commitment, so the argument of the previous paragraph covers what a verifier sees of them: the commitment is hiding and the argument is zero-knowledge, so a service or a contract learns nothing about a slot the user has not disclosed, and about a disclosed slot exactly the range in the public inputs. What attestation changes is the AA\'s view. The AA knows the attribute values, since they are its own account record and π_issue verifies the commitment against that record; what it does not learn is where they are used. At issuance it sees neither the service nor the pseudonym (G4) and at login nothing at all, and a disclosure reaches the AA only if it appears on the chain, where it is a public input of a transaction from an account the AA cannot locate. The property is therefore attribute-use unobservability rather than attribute confidentiality: the AA cannot tell to which service, or in which transaction, a given user disclosed a given slot, beyond what the next paragraph concedes. The session key is as before: it is inside the commitment, never disclosed, and unknown to the AA.')
log(p, old, p.text)
q = clone_after(p,
    'A disclosed range shrinks the anonymity set of Section VII-C. The AA holds every user\'s attributes, so from a transaction that discloses ranges it can compute the set of users whose recorded attributes satisfy them; intersected with the users issued in the same (chainid, allowAgent, max_height) window, this set, rather than the window alone, is the anonymity set of that transaction against the AA. An exact value (disc_lo = disc_hi) shrinks the set more than a range does, and the disclosed values are public for as long as the chain is, bound to the account and thus to the service. This is the price of attestation, and the user pays it only for the slots and the transactions where a predicate is required.',
    template=PARA_T)

# ---------------------------------------------------------------- VIII-A prototype
p = find('The prototype consists of four Node.js processes and four contracts')
rep(p, 'four Node.js processes and four contracts: the revocation log, the Groth16 verifier generated from the circuit, a per-service account factory, and the account contract it creates.',
    'four Node.js processes and five contracts: the revocation log, the Groth16 verifier generated from the circuit, a per-service account factory, the account contract it creates, and a demonstration target, AttrGate, that reads disclosed attributes.')
rep(p, 'The AA exposes registration, service registration with operator approval, issuance, revocation, publication, the opening endpoints, an administrator page with approval controls, and a user account page with self-revocation.',
    'The AA exposes registration, which returns the account\'s attributes, an endpoint from which the wallet re-fetches them, service registration with operator approval, issuance, revocation, publication, the opening endpoints, an administrator page with approval controls and per-account attribute editing, and a user account page with self-revocation.')
rep(p, 'produces and caches proofs, and signs session requests; a browser page drives it.',
    'produces and caches proofs, signs session requests, and takes a per-transaction disclosure choice, a lower and an upper bound per slot, which it checks against its own attributes before proving; a browser page drives it.')
rep(p, 'The trusted setup uses a public powers-of-tau file of size 2^21.',
    'The trusted setup uses a public powers-of-tau file of size 2^21 and a single phase-2 contribution made on the development machine (Section IX).')

# ---------------------------------------------------------------- VIII-B cost
p = find('Two earlier versions of the circuit give the cost of each addition.')
rep(p, 'The Pedersen commitment remains the largest single block;',
    'Narrowing the four attribute slots from 250 to 64 bits and adding the disclosure predicates of Section V-E brought the circuit to the 25,369 constraints of Table 3: the predicates add about 1,100 constraints and the narrower attribute scalars remove more, so the circuit is slightly smaller than before; we did not itemize the difference further. The Pedersen commitment remains the largest single block;')

p = find('Per-login cost is one issuance round trip')
rep(p, 'about 0.3 s of sigma-protocol work split between wallet and AA, plus one Groth16 proof of about 0.9 s',
    'about 0.18 s of sigma-protocol work split between wallet and AA, plus one Groth16 proof of about 0.8 s')
rep(p, 'plus 340,000 to 390,000 gas for verification and execution depending on the inner call, sixteen to eighteen times a plain transfer; this is the cost of verifying at every transaction that Section V-I accepts.',
    'plus 402,000 gas for verification and a call without value, 451,000 for a value transfer to a fresh address, and 445,000 for a call that discloses two attribute ranges to the demonstration target, nineteen to twenty-one times a plain transfer; this is the cost of verifying at every transaction that Section V-I accepts. Of this, the nine disclosure inputs cost about 62,800 gas in the verifier over the 14-input version, and the disclosure path, the trailing words, the event, and the target\'s own checks, about 42,400 over a call without disclosure.')
rep(p, 'a login with issuance took 1.39 s (issuance 0.42 s, proof 0.90 s), a re-validation with the cached proof 33 ms, and an on-chain transaction with the cached proof 157 ms including mining on the local node.',
    'a login took 1.05 s (Groth16 proof 0.86 s), a re-validation with the cached proof 31 ms, an on-chain transaction with the cached proof 151 ms including mining on the local node, and a transaction that discloses two attribute ranges, and therefore needs a fresh proof, 1.00 s.')

# ---------------------------------------------------------------- VIII-C validation
p = find('An end-to-end suite runs the AA, wallet agent, service, and chain in isolation')
rep(p, 'and a tag whose ciphertext, key, or randomness has been altered.',
    'and a tag whose ciphertext, key, or randomness has been altered; for disclosure it checks a mask of zero, an exact value, a range, a value outside the range, a bound of 2^64 or more, and a mask of 16 or more, the last three failing at witness generation. The end-to-end suite further covers selective disclosure: a claim at AttrGate by a user whose recorded country and birth year satisfy its policy; a second claim by the same account, refused; the same claim by a user with another country, refused inside the call with the nonce consumed; a range that the user\'s own attribute does not satisfy, refused by the wallet before any proof is made; an administrator change of one attribute, after which the wallet\'s next issuance fails, re-fetches the record, re-issues, and logs in under the same address; and a login whose disclosure the service records with the session. The contract suite adds the 23-input verifier, the trailing disclosure words and the Disclosure event, a signature whose digest omits the disclosure inputs, a plain transfer to a receive-only contract, and AttrGate\'s own checks: acceptance, a country mismatch, an age mismatch, and a direct call from an externally owned account.')

p = find('Publication cost was measured on the local Hardhat node')
rep(p, 'An execute() call with a valid proof consumed 387,905 gas for a value transfer to a fresh address and 339,341 gas for a call without value to an existing account, the bulk of it the pairing check of the Groth16 verifier; deploying an account through the factory consumed 713,542 gas once; a heartbeat is a publication with no leaves and consumed 40,261 gas.',
    'An execute() call with a valid proof and no disclosure consumed 402,130 gas for a call without value to an existing account and 419,262 gas for the first transaction of an account; the contract suite\'s value transfer to a fresh address consumed 450,675 gas; and a call disclosing two attribute ranges to AttrGate, including its claim logic, consumed 444,509 gas; in every case the bulk is the pairing check of the Groth16 verifier. Deploying an account through the factory consumed 852,198 gas once; a heartbeat is a publication with no leaves and consumed 40,285 to 40,305 gas. All figures are medians of ten runs measured on 2026-09-22 (contract suite: one run; individual runs differ by tens of gas with the proof encoding).')

# ---------------------------------------------------------------- IX limitations
p = find('The AA sees the chain identifier and, through the chain, the quantized expiry of every transaction')
rep(p, 'and a small deployment gives a transaction a small anonymity set against the AA.',
    'and a small deployment gives a transaction a small anonymity set against the AA. A transaction that discloses attribute ranges narrows that set further, to the users of the window whose recorded attributes satisfy the disclosed predicate, since the AA holds every user\'s attributes; an exact value narrows it more than a range, and the disclosed values remain public on the chain (Section VII-D).')

p = find('Attributes are hidden but not attested.')
old = p.text
set_text(p, 'Attribute predicates are ranges over single slots, with equality as the degenerate range, over a 64-bit integer encoding of each attribute. A predicate relative to the current time, such as an age computed from a birth year at the block in which the proof is verified, must be expressed by the wallet as a bound on the stored value and goes stale as time passes; set membership and predicates across slots are not supported. The attribute record is the AA\'s: a user cannot carry attributes the AA does not hold, and a change to the record retires the user\'s current statements and takes effect at the next issuance, which the same publication window exposes as a revocation.')
log(p, old, p.text)

p = find('Service registration is approved by an operator, but the criteria for approval are operational policy')
q = clone_after(p,
    'Revocation is a timing channel. Revoking an account inserts the leaves of all its unexpired statements in one publication, so every session of that user, at every service, stops at the same root transition; two services that compare which of their sessions a publication ended can link them, and the anonymity set of a revoked user is the number of leaves in that publication. Batching revocations over a longer window enlarges the set; the prototype publishes on command and does not batch for this purpose.',
    template=PARA_T)

p = find('Expiry is a block height.')
rep(p, 'The prototype\'s wallet asks for no confirmation before a transaction, and its relayer sees every transaction it submits.',
    'The prototype\'s wallet asks for no confirmation before a transaction, and its relayer sees every transaction it submits. The endpoint from which the wallet re-fetches its attributes checks the user\'s signature but not the freshness of the request, so a captured request can be replayed to read the same attributes again; and the service\'s session and login listings, which now include the disclosed ranges, are unauthenticated demonstration interfaces.')

p = find('The commitment is computationally rather than perfectly hiding')
rep(p, 'and Groth16 requires a trusted setup, whose compromise would allow forged proofs.',
    'and Groth16 requires a trusted setup, whose compromise would allow forged proofs; the prototype\'s phase-2 contribution is a single contribution with fixed entropy made on the development machine, so the operator of that machine could forge proofs, and a deployment needs a multi-party ceremony.')

# ---------------------------------------------------------------- Conclusion
p = find('This paper separated two things that Web3 services conflate')
rep(p, 'while keeping the service, the address, and the user\'s attributes out of the authority\'s view,',
    'while keeping the service and the address out of the authority\'s view and the use of the user\'s attributes unobservable to it,')
p = find('The prototype shows that the whole flow fits in one Groth16 argument')
rep(p, 'and one verification per transaction on chain, at about 390,000 gas.',
    'and one verification per transaction on chain, at about 400,000 gas; the same argument discloses a chosen range of an attested attribute to a contract for about 42,000 gas more.')
rep(p, 'Future work is to add attribute predicate proofs over the committed slots,',
    'Future work is to extend the range predicates over the committed slots to set membership and time-relative predicates,')

# ---------------------------------------------------------------- Tables
for t in d.tables:
    heads = [r.cells[0].text.strip() for r in t.rows]
    if 'a_1..a_4, blind' in heads:
        set_row(t, 'a_1..a_4, blind', ['a_1..a_4, blind', 'user attributes (the AA\'s account record) and commitment blinding', 'yes', 'a_1..a_4', 'disclosed ranges only'])
    if '6' in heads and heads[0] == '#':
        row_after(t, '6', ['7', 'disc_mask < 16;  for each k with bit k of disc_mask set:  disc_lo_k ≤ a_k ≤ disc_hi_k,  with a_k, disc_lo_k, disc_hi_k < 2^64',
                           'the disclosed slots lie in the ranges given as public inputs; undisclosed slots are unconstrained'])
    if 'R1CS constraints' in heads:
        set_row(t, 'R1CS constraints', ['R1CS constraints', '25,369'])
        set_row(t, 'Public inputs / private inputs', ['Public inputs / private inputs', '23 / 79'])
        set_row(t, 'Proving key (zkey) / witness generator (wasm)', ['Proving key (zkey) / witness generator (wasm)', '15.4 MB / 4.7 MB'])
        set_row(t, 'Groth16 proving time (median)', ['Groth16 proving time (median)', '809 ms'])
        set_row(t, 'Groth16 verification time (off chain)', ['Groth16 verification time (off chain)', '10.7 ms'])
        set_row(t, 'execute() gas on chain (verification + call)', ['execute() gas on chain (verification + call)', '402,130 (call) / 450,675 (value transfer) / 444,509 (call disclosing two slots, AttrGate)'])
        set_row(t, 'Issuance proof π_issue: prove / verify', ['Issuance proof π_issue: prove / verify', '80 ms / 99 ms'])
        row_after(t, 'Issuance proof π_issue: prove / verify', ['Disclosure inputs / verifier gas over the 14-input version', '9 (mask, 4 lower, 4 upper bounds) / about 62,800'])

d.save(DST)
print('saved', DST)
for i, kind, before, after in LOG:
    b = (before[:70] + '…') if len(before) > 70 else before
    a = (after[:70] + '…') if len(after) > 70 else after
    print(f'[{i}] {kind}: {b!r} -> {a!r}')
